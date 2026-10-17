"""Excel quantity report. Direct cost only. No price book and no routing."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from adel_takeoff.confidence import ConfidenceThresholds
from adel_takeoff.constants import (
    CONTRACTOR,
    DISCLAIMER,
    DISCREPANCY_CONFIDENCE_CAP,
    LABOR_RATES,
)
from adel_takeoff.costing import journeyman_warning
from adel_takeoff.present import build_review


def _number(value: Decimal | None):
    if value is None:
        return None
    return float(value)


def _score(value):
    if value is None or value == "":
        return None
    return float(value)


QUANTITY_HEADERS = [
    "Line",
    "Item IDs",
    "Item",
    "Assembly",
    "Assembly version",
    "Assembly labor hours",
    "Assembly confirmed",
    "Category",
    "Scope",
    "Measured quantity",
    "Count correction",
    "Effective quantity",
    "Sheet number",
    "Sheet title",
    "PDF page",
    "Location",
    "How counted",
    "Confidence",
    "Detector score",
    "Sheets",
    "Panel",
    "Circuit",
    "Type",
    "Type description",
    "Mounting",
    "Mounting height",
    "Voltage",
    "Circuit description",
    "Breaker",
    "Poles",
    "Load",
    "Review flag",
    "Deduped",
    "Material unit cost",
    "Labor hours",
    "Labor rate",
    "Direct cost",
]


def build_workbook(job: dict, items, allowances, rfis, claims) -> bytes:
    review = build_review(job, items, allowances, rfis, claims)
    workbook = Workbook()
    workbook.properties.creator = CONTRACTOR
    workbook.properties.title = "Quantity takeoff"
    sheet = workbook.active
    sheet.title = "Quantities"

    sheet["A1"] = CONTRACTOR
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = "Quantity takeoff"
    sheet["A3"] = DISCLAIMER
    sheet["A4"] = f"Job: {job['name']}"
    permit = job.get("permit_issue_date") or ""
    sheet["A5"] = (
        f"Permit-issue date: {permit or 'not set'}. "
        f"Code edition basis: {job['code_basis']}. "
        f"{job['code_fbc']}. {job['code_nec']}."
    )
    if review["totals"]["markup"] is None:
        sheet["A6"] = "Markup: not applied"
    else:
        sheet["A6"] = f"Markup: {format(review['totals']['markup'], 'f')}"
    warning = journeyman_warning(Decimal(job["rate_journeyman"]))
    sheet["A7"] = warning
    sheet["A8"] = (
        f"Power devices: {review['totals']['power_devices']}. "
        f"Lighting fixtures: {review['totals']['lighting_fixtures']}."
    )
    limits = ConfidenceThresholds.from_config()
    sheet["A9"] = (
        f"Confidence: high at or above {limits.high_min}, "
        f"check at or above {limits.check_min}, otherwise low. "
        "Blank means the detector recorded no score. "
        f"A schedule discrepancy lowers a recorded flag to {DISCREPANCY_CONFIDENCE_CAP} "
        "and does not create a flag when there is no score."
    )
    sheet["A10"] = (
        "Numbered marks on the PDF use the item id. "
        "A quantity line with one item uses that same id."
    )
    sheet["A11"] = (
        "Fixture and panel schedule cells are copied onto the quantity lines. "
        "A blank schedule cell was not printed. The effective quantity stays the plan count. "
        "Where a schedule count differs, both numbers are on the Discrepancies sheet."
    )

    sheet["A12"] = review.get("assembly_warning") or (
        "No counted line uses an assembly."
    )
    sheet["A13"] = "Labor rates"
    rate_row = 14
    stored = {
        "journeyman": job["rate_journeyman"],
        "foreman": job["rate_foreman"],
        "pm": job["rate_pm"],
        "apprentice": job["rate_apprentice"],
    }
    for key, label, _default in LABOR_RATES:
        sheet.cell(rate_row, 1, label)
        sheet.cell(rate_row, 2, float(Decimal(stored[key])))
        sheet.cell(rate_row, 3, "per hour")
        rate_row += 1

    header_row = rate_row + 1
    for column, header in enumerate(QUANTITY_HEADERS, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(bold=True)

    row_index = header_row + 1
    rate_labels = {key: label for key, label, _default in LABOR_RATES}
    for line in review["lines"]:
        _write_quantity_row(sheet, row_index, line, rate_labels)
        row_index += 1

    row_index += 1
    sheet.cell(row_index, 1, "Allowances").font = Font(bold=True)
    row_index += 1
    for column, header in enumerate(
        ["Description", "Quantity", "Reason", "Category", "Material unit cost", "Labor hours", "Labor rate", "Direct cost"],
        start=1,
    ):
        sheet.cell(row_index, column, header).font = Font(bold=True)
    row_index += 1
    for allowance in review["allowances"]:
        values = [
            allowance["description"],
            _number(allowance["quantity"]),
            allowance["reason"],
            allowance["category"],
            _number(_decimal_or_none(allowance["material_unit_cost"])),
            _number(_decimal_or_none(allowance["labor_hours"])),
            rate_labels.get(allowance["labor_rate_key"], allowance["labor_rate_key"]),
            _number(allowance["direct_cost"]),
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)
        row_index += 1

    row_index += 1
    sheet.cell(row_index, 1, "Job direct cost")
    sheet.cell(row_index, 2, _number(review["totals"]["direct_cost"]))
    row_index += 1
    sheet.cell(row_index, 1, "Marked-up total")
    if review["totals"]["marked_up_total"] is None:
        sheet.cell(row_index, 2, "not applied")
    else:
        sheet.cell(row_index, 2, _number(review["totals"]["marked_up_total"]))

    excluded = workbook.create_sheet("Excluded")
    excluded_headers = [
        "Item ID",
        "Item",
        "Category",
        "Sheet number",
        "Sheet title",
        "PDF page",
        "Location",
        "Reason",
        "How counted",
        "Confidence",
        "Detector score",
        "Panel",
        "Circuit",
    ]
    for column, header in enumerate(excluded_headers, start=1):
        cell = excluded.cell(1, column, header)
        cell.font = Font(bold=True)
    for index, line in enumerate(review["excluded"], start=2):
        values = [
            ", ".join(line["item_ids"]),
            line["label"],
            line["category"],
            line["sheet_number"],
            line["sheet_title"],
            line["pdf_page"],
            line["location"],
            line["excluded_reason"],
            line["how_counted"],
            line["confidence"] or None,
            _score(line["detector_score"]),
            line["panel"],
            line["circuit"],
        ]
        for column, value in enumerate(values, start=1):
            excluded.cell(index, column, value if value != "" else None)

    rfi_sheet = workbook.create_sheet("RFIs")
    rfi_sheet.append(["Topic", "Disposition", "Sections", "Detail"])
    for rfi in rfis:
        rfi_sheet.append([rfi["topic"], rfi["disposition"], rfi["sections"], rfi["detail"]])

    claim_sheet = workbook.create_sheet("Code claims")
    claim_sheet.append(["Statement", "Section", "Rendered"])
    for claim in claims:
        claim_sheet.append([claim["statement"], claim["section"], claim["rendered"]])

    _write_by_type(workbook, review["by_type"])
    _write_discrepancies(workbook, review["discrepancies"])
    _write_materials(workbook, review["materials"])
    _write_labor(workbook, job, review["labor"], review.get("assembly_warning") or "")
    _write_measured(workbook, review.get("scales") or [], review.get("measurements") or [])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_quantity_row(sheet, row_index: int, line: dict, rate_labels: dict) -> None:
    rate_key = line["labor_rate_key"]
    values = [
        line["line_id"],
        ", ".join(line["item_ids"]),
        line["label"],
        line.get("assembly_name") or None,
        line.get("assembly_version") or None,
        _number(_decimal_or_none(line.get("assembly_labor_hours"))),
        "yes" if line.get("assembly_confirmed") else ("no" if line.get("assembly_id") else None),
        line["category"],
        line["scope"],
        line["measured"],
        line["override"],
        line["effective"],
        line["sheet_number"],
        line["sheet_title"] or None,
        line["pdf_page"],
        line["location"],
        line["how_counted"],
        line["confidence"] or None,
        _score(line["detector_score"]),
        ", ".join(line["sheets"]),
        line["panel"],
        line["circuit"],
        line.get("type_mark") or None,
        line.get("type_description") or None,
        line.get("mounting") or None,
        line.get("mounting_height") or None,
        line.get("voltage") or None,
        line.get("circuit_description") or None,
        line.get("breaker") or None,
        line.get("poles") or None,
        line.get("schedule_load") or None,
        line["link_flag"],
        "counted once" if line["deduped"] else "",
        _number(_decimal_or_none(line["material_unit_cost"])),
        _number(_decimal_or_none(line["labor_hours"])),
        rate_labels.get(rate_key, rate_key) if rate_key else None,
        _number(line["direct_cost"]),
    ]
    for column, value in enumerate(values, start=1):
        sheet.cell(row_index, column, value if value is not None else None)


def _write_by_type(workbook: Workbook, rows: list[dict]) -> None:
    sheet = workbook.create_sheet("By type")
    sheet["A1"] = (
        "Each row is one fixture type on one sheet. "
        "Plan count on sheet counts devices drawn on that sheet. "
        "Plan count all sheets counts each device once. "
        "Adding the per-sheet counts can count an overlap device twice. "
        "The schedule count is the number printed for the type, not a per-sheet number. "
        "Comparison uses the unique device count."
    )
    headers = [
        "Type",
        "Type description",
        "Sheet number",
        "Floor",
        "Plan count on sheet",
        "Plan count all sheets",
        "Schedule count",
        "Voltage",
        "Mounting",
        "Mounting height",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(3, column, header)
        cell.font = Font(bold=True)
    for index, row in enumerate(rows, start=4):
        values = [
            row["type"],
            row["type_description"] or None,
            row["sheet_number"],
            row["floor"] or None,
            row["plan_count_on_sheet"],
            row["plan_count_all_sheets"],
            row["schedule_count"],
            row["voltage"] or None,
            row["mounting"] or None,
            row["mounting_height"] or None,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(index, column, value if value is not None else None)


def _write_discrepancies(workbook: Workbook, rows: list[dict]) -> None:
    sheet = workbook.create_sheet("Discrepancies")
    sheet["A1"] = (
        "Both numbers are recorded. A schedule count does not replace the plan count."
    )
    headers = ["Kind", "Detail", "Plan", "Schedule", "Item IDs", "Sheet number"]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(2, column, header)
        cell.font = Font(bold=True)
    for index, row in enumerate(rows, start=3):
        values = [
            row["kind"],
            row["detail"],
            row["plan"] if row["plan"] != "" else None,
            row["schedule"] if row["schedule"] != "" else None,
            ", ".join(row["item_ids"]) or None,
            row["sheet_number"] or None,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(index, column, value if value is not None else None)


def _write_materials(workbook: Workbook, rows: list[dict]) -> None:
    sheet = workbook.create_sheet("Materials")
    sheet["A1"] = (
        "Part quantity is the takeoff count times the assembly per-unit quantity. "
        "Raceway and conductor lengths are allowances, not measured lengths. "
        "A measured length is on the Measured sheet and does not replace an allowance. "
        "Unconfirmed rows are examples and are not a published labor standard."
    )
    headers = [
        "Assembly",
        "Assembly version",
        "Part",
        "Description",
        "Unit",
        "Size",
        "Conductor count",
        "Allowance",
        "Quantity",
        "Source item IDs",
        "Confirmed",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(3, column, header)
        cell.font = Font(bold=True)
    for index, row in enumerate(rows, start=4):
        values = [
            row["assembly"],
            row["assembly_version"],
            row["role"],
            row["description"],
            row["unit"],
            row["size"] or None,
            row["conductor_count"] or None,
            row["allowance"] or None,
            _number(row["quantity"]),
            ", ".join(row["item_ids"]),
            "yes" if row["confirmed"] else "no",
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(index, column, value if value is not None else None)


def _write_measured(workbook: Workbook, sheets: list[dict], measurements: list[dict]) -> None:
    sheet = workbook.create_sheet("Measured")
    sheet["A1"] = (
        "A length is measured only where the sheet has one valid scale. "
        "Missing, NTS, ambiguous, and disagreeing scales are refused. "
        "Measured lengths stay off the assembly allowance lines. "
        "There is no routing and there are no home runs."
    )
    headers = [
        "Sheet",
        "PDF page",
        "Status",
        "Scale",
        "Source",
        "Read from",
        "Applies to",
        "Reason",
        "Set by",
        "Set at",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(3, column, header)
        cell.font = Font(bold=True)
    row_index = 4
    for page in sheets:
        regions = page["regions"]
        detected = page["detected"] if page.get("overridden") else []
        for region in regions:
            read_from = "; ".join(
                (
                    f"{item['kind']} {item['bbox'][0]:.2f},{item['bbox'][1]:.2f} to {item['bbox'][2]:.2f},{item['bbox'][3]:.2f}"
                    if item.get("bbox")
                    else item["kind"]
                )
                for item in region.get("read_from") or []
            )
            applies = "sheet"
            if region.get("viewport"):
                box = region["viewport"]
                applies = f"viewport {box[0]:.2f},{box[1]:.2f} to {box[2]:.2f},{box[3]:.2f}"
            values = [
                page["sheet"],
                page["page_index"] + 1,
                region["status"],
                region.get("label") or None,
                region.get("source") or None,
                read_from or None,
                applies,
                region.get("reason") or None,
                region.get("set_by") or None,
                region.get("set_at") or None,
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row_index, column, value if value is not None else None)
            row_index += 1
        for region in detected:
            values = [
                page["sheet"],
                page["page_index"] + 1,
                "overridden",
                region.get("label") or None,
                region.get("source") or None,
                None,
                "detected",
                region.get("reason") or "Detection was overridden by a manual scale.",
                None,
                None,
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row_index, column, value if value is not None else None)
            row_index += 1
    row_index += 1
    sheet.cell(row_index, 1, "Measured lengths").font = Font(bold=True)
    row_index += 1
    length_headers = [
        "Sheet",
        "PDF page",
        "Scale",
        "Scale source",
        "Method",
        "Length (ft)",
        "Scale read from",
        "Points",
        "Set by",
        "Set at",
    ]
    for column, header in enumerate(length_headers, start=1):
        cell = sheet.cell(row_index, column, header)
        cell.font = Font(bold=True)
    row_index += 1
    if not measurements:
        sheet.cell(row_index, 1, "No measured lengths")
        return
    for measurement in measurements:
        points = "; ".join(
            f"{point[0]},{point[1]}" for point in measurement.get("points") or []
        )
        length = measurement.get("length_ft")
        values = [
            measurement.get("sheet"),
            (measurement.get("page_index") or 0) + 1,
            measurement.get("scale_label"),
            measurement.get("scale_source"),
            measurement.get("method") or "measured",
            _number(Decimal(length)) if length not in (None, "") else None,
            measurement.get("scale_read_from") or None,
            points or None,
            measurement.get("set_by"),
            measurement.get("set_at"),
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value if value is not None else None)
        row_index += 1


def _write_labor(workbook: Workbook, job: dict, rows: list[dict], warning: str) -> None:
    sheet = workbook.create_sheet("Labor")
    sheet["A1"] = warning or "No counted line uses an assembly."
    sheet["A2"] = (
        "Rates are separate lines. Assembly hours are the takeoff count times "
        "the hours per unit. Confirmed hours are in direct cost. Example hours are not."
    )
    sheet["A4"] = "Labor rates"
    sheet["A4"].font = Font(bold=True)
    stored = {
        "journeyman": job["rate_journeyman"],
        "foreman": job["rate_foreman"],
        "pm": job["rate_pm"],
        "apprentice": job["rate_apprentice"],
    }
    rate_row = 5
    for key, label, _default in LABOR_RATES:
        sheet.cell(rate_row, 1, label)
        sheet.cell(rate_row, 2, float(Decimal(stored[key])))
        sheet.cell(rate_row, 3, "per hour")
        rate_row += 1
    sheet.cell(rate_row, 1, journeyman_warning(Decimal(job["rate_journeyman"])) or None)

    header_row = rate_row + 2
    headers = [
        "Line",
        "Item IDs",
        "Item",
        "Assembly",
        "Assembly version",
        "Count",
        "Hours per unit",
        "Hours",
        "Rate",
        "Confirmed",
        "Direct labor",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(bold=True)
    total_hours = Decimal("0")
    confirmed_hours = Decimal("0")
    direct = Decimal("0")
    row_index = header_row + 1
    for row in rows:
        hours = row["hours"] if row["hours"] is not None else Decimal("0")
        total_hours += hours
        if row["confirmed"]:
            confirmed_hours += hours
            if row["direct_labor"] is not None:
                direct += row["direct_labor"]
        values = [
            row["line_id"],
            ", ".join(row["item_ids"]),
            row["item"],
            row["assembly"],
            row["assembly_version"],
            row["count"],
            _number(row["hours_per_unit"]),
            _number(row["hours"]),
            row["rate"],
            "yes" if row["confirmed"] else "no",
            _number(row["direct_labor"]),
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value if value is not None else None)
        row_index += 1
    row_index += 1
    sheet.cell(row_index, 1, "Example and confirmed hours")
    sheet.cell(row_index, 2, _number(total_hours))
    row_index += 1
    sheet.cell(row_index, 1, "Confirmed hours")
    sheet.cell(row_index, 2, _number(confirmed_hours))
    row_index += 1
    sheet.cell(row_index, 1, "Confirmed direct labor")
    sheet.cell(row_index, 2, _number(direct))


def _decimal_or_none(value):
    if value is None or value == "":
        return None
    return Decimal(str(value))
