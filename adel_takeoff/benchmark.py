"""Score a takeoff against ground truth for one plan set.

Truth is a folder the estimator owns. Sheet files list the devices that
should be found, each with a tag and a point. A totals file per tag is enough
when only the count is known. The scorer pairs engine marks to those points
by tag and distance. The same inputs always produce the same pairs.

Real client drawings stay outside the repository. Point the command at a
folder on the estimator's machine.
"""

from __future__ import annotations

import json
import re
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from adel_takeoff.extract import read_pdf
from adel_takeoff.schedule_link import attach_schedules
from adel_takeoff.takeoff import run_takeoff


DEFAULT_TOLERANCE_PT = 18.0
SCHEDULE_FIELDS = (
    "panel",
    "circuit",
    "type_mark",
    "voltage",
    "mounting",
    "mounting_height",
)
_SLUG = re.compile(r"[^a-z0-9]+")


class BenchmarkError(ValueError):
    pass


@dataclass
class Report:
    set_name: str
    pdf_name: str
    precision: float
    recall: float
    max_abs_count_error: int
    schedule_recall: float
    matched: int
    predicted: int
    truth_points: int
    by_tag: list[dict] = field(default_factory=list)
    by_sheet: list[dict] = field(default_factory=list)
    by_category: list[dict] = field(default_factory=list)
    missed: list[dict] = field(default_factory=list)
    extra: list[dict] = field(default_factory=list)
    schedule_mismatches: list[dict] = field(default_factory=list)
    unmatched: list[dict] = field(default_factory=list)
    totals: list[dict] = field(default_factory=list)
    location_scored: bool = False
    totals_scored: bool = False

    def passes(self, thresholds: dict | None) -> bool:
        if not thresholds:
            return True
        if "precision" in thresholds and self.precision + 1e-9 < float(thresholds["precision"]):
            return False
        if "recall" in thresholds and self.recall + 1e-9 < float(thresholds["recall"]):
            return False
        if "max_abs_count_error" in thresholds and self.max_abs_count_error > int(
            thresholds["max_abs_count_error"]
        ):
            return False
        if "schedule_recall" in thresholds and self.schedule_recall + 1e-9 < float(
            thresholds["schedule_recall"]
        ):
            return False
        return True


def find_pdf(set_dir: Path) -> Path:
    folder = Path(set_dir)
    for name in ("drawing.pdf", "set.pdf"):
        candidate = folder / name
        if candidate.is_file():
            return candidate
    pdfs = sorted(path for path in folder.glob("*.pdf") if path.is_file())
    if len(pdfs) == 1:
        return pdfs[0]
    if not pdfs:
        raise BenchmarkError(f"{folder} needs a drawing PDF named drawing.pdf.")
    raise BenchmarkError(f"{folder} has more than one PDF. Name the drawing drawing.pdf.")


def load_truth(set_dir: Path) -> dict:
    folder = Path(set_dir) / "truth"
    sheets = []
    sheets_dir = folder / "sheets"
    if sheets_dir.is_dir():
        for path in sorted(sheets_dir.glob("*.json")):
            sheets.append(_sheet_document(path))
    totals = []
    totals_dir = folder / "totals"
    if totals_dir.is_dir():
        for path in sorted(totals_dir.glob("*.json")):
            totals.append(_total_document(path))
    if not sheets and not totals:
        raise BenchmarkError(
            f"{folder} needs truth/sheets/<sheet>.json or truth/totals/<tag>.json."
        )
    return {"sheets": sheets, "totals": totals}


def load_thresholds(set_dir: Path) -> dict | None:
    path = Path(set_dir) / "thresholds.json"
    if not path.is_file():
        return None
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise BenchmarkError("thresholds.json must be an object.")
    return data


def engine_items(pdf_bytes: bytes, cleanup: bool = False) -> tuple[list[dict], list[dict]]:
    """Return tagged predictions and unmatched marks from a fresh takeoff."""

    result = run_takeoff(pdf_bytes, cleanup=cleanup)
    rows = []
    for item in result.items:
        rows.append(
            {
                "id": item.item_id,
                "label": item.label,
                "category": item.category,
                "measurement": item.measurement,
                "measured_quantity": item.measured_quantity,
                "panel": item.panel or "",
                "circuit": item.circuit or "",
                "occurrences": [occurrence.as_dict() for occurrence in item.occurrences],
                "type_mark": "",
                "voltage": "",
                "mounting": "",
                "mounting_height": "",
            }
        )
    attach_schedules(rows, read_pdf(pdf_bytes))
    predicted = []
    unmatched = []
    for row in rows:
        if row["measurement"] != "matched":
            for occurrence in row["occurrences"]:
                unmatched.append(_point_from_occurrence(row, occurrence))
            continue
        quantity = row["measured_quantity"] if row["measured_quantity"] is not None else 1
        predicted.append(
            {
                "id": str(row["id"]),
                "tag": _norm_display(row["label"]),
                "tag_key": _norm_key(row["label"]),
                "category": row["category"] or "",
                "quantity": int(quantity),
                "schedule": _schedule_from_row(row),
                "occurrences": [
                    _point_from_occurrence(row, occurrence) for occurrence in row["occurrences"]
                ],
            }
        )
    return predicted, unmatched


def score_set(set_dir: Path, cleanup: bool = False) -> Report:
    folder = Path(set_dir)
    pdf_path = find_pdf(folder)
    truth = load_truth(folder)
    predicted, unmatched = engine_items(pdf_path.read_bytes(), cleanup=cleanup)
    report = score_predictions(
        predicted,
        unmatched,
        truth,
        set_name=folder.name,
        pdf_name=pdf_path.name,
    )
    return report


def score_predictions(
    predicted: list[dict],
    unmatched: list[dict],
    truth: dict,
    *,
    set_name: str = "",
    pdf_name: str = "",
) -> Report:
    truth_points = _truth_points(truth)
    pairs = _match_points(truth_points, predicted) if truth_points else []
    matched_truth = {pair["truth_index"] for pair in pairs}
    matched_occ = {pair["occurrence_id"] for pair in pairs}
    missed = [
        _miss_row(point)
        for index, point in enumerate(truth_points)
        if index not in matched_truth
    ]
    extra = []
    if truth_points:
        for item in predicted:
            for occurrence in item["occurrences"]:
                if occurrence["occurrence_id"] in matched_occ:
                    continue
                extra.append(
                    {
                        "item_id": item["id"],
                        "sheet": occurrence["sheet"],
                        "tag": item["tag"],
                        "category": item["category"],
                        "x": occurrence["x"],
                        "y": occurrence["y"],
                    }
                )
    mismatches = []
    schedule_checks = 0
    schedule_hits = 0
    for pair in pairs:
        for field_name in SCHEDULE_FIELDS:
            expected = pair["expected_schedule"].get(field_name)
            if expected is None:
                continue
            schedule_checks += 1
            actual = pair["actual_schedule"].get(field_name, "")
            if _same_schedule(field_name, expected, actual):
                schedule_hits += 1
                continue
            mismatches.append(
                {
                    "sheet": pair["sheet"],
                    "tag": pair["tag"],
                    "x": pair["x"],
                    "y": pair["y"],
                    "field": field_name,
                    "expected": expected,
                    "actual": actual,
                }
            )
    predicted_count = sum(len(item["occurrences"]) for item in predicted)
    location = bool(truth_points)
    if location:
        precision = _ratio(len(pairs), predicted_count)
        recall = _ratio(len(pairs), len(truth_points))
    else:
        precision, recall = _totals_ratios(predicted, truth.get("totals") or [])
    totals_rows = _totals_rows(predicted, truth.get("totals") or [])
    count_errors = [abs(int(row["count_error"])) for row in totals_rows if row["expected"] is not None]
    if truth_points:
        by_tag = _group_rows(truth_points, predicted, pairs, key="tag")
        _attach_count_errors(by_tag, totals_rows)
        by_sheet = _group_rows(truth_points, predicted, pairs, key="sheet")
        by_category = _group_rows(truth_points, predicted, pairs, key="category")
    else:
        by_tag = _tag_rows_from_totals(totals_rows)
        by_sheet = []
        by_category = []
    return Report(
        set_name=set_name,
        pdf_name=pdf_name,
        precision=precision,
        recall=recall,
        max_abs_count_error=max(count_errors, default=0),
        schedule_recall=_ratio(schedule_hits, schedule_checks) if schedule_checks else 1.0,
        matched=len(pairs),
        predicted=predicted_count,
        truth_points=len(truth_points),
        by_tag=by_tag,
        by_sheet=by_sheet,
        by_category=by_category,
        missed=missed,
        extra=extra,
        schedule_mismatches=mismatches,
        unmatched=unmatched,
        totals=totals_rows,
        location_scored=location,
        totals_scored=bool(truth.get("totals")),
    )


def write_scorecard(report: Report, out_dir: Path) -> tuple[Path, Path]:
    folder = Path(out_dir)
    folder.mkdir(parents=True, exist_ok=True)
    markdown_path = folder / "scorecard.md"
    excel_path = folder / "scorecard.xlsx"
    _write_text(markdown_path, render_markdown(report))
    excel_path.write_bytes(render_excel(report))
    return markdown_path, excel_path


def render_markdown(report: Report) -> str:
    lines = [
        f"# Benchmark scorecard: {report.set_name or 'set'}",
        "",
        f"Drawing: {report.pdf_name}",
        "",
        "## Overall",
        "",
        f"- Precision: {report.precision:.3f}",
        f"- Recall: {report.recall:.3f}",
        (
            f"- Matched locations: {report.matched} of {report.truth_points} truth points, "
            f"{report.predicted} engine locations"
            if report.location_scored
            else "- Locations: not scored. Precision and recall compare tag totals."
        ),
        f"- Largest absolute count error: {report.max_abs_count_error}",
        f"- Schedule recall: {report.schedule_recall:.3f}",
        "",
    ]
    if not report.location_scored:
        lines.append(
            "This set lists tag totals only. Sheet and category location scores are omitted."
        )
        lines.append("")
    lines.extend(
        [
        "## By tag",
        "",
        _markdown_table(
            ["Tag", "Truth", "Engine locations", "Matched", "Precision", "Recall", "Count error"],
            [
                [
                    row["key"],
                    row["truth"],
                    row["predicted"],
                    row["matched"],
                    _fmt_ratio(row["precision"]),
                    _fmt_ratio(row["recall"]),
                    "" if row.get("count_error") is None else row["count_error"],
                ]
                for row in report.by_tag
            ],
        ),
        "",
        "## By sheet",
        "",
        _markdown_table(
            ["Sheet", "Truth", "Engine locations", "Matched", "Precision", "Recall"],
            [
                [
                    row["key"],
                    row["truth"],
                    row["predicted"],
                    row["matched"],
                    _fmt_ratio(row["precision"]),
                    _fmt_ratio(row["recall"]),
                ]
                for row in report.by_sheet
            ],
        ),
        "",
        "## By category",
        "",
        _markdown_table(
            ["Category", "Truth", "Engine locations", "Matched", "Precision", "Recall"],
            [
                [
                    row["key"],
                    row["truth"],
                    row["predicted"],
                    row["matched"],
                    _fmt_ratio(row["precision"]),
                    _fmt_ratio(row["recall"]),
                ]
                for row in report.by_category
            ],
        ),
        "",
        "## Totals",
        "",
        _markdown_table(
            ["Tag", "Category", "Expected", "Engine", "Count error"],
            [
                [row["tag"], row["category"], row["expected"], row["engine"], row["count_error"]]
                for row in report.totals
            ],
        ),
        "",
        "## Missed",
        "",
        _markdown_table(
            ["Sheet", "Tag", "Category", "X", "Y"],
            [
                [row["sheet"], row["tag"], row["category"], row["x"], row["y"]]
                for row in report.missed
            ],
        ),
        "",
        "## Extra",
        "",
        _markdown_table(
            ["Item", "Sheet", "Tag", "Category", "X", "Y"],
            [
                [row["item_id"], row["sheet"], row["tag"], row["category"], row["x"], row["y"]]
                for row in report.extra
            ],
        ),
        "",
        "## Schedule mismatches",
        "",
        _markdown_table(
            ["Sheet", "Tag", "Field", "Expected", "Engine", "X", "Y"],
            [
                [
                    row["sheet"],
                    row["tag"],
                    row["field"],
                    row["expected"],
                    row["actual"],
                    row["x"],
                    row["y"],
                ]
                for row in report.schedule_mismatches
            ],
        ),
        "",
        "## Unmatched engine marks",
        "",
        "These marks have no tag, so they are not counted as extra devices.",
        "",
        _markdown_table(
            ["Sheet", "X", "Y"],
            [[row["sheet"], row["x"], row["y"]] for row in report.unmatched],
        ),
        "",
        ]
    )
    return "\n".join(lines)


def render_excel(report: Report) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "Value"])
    summary.append(["Set", report.set_name])
    summary.append(["Drawing", report.pdf_name])
    summary.append(["Precision", report.precision])
    summary.append(["Recall", report.recall])
    summary.append(["Matched locations", report.matched])
    summary.append(["Truth points", report.truth_points])
    summary.append(["Engine locations", report.predicted])
    summary.append(["Largest absolute count error", report.max_abs_count_error])
    summary.append(["Schedule recall", report.schedule_recall])
    _sheet_rows(
        workbook.create_sheet("By tag"),
        ["Tag", "Truth", "Engine locations", "Matched", "Precision", "Recall", "Count error"],
        [
            [
                row["key"],
                row["truth"],
                row["predicted"],
                row["matched"],
                row["precision"],
                row["recall"],
                row.get("count_error"),
            ]
            for row in report.by_tag
        ],
    )
    _sheet_rows(
        workbook.create_sheet("By sheet"),
        ["Sheet", "Truth", "Engine locations", "Matched", "Precision", "Recall"],
        [
            [row["key"], row["truth"], row["predicted"], row["matched"], row["precision"], row["recall"]]
            for row in report.by_sheet
        ],
    )
    _sheet_rows(
        workbook.create_sheet("By category"),
        ["Category", "Truth", "Engine locations", "Matched", "Precision", "Recall"],
        [
            [row["key"], row["truth"], row["predicted"], row["matched"], row["precision"], row["recall"]]
            for row in report.by_category
        ],
    )
    _sheet_rows(
        workbook.create_sheet("Totals"),
        ["Tag", "Category", "Expected", "Engine", "Count error"],
        [
            [row["tag"], row["category"], row["expected"], row["engine"], row["count_error"]]
            for row in report.totals
        ],
    )
    _sheet_rows(
        workbook.create_sheet("Missed"),
        ["Sheet", "Tag", "Category", "X", "Y"],
        [[row["sheet"], row["tag"], row["category"], row["x"], row["y"]] for row in report.missed],
    )
    _sheet_rows(
        workbook.create_sheet("Extra"),
        ["Item", "Sheet", "Tag", "Category", "X", "Y"],
        [
            [row["item_id"], row["sheet"], row["tag"], row["category"], row["x"], row["y"]]
            for row in report.extra
        ],
    )
    _sheet_rows(
        workbook.create_sheet("Schedule"),
        ["Sheet", "Tag", "Field", "Expected", "Engine", "X", "Y"],
        [
            [row["sheet"], row["tag"], row["field"], row["expected"], row["actual"], row["x"], row["y"]]
            for row in report.schedule_mismatches
        ],
    )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def documents_from_rows(rows: list[dict], tolerance_pt: float = DEFAULT_TOLERANCE_PT) -> dict[str, dict]:
    """Turn accepted, corrected, and added review rows into truth files.

    Rejected marks, excluded marks, and unmatched marks without a correction
    are left out. One device drawn on two sheets becomes one total and a
    point on each sheet.
    """

    grouped: dict[str, list[dict]] = {}
    quantities: dict[str, dict] = {}
    for row in rows:
        if not include_in_truth(row):
            continue
        tag = _norm_display(row.get("label") or "")
        if not tag:
            continue
        category = row.get("category") or ""
        key = _norm_key(tag)
        bucket = quantities.setdefault(key, {"tag": tag, "category": category, "quantity": 0})
        measured = row.get("measured")
        if measured is None:
            measured = row.get("measured_quantity")
        bucket["quantity"] += int(measured if measured is not None else 1)
        if row.get("category"):
            bucket["category"] = row["category"]
        schedule = _schedule_from_row(row)
        for occurrence in row.get("occurrences") or []:
            sheet = occurrence.get("sheet") or f"Page {int(occurrence.get('page_index') or 0) + 1}"
            x0 = float(occurrence.get("x0") or 0)
            y0 = float(occurrence.get("y0") or 0)
            x1 = float(occurrence.get("x1") or x0)
            y1 = float(occurrence.get("y1") or y0)
            grouped.setdefault(sheet, []).append(
                {
                    "category": category,
                    "schedule": schedule,
                    "tag": tag,
                    "x": round((x0 + x1) / 2, 2),
                    "y": round((y0 + y1) / 2, 2),
                }
            )
    documents = {}
    for sheet in sorted(grouped):
        items = sorted(grouped[sheet], key=lambda item: (item["y"], item["x"], item["tag"]))
        documents[f"truth/sheets/{_file_slug(sheet)}.json"] = {
            "items": items,
            "sheet": sheet,
            "tolerance_pt": tolerance_pt,
        }
    for key in sorted(quantities):
        record = quantities[key]
        documents[f"truth/totals/{tag_slug(record['tag'])}.json"] = {
            "category": record["category"],
            "quantity": record["quantity"],
            "tag": record["tag"],
        }
    return documents


def include_in_truth(row: dict) -> bool:
    if row.get("exclusion"):
        return False
    status = row.get("review_status") or "accepted"
    if status in {"rejected", "pending"}:
        return False
    if status == "corrected":
        return bool(str(row.get("label") or "").strip())
    if status != "accepted":
        return False
    if row.get("measurement") == "matched" or row.get("method") == "manual_add":
        return True
    return False


def write_truth_dir(out_dir: Path, documents: dict[str, dict]) -> None:
    root = Path(out_dir)
    for relative, payload in documents.items():
        path = root / relative
        _write_text(path, json.dumps(payload, indent=2, sort_keys=True) + "\n")


def truth_zip_bytes(rows: list[dict], pdf_bytes: bytes | None = None) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        if pdf_bytes:
            archive.writestr("drawing.pdf", pdf_bytes)
        for relative, payload in sorted(documents_from_rows(rows).items()):
            archive.writestr(relative, json.dumps(payload, indent=2, sort_keys=True) + "\n")
    return buffer.getvalue()


def tag_slug(tag: str) -> str:
    slug = _SLUG.sub("-", tag.casefold()).strip("-")
    return slug or "tag"


def _truth_points(truth: dict) -> list[dict]:
    points = []
    for sheet in truth.get("sheets") or []:
        sheet_name = sheet.get("sheet") or ""
        sheet_tolerance = float(sheet.get("tolerance_pt") or DEFAULT_TOLERANCE_PT)
        for item in sheet.get("items") or []:
            tolerance = item.get("tolerance_pt")
            points.append(
                {
                    "sheet": item.get("sheet") or sheet_name,
                    "tag": _norm_display(item.get("tag") or ""),
                    "tag_key": _norm_key(item.get("tag") or ""),
                    "category": item.get("category") or "",
                    "x": float(item["x"]),
                    "y": float(item["y"]),
                    "tolerance_pt": float(tolerance if tolerance is not None else sheet_tolerance),
                    "schedule": item.get("schedule") or {},
                }
            )
    return points


def _match_points(truth_points: list[dict], predicted: list[dict]) -> list[dict]:
    occurrences = []
    for item in predicted:
        for occurrence in item["occurrences"]:
            occurrences.append((item, occurrence))
    candidates = []
    for truth_index, point in enumerate(truth_points):
        for engine_index, (item, occurrence) in enumerate(occurrences):
            if item["tag_key"] != point["tag_key"]:
                continue
            if point["sheet"] and occurrence["sheet"] and point["sheet"] != occurrence["sheet"]:
                continue
            distance = _distance(point["x"], point["y"], occurrence["x"], occurrence["y"])
            if distance > point["tolerance_pt"] + 1e-6:
                continue
            candidates.append((round(distance, 4), truth_index, engine_index))
    candidates.sort()
    used_truth = set()
    used_engine = set()
    pairs = []
    for _pair_distance, truth_index, engine_index in candidates:
        if truth_index in used_truth or engine_index in used_engine:
            continue
        used_truth.add(truth_index)
        used_engine.add(engine_index)
        point = truth_points[truth_index]
        item, occurrence = occurrences[engine_index]
        pairs.append(
            {
                "truth_index": truth_index,
                "occurrence_id": occurrence["occurrence_id"],
                "sheet": point["sheet"],
                "tag": point["tag"],
                "category": point["category"] or item["category"],
                "x": point["x"],
                "y": point["y"],
                "expected_schedule": point["schedule"] if isinstance(point["schedule"], dict) else {},
                "actual_schedule": item["schedule"],
            }
        )
    return pairs


def _group_rows(truth_points, predicted, pairs, *, key: str) -> list[dict]:
    truth_counts: dict[str, int] = {}
    predicted_counts: dict[str, int] = {}
    matched_counts: dict[str, int] = {}
    for point in truth_points:
        name = point[key] or "(blank)"
        truth_counts[name] = truth_counts.get(name, 0) + 1
    for item in predicted:
        for occurrence in item["occurrences"]:
            name = item["tag"] if key == "tag" else item["category"] if key == "category" else occurrence["sheet"]
            name = name or "(blank)"
            predicted_counts[name] = predicted_counts.get(name, 0) + 1
    for pair in pairs:
        if key == "tag":
            name = pair["tag"]
        elif key == "category":
            name = pair["category"]
        else:
            name = pair["sheet"]
        name = name or "(blank)"
        matched_counts[name] = matched_counts.get(name, 0) + 1
    names = sorted(set(truth_counts) | set(predicted_counts) | set(matched_counts))
    rows = []
    for name in names:
        truth_count = truth_counts.get(name, 0)
        predicted_count = predicted_counts.get(name, 0)
        matched_count = matched_counts.get(name, 0)
        rows.append(
            {
                "key": name,
                "truth": truth_count,
                "predicted": predicted_count,
                "matched": matched_count,
                "precision": _ratio(matched_count, predicted_count),
                "recall": _ratio(matched_count, truth_count),
                "count_error": None,
            }
        )
    return rows


def _totals_rows(predicted: list[dict], totals: list[dict]) -> list[dict]:
    engine: dict[str, dict] = {}
    for item in predicted:
        bucket = engine.setdefault(
            item["tag_key"],
            {"tag": item["tag"], "category": item["category"], "quantity": 0},
        )
        bucket["quantity"] += int(item["quantity"])
    rows = []
    seen = set()
    for total in totals:
        key = _norm_key(total.get("tag") or "")
        seen.add(key)
        expected = int(total["quantity"])
        actual = engine.get(key, {}).get("quantity", 0)
        rows.append(
            {
                "tag": total.get("tag") or "",
                "category": total.get("category") or engine.get(key, {}).get("category", ""),
                "expected": expected,
                "engine": actual,
                "count_error": actual - expected,
            }
        )
    for key, bucket in sorted(engine.items()):
        if key in seen:
            continue
        if totals:
            rows.append(
                {
                    "tag": bucket["tag"],
                    "category": bucket["category"],
                    "expected": None,
                    "engine": bucket["quantity"],
                    "count_error": None,
                }
            )
    rows.sort(key=lambda row: row["tag"])
    return rows


def _tag_rows_from_totals(totals_rows: list[dict]) -> list[dict]:
    rows = []
    for row in totals_rows:
        expected = row["expected"]
        engine = int(row["engine"])
        if expected is None:
            matched = 0
            precision = None
            recall = None
            truth_count = 0
        else:
            matched = min(engine, int(expected))
            precision = _ratio(matched, engine)
            recall = _ratio(matched, int(expected))
            truth_count = int(expected)
        rows.append(
            {
                "key": row["tag"],
                "truth": truth_count,
                "predicted": engine,
                "matched": matched,
                "precision": precision,
                "recall": recall,
                "count_error": row["count_error"],
            }
        )
    return rows


def _attach_count_errors(by_tag: list[dict], totals_rows: list[dict]) -> None:
    errors = {_norm_key(row["tag"]): row["count_error"] for row in totals_rows}
    for row in by_tag:
        row["count_error"] = errors.get(_norm_key(row["key"]))


def _totals_ratios(predicted: list[dict], totals: list[dict]) -> tuple[float, float]:
    rows = [row for row in _totals_rows(predicted, totals) if row["expected"] is not None]
    matched = sum(min(int(row["engine"]), int(row["expected"])) for row in rows)
    expected = sum(int(row["expected"]) for row in rows)
    engine = sum(int(row["engine"]) for row in rows)
    return _ratio(matched, engine), _ratio(matched, expected)


def _sheet_document(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or not isinstance(data.get("items"), list):
        raise BenchmarkError(f"{path} needs an items list.")
    if not data.get("sheet"):
        data["sheet"] = path.stem
    return data


def _total_document(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or "quantity" not in data or not data.get("tag"):
        raise BenchmarkError(f"{path} needs a tag and a quantity.")
    return data


def _point_from_occurrence(row: dict, occurrence: dict) -> dict:
    x0 = float(occurrence.get("x0") or 0)
    y0 = float(occurrence.get("y0") or 0)
    x1 = float(occurrence.get("x1") or x0)
    y1 = float(occurrence.get("y1") or y0)
    sheet = occurrence.get("sheet") or ""
    item_id = str(row.get("id") or "")
    return {
        "occurrence_id": f"{item_id}:{sheet}:{round((x0 + x1) / 2, 2)}:{round((y0 + y1) / 2, 2)}",
        "sheet": sheet,
        "x": round((x0 + x1) / 2, 2),
        "y": round((y0 + y1) / 2, 2),
    }


def _schedule_from_row(row: dict) -> dict:
    return {name: "" if row.get(name) is None else str(row.get(name) or "") for name in SCHEDULE_FIELDS}


def _miss_row(point: dict) -> dict:
    return {
        "sheet": point["sheet"],
        "tag": point["tag"],
        "category": point["category"],
        "x": point["x"],
        "y": point["y"],
    }


def _same_schedule(field_name: str, expected, actual) -> bool:
    left = " ".join(str(expected).split())
    right = " ".join(str(actual or "").split())
    if field_name in {"panel", "circuit", "type_mark"}:
        return left.casefold() == right.casefold()
    return left == right


def _distance(x0: float, y0: float, x1: float, y1: float) -> float:
    return ((x0 - x1) ** 2 + (y0 - y1) ** 2) ** 0.5


def _ratio(numerator: int, denominator: int) -> float:
    if denominator == 0:
        return 1.0
    return numerator / denominator


def _fmt_ratio(value) -> str:
    if value is None:
        return ""
    return f"{float(value):.3f}"


def _norm_key(value: str) -> str:
    return " ".join(str(value).casefold().split())


def _norm_display(value: str) -> str:
    return " ".join(str(value).split())


def _file_slug(value: str) -> str:
    slug = _SLUG.sub("-", value.casefold()).strip("-")
    return slug or "sheet"


def _markdown_table(headers: list[str], rows: list[list]) -> str:
    head = "| " + " | ".join(headers) + " |"
    rule = "| " + " | ".join("---" for _ in headers) + " |"
    body = ["| " + " | ".join(_cell(value) for value in row) + " |" for row in rows]
    if not body:
        body = ["| " + " | ".join("" for _ in headers) + " |"]
    return "\n".join([head, rule, *body])


def _cell(value) -> str:
    return str(value if value is not None else "").replace("|", "/")


def _sheet_rows(sheet, headers: list[str], rows: list[list]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))


def _write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(text, encoding="utf-8")
    temporary.replace(path)
