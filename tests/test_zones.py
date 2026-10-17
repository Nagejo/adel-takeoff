"""Zones for floors, phases, areas, alternates, and by-others work."""

import re
from io import BytesIO

import fitz
import openpyxl
from fastapi.testclient import TestClient

from adel_takeoff.main import create_app
from adel_takeoff.zones import assign_zones, point_in_polygon, zone_rollups
from tests.drawings import build_sample_pdf, draw_duplex


def _text(page, x, y, text, size=9):
    page.insert_text(fitz.Point(x, y), text, fontsize=size, fontname="helv")


def _occurrence(page_index, x, y, sheet):
    return {
        "page_index": page_index,
        "sheet": sheet,
        "x0": x - 8,
        "y0": y - 8,
        "x1": x + 8,
        "y1": y + 8,
    }


def _row(item_id, occurrences, effective=1):
    return {
        "id": item_id,
        "label": "DUPLEX RECEPTACLE",
        "category": "power",
        "system": "power",
        "effective": effective,
        "occurrences": occurrences,
    }


def _zone(zone_id, name, zone_type, page_index, sheet, corners, alternate=""):
    x0, y0, x1, y1 = corners
    return {
        "id": zone_id,
        "name": name,
        "type": zone_type,
        "alternate": alternate,
        "sheet": sheet,
        "page_index": page_index,
        "shape": "rectangle",
        "points": [[x0, y0], [x1, y0], [x1, y1], [x0, y1]],
    }


def build_zone_pdf() -> bytes:
    document = fitz.open()
    legend = document.new_page(width=612, height=792)
    _text(legend, 72, 48, "SYMBOL LEGEND", size=12)
    draw_duplex(legend, 90, 100)
    _text(legend, 118, 103, "DUPLEX RECEPTACLE")

    sheet = document.new_page(width=792, height=792)
    _text(sheet, 72, 36, "SHEET E2", size=12)
    _text(sheet, 220, 36, "TITLE LEVEL 2 POWER", size=12)
    draw_duplex(sheet, 200, 220)
    draw_duplex(sheet, 400, 220)
    draw_duplex(sheet, 200, 420)
    draw_duplex(sheet, 520, 420)

    other = document.new_page(width=792, height=792)
    _text(other, 72, 36, "SHEET E3", size=12)
    _text(other, 220, 36, "TITLE LEVEL 2 LIGHTING", size=12)
    draw_duplex(other, 200, 220)
    draw_duplex(other, 400, 420)
    payload = document.tobytes()
    document.close()
    return payload


def _upload(tmp_path, payload):
    client = TestClient(create_app(tmp_path))
    response = client.post(
        "/jobs",
        files={"pdf": ("zones.pdf", payload, "application/pdf")},
        follow_redirects=False,
    )
    assert response.status_code == 303
    job_id = response.headers["location"].rstrip("/").split("/")[-1]
    return client, job_id


def _save_zone(client, job_id, **fields):
    response = client.post(f"/jobs/{job_id}/zones", data=fields, follow_redirects=False)
    assert response.status_code == 303, response.text
    return response


def test_point_in_rectangle_and_polygon_includes_the_boundary():
    rectangle = [[10, 20], [110, 20], [110, 80], [10, 80]]
    assert point_in_polygon(40, 40, rectangle)
    assert point_in_polygon(10, 40, rectangle)
    assert point_in_polygon(110, 80, rectangle)
    assert not point_in_polygon(9.9, 40, rectangle)
    polygon = [[0, 0], [100, 0], [50, 80]]
    assert point_in_polygon(50, 20, polygon)
    assert point_in_polygon(0, 0, polygon)
    assert not point_in_polygon(10, 70, polygon)


def test_same_zone_name_on_two_sheets_rolls_up_once_per_item():
    row = _row(
        "1",
        [
            _occurrence(1, 200, 220, "E2"),
            _occurrence(2, 200, 220, "E3"),
        ],
    )
    zones = [
        _zone("z1", "Level 2", "floor", 1, "E2", (160, 180, 240, 260)),
        _zone("z2", "Level 2", "floor", 2, "E3", (160, 180, 240, 260)),
    ]
    assign_zones([row], zones)
    assert row["zone_label"] == "Floor: Level 2"
    assert row["zone_conflicts"] == []
    assert row["counts_in_base"] is True
    rollups = zone_rollups([row])
    assert len(rollups) == 1
    assert rollups[0]["name"] == "Level 2"
    assert rollups[0]["quantity"] == 1
    assert rollups[0]["sheets"] == ["E2", "E3"]


def test_different_types_both_assign_and_same_type_overlap_does_not():
    both = _row("1", [_occurrence(1, 200, 220, "E2")])
    assign_zones(
        [both],
        [
            _zone("z1", "Level 2", "floor", 1, "E2", (160, 180, 240, 260)),
            _zone("z2", "Phase 1", "phase", 1, "E2", (150, 170, 250, 270)),
        ],
    )
    assert both["zone_label"] == "Floor: Level 2; Phase: Phase 1"
    assert both["zone_hold"] == ""

    overlap = _row("2", [_occurrence(1, 520, 420, "E2")])
    assign_zones(
        [overlap],
        [
            _zone("z3", "North", "floor", 1, "E2", (490, 390, 560, 460)),
            _zone("z4", "South", "floor", 1, "E2", (500, 400, 580, 470)),
        ],
    )
    assert overlap["zones"] == []
    assert overlap["zone_hold"] == ""
    assert overlap["counts_in_base"] is True
    assert "North" in overlap["zone_conflict_text"]
    assert "South" in overlap["zone_conflict_text"]
    assert "Not assigned to either" in overlap["zone_conflict_text"]
    assert "violation" not in overlap["zone_conflict_text"].lower()
    assert zone_rollups([overlap]) == []

    split = _row(
        "3",
        [
            _occurrence(1, 200, 220, "E2"),
            _occurrence(2, 400, 220, "E3"),
        ],
    )
    assign_zones(
        [split],
        [
            _zone("z5", "Level 2", "floor", 1, "E2", (160, 180, 240, 260)),
            _zone("z6", "Level 3", "floor", 2, "E3", (360, 180, 440, 260)),
        ],
    )
    assert split["zone_conflicts"] == []
    assert {zone["name"] for zone in split["zones"]} == {"Level 2", "Level 3"}
    quantities = {item["name"]: item["quantity"] for item in zone_rollups([split])}
    assert quantities == {"Level 2": 1, "Level 3": 1}


def test_by_others_and_alternates_leave_the_base_and_stay_listed():
    nic = _row("1", [_occurrence(1, 400, 220, "E2")])
    added = _row("2", [_occurrence(1, 200, 420, "E2")])
    deducted = _row("3", [_occurrence(2, 400, 420, "E3")])
    assign_zones(
        [nic, added, deducted],
        [
            _zone("z1", "NIC", "by_others", 1, "E2", (360, 180, 440, 260)),
            _zone("z2", "Lobby", "alternate", 1, "E2", (160, 380, 240, 460), "add"),
            _zone("z3", "Owner deduct", "alternate", 2, "E3", (360, 380, 440, 460), "deduct"),
        ],
    )
    assert nic["effective"] == 1
    assert nic["zone_hold"] == "by_others"
    assert nic["counts_in_base"] is False
    assert added["zone_hold"] == "alternate" and added["alternate_side"] == "add"
    assert deducted["alternate_side"] == "deduct"
    sides = {(item["name"], item["side"], item["in_base"]) for item in zone_rollups([nic, added, deducted])}
    assert ("NIC", "", False) in sides
    assert ("Lobby", "add", False) in sides
    assert ("Owner deduct", "deduct", False) in sides


def test_assignment_is_deterministic():
    row = _row("9", [_occurrence(1, 200, 220, "E2")])
    zones = [
        _zone("z2", "Phase 1", "phase", 1, "E2", (150, 170, 250, 270)),
        _zone("z1", "Level 2", "floor", 1, "E2", (160, 180, 240, 260)),
    ]
    first = dict(row)
    second = dict(row)
    assign_zones([first], list(reversed(zones)))
    assign_zones([second], zones)
    assert first["zone_label"] == second["zone_label"]
    assert first["zone_key"] == second["zone_key"]


def test_zone_page_review_excel_and_pdf(tmp_path):
    client, job_id = _upload(tmp_path, build_zone_pdf())
    bare = client.get(f"/jobs/{job_id}/review").text
    assert 'id="power-device-count">6<' in bare
    assert 'id="zones-empty"' in bare

    saved = [
        ("Level 2", "floor", "", "1", "170,190 230,250"),
        ("Level 2", "floor", "", "2", "170,190 230,250"),
        ("NIC closet", "by_others", "", "1", "370,190 430,250"),
        ("Lobby alternate", "alternate", "add", "1", "170,390 230,450"),
        ("Owner deduct", "alternate", "deduct", "2", "370,390 430,450"),
        ("North", "floor", "", "1", "490,390 560,460"),
        ("South", "floor", "", "1", "500,400 580,470"),
    ]
    for name, zone_type, side, page, points in saved:
        response = _save_zone(
            client,
            job_id,
            name=name,
            type=zone_type,
            alternate=side,
            shape="rectangle",
            page_index=page,
            points=points,
        )
        assert "error=" not in response.headers["location"]

    zones_page = client.get(f"/jobs/{job_id}/zones").text
    assert "Level 2" in zones_page
    assert "NIC closet" in zones_page
    png = client.get(f"/jobs/{job_id}/zones/page.png?page=1")
    assert png.status_code == 200
    assert png.headers["content-type"] == "image/png"

    review = client.get(f"/jobs/{job_id}/review").text
    assert 'id="power-device-count">3<' in review
    assert "Level 2" in review
    assert "Needs review" in review
    assert "In more than one floor zone" in review
    assert "not in base" in review
    assert "Lobby alternate" in review
    assert "violation" not in review.lower()
    assert "zone-conflict" in review

    rejected = _save_zone(
        client,
        job_id,
        name="",
        type="floor",
        alternate="",
        shape="rectangle",
        page_index="1",
        points="1,1 2,2",
    )
    assert "error=" in rejected.headers["location"]

    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    assert "Zones" in workbook.sheetnames
    assert "Routing" not in workbook.sheetnames
    quantities = _sheet_text(workbook["Quantities"])
    assert "Zone" in quantities
    assert "Power devices: 3" in quantities
    assert "Alternate add / deduct" in quantities
    assert "By others / NIC" in quantities
    assert "Lobby alternate" in quantities
    zones_text = _sheet_text(workbook["Zones"])
    assert "Level 2" in zones_text
    assert "no" in zones_text
    header = _header(workbook["Quantities"])
    assert "Zone" in header
    assert "How counted" in header

    pdf_text = _pdf_text(client.get(f"/jobs/{job_id}/export.pdf").content)
    assert "Level 2" in pdf_text
    assert "NIC closet" in pdf_text

    nic_id = _item_id(review, "NIC closet")
    client.post(
        f"/jobs/{job_id}/items/{nic_id}/correction",
        data={"material_unit_cost": "100", "labor_hours": "0", "labor_rate_key": "journeyman"},
        follow_redirects=False,
    )
    priced = client.get(f"/jobs/{job_id}/review").text
    assert 'id="power-device-count">3<' in priced
    cost = re.search(r'id="direct-cost">([^<]+)<', priced).group(1)
    assert cost in {"0", "0.00", "0.0"}


def test_sample_job_without_zones_keeps_its_counts(tmp_path):
    client, job_id = _upload(tmp_path, build_sample_pdf())
    review = client.get(f"/jobs/{job_id}/review").text
    assert 'id="power-device-count">4<' in review
    assert 'id="zones-empty"' in review
    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    assert "No zones" in _sheet_text(workbook["Zones"])
    assert "Power devices: 4" in _sheet_text(workbook["Quantities"])


def test_polygon_zone_and_delete(tmp_path):
    client, job_id = _upload(tmp_path, build_zone_pdf())
    response = _save_zone(
        client,
        job_id,
        name="Wing",
        type="area",
        alternate="",
        shape="polygon",
        page_index="1",
        points="160,180 280,180 220,300",
    )
    assert "error=" not in response.headers["location"]
    page = client.get(f"/jobs/{job_id}/zones").text
    assert "Wing" in page
    zone_id = re.search(r'id="zone-(z\d+)"', page).group(1)
    deleted = client.post(f"/jobs/{job_id}/zones/{zone_id}/delete", follow_redirects=False)
    assert deleted.status_code == 303
    assert "Wing" not in client.get(f"/jobs/{job_id}/zones").text


def _sheet_text(sheet) -> str:
    parts = []
    for row in sheet.iter_rows(values_only=True):
        parts.append(" ".join("" if value is None else str(value) for value in row))
    return "\n".join(parts)


def _header(sheet) -> list:
    for row in sheet.iter_rows(values_only=True):
        if row and "How counted" in row:
            return list(row)
    raise AssertionError("header row not found")


def _pdf_text(payload: bytes) -> str:
    document = fitz.open(stream=payload, filetype="pdf")
    try:
        return "\n".join(page.get_text() for page in document)
    finally:
        document.close()


def _item_id(html: str, label: str) -> str:
    for match in re.finditer(r'<tr id="item-([^"]+)"[^>]*>(.*?)</tr>', html, re.DOTALL):
        if label in match.group(2):
            return match.group(1)
    raise AssertionError(label)
