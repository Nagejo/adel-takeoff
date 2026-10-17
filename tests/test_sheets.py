"""Sheet classes, the steps each class runs, and background cleanup."""

import json
from io import BytesIO
from pathlib import Path

import fitz
import openpyxl
from fastapi.testclient import TestClient

from adel_takeoff.cleanup import BACKGROUND_CLEANUP, measure_cleanup, suppress_background
from adel_takeoff.extract import Primitive, Word, read_pdf
from adel_takeoff.main import create_app
from adel_takeoff.sheets import classify_page, classify_pages
from adel_takeoff.takeoff import run_takeoff
from tests.drawings import build_sample_pdf, build_scheduled_pdf, draw_duplex


ROOT = Path(__file__).resolve().parents[1]


def _text(page, x, y, text, size=9):
    page.insert_text(fitz.Point(x, y), text, fontsize=size, fontname="helv")


def _word(text, x, y, source="pdf", confidence=None):
    return Word(x, y, x + max(8, len(text) * 4.5), y + 8, text, source=source, confidence=confidence)


def _page(words, sheet="E1", title="", width=792, height=792, vertical_lines=0):
    from adel_takeoff.extract import PageData

    return PageData(
        index=0,
        width=width,
        height=height,
        words=words,
        sheet=sheet,
        sheet_title=title,
        text_source="ocr" if any(word.source == "ocr" for word in words) else "pdf",
        vertical_lines=vertical_lines,
    )


def _line(x0, y0, x1, y1, color=(0, 0, 0), width=1.0):
    return Primitive(
        "l",
        [(x0, y0), (x1, y1)],
        (min(x0, x1), min(y0, y1), max(x0, x1), max(y0, y1)),
        color=color,
        width=width,
    )


def _curve(x, y, color=(0, 0, 0)):
    return Primitive(
        "c",
        [(x, y), (x + 4, y + 2), (x + 8, y), (x + 4, y - 2)],
        (x, y - 2, x + 8, y + 2),
        color=color,
    )


def test_title_and_prefix_choose_a_class_with_reasons():
    power = classify_page(
        _page(
            [_word("SHEET", 72, 36), _word("E1", 120, 36), _word("TITLE", 220, 36), _word("FIRST", 280, 36), _word("FLOOR", 330, 36), _word("POWER", 390, 36)],
            sheet="E1",
            title="FIRST FLOOR POWER",
            vertical_lines=5,
        )
    )
    assert power["sheet_class"] == "power_plan"
    assert power["confidence"] == "high"
    assert any("power" in reason.lower() for reason in power["reasons"])
    assert any("Column grid" in reason for reason in power["reasons"])
    assert "Symbol detection" in power["step_labels"]
    assert "Fixture schedule" not in power["step_labels"]

    lighting = classify_page(
        _page([_word("EL2", 72, 30)], sheet="EL2", title="")
    )
    assert lighting["sheet_class"] == "lighting_plan"
    assert lighting["confidence"] == "medium"
    assert any("EL" in reason for reason in lighting["reasons"])

    unknown = classify_page(_page([_word("SHEET", 72, 36), _word("E9", 120, 36)], sheet="E9"))
    assert unknown["sheet_class"] == "unknown"
    assert unknown["confidence"] == "low"
    assert unknown["steps"] == []


def test_low_ocr_title_is_not_applied_and_high_ocr_is():
    low = classify_page(
        _page([_word("POWER", 72, 36, source="ocr", confidence=0.2)], sheet="E1", title="POWER")
    )
    assert low["sheet_class"] == "unknown"
    assert any("not applied" in reason.lower() for reason in low["reasons"])

    high = classify_page(
        _page([_word("POWER", 72, 36, source="ocr", confidence=0.84)], sheet="E4", title="POWER")
    )
    assert high["sheet_class"] == "power_plan"
    assert any("OCR" in reason for reason in high["reasons"])


def test_schedule_sheet_runs_schedules_and_a_plan_runs_symbols():
    pages = read_pdf(build_scheduled_pdf())
    rows = {row["sheet"]: row for row in classify_pages(pages)}
    assert rows["E1"]["sheet_class"] == "power_plan"
    assert rows["E2"]["sheet_class"] == "lighting_plan"
    schedule = rows["E3"]
    assert schedule["sheet_class"] in {"fixture_schedule", "panel_schedule"}
    assert "Fixture schedule" in schedule["step_labels"]
    assert "Panel schedule" in schedule["step_labels"]
    assert "Symbol detection" not in schedule["step_labels"]
    again = classify_pages(pages)
    assert [(row["sheet_class"], row["reasons"]) for row in rows.values()] == [
        (row["sheet_class"], row["reasons"]) for row in again
    ]


def test_symbol_detection_skips_a_schedule_sheet(tmp_path):
    document = fitz.open()
    legend = document.new_page(width=612, height=792)
    _text(legend, 72, 48, "SYMBOL LEGEND", size=12)
    draw_duplex(legend, 90, 140)
    _text(legend, 118, 143, "DUPLEX RECEPTACLE")
    sheet = document.new_page(width=792, height=792)
    _text(sheet, 72, 36, "SHEET E3", size=12)
    _text(sheet, 220, 36, "TITLE PANEL SCHEDULE", size=12)
    _text(sheet, 72, 80, "PANEL SCHEDULE LP1", size=12)
    draw_duplex(sheet, 400, 400)
    payload = document.tobytes()
    document.close()
    result = run_takeoff(payload)
    assert [item.label for item in result.items if item.measurement == "matched"] == []


def test_sample_classes_keep_the_counts(tmp_path):
    client = TestClient(create_app(tmp_path))
    response = client.post(
        "/jobs",
        files={"pdf": ("set.pdf", build_sample_pdf(), "application/pdf")},
        follow_redirects=False,
    )
    job_id = response.headers["location"].rstrip("/").split("/")[-1]
    page = client.get(f"/jobs/{job_id}/sheets").text
    assert "Power plan" in page
    assert "Lighting plan" in page
    assert "Legend / symbols" in page
    assert "Symbol detection" in page
    assert 'id="cleanup-note"' in page
    assert "off" in page
    review = client.get(f"/jobs/{job_id}/review").text
    assert 'id="power-device-count">4<' in review
    saved = client.post(
        f"/jobs/{job_id}/sheets/1",
        data={"sheet_class": "cover"},
        follow_redirects=False,
    )
    assert saved.status_code == 303
    overridden = client.get(f"/jobs/{job_id}/sheets").text
    assert "Estimator override" in overridden
    assert "Cover / index" in overridden
    recounted = client.get(f"/jobs/{job_id}/review").text
    assert 'id="power-device-count">4<' not in recounted
    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    assert "Sheet Index" in workbook.sheetnames
    assert "Routing" not in workbook.sheetnames
    text = " ".join(
        str(cell.value)
        for row in workbook["Sheet Index"].iter_rows()
        for cell in row
        if cell.value
    )
    assert "Cover / index" in text
    assert "Power plan" in text


def test_cleanup_drops_background_and_stays_off_for_shipped_benchmarks(tmp_path):
    border = _line(10, 10, 700, 10)
    faint = _line(100, 100, 108, 100, color=(0.85, 0.85, 0.85))
    dark = _line(100, 200, 108, 200, color=(0, 0, 0))
    clouds = [_curve(20 + index * 6, 400) for index in range(8)]
    kept = suppress_background([border, faint, dark, *clouds], 792, 792)
    assert dark in kept
    assert border not in kept
    assert faint not in kept
    assert clouds[0] not in kept

    plain = run_takeoff(build_sample_pdf(), cleanup=False)
    cleaned = run_takeoff(build_sample_pdf(), cleanup=True)
    assert [item.label for item in plain.items] == [item.label for item in cleaned.items]

    shipped = measure_cleanup(
        [ROOT / "benchmarks" / "synthetic-scheduled", ROOT / "benchmarks" / "synthetic-totals"]
    )
    assert shipped["improved"] is False
    assert BACKGROUND_CLEANUP is False
    assert shipped["enabled"] is False
    assert all(row["recall_on"] + 1e-9 >= row["recall_off"] for row in shipped["sets"])

    noisy = _noisy_set(tmp_path)
    effect = measure_cleanup([noisy])
    assert effect["improved"] is True
    assert effect["sets"][0]["recall_on"] + 1e-9 >= effect["sets"][0]["recall_off"]
    assert effect["sets"][0]["precision_on"] > effect["sets"][0]["precision_off"]


def _noisy_set(folder: Path) -> Path:
    document = fitz.open()
    legend = document.new_page(width=612, height=792)
    _text(legend, 72, 48, "SYMBOL LEGEND", size=12)
    draw_duplex(legend, 90, 140)
    _text(legend, 118, 143, "DUPLEX RECEPTACLE")
    sheet = document.new_page(width=792, height=792)
    _text(sheet, 72, 36, "SHEET E1", size=12)
    _text(sheet, 220, 36, "TITLE FIRST FLOOR POWER", size=12)
    draw_duplex(sheet, 200, 300)
    sheet.draw_circle(fitz.Point(460, 300), 8, color=(0.85, 0.85, 0.85), width=0.8)
    sheet.draw_line(fitz.Point(456, 298), fitz.Point(464, 298), color=(0.85, 0.85, 0.85), width=0.6)
    sheet.draw_line(fitz.Point(456, 302), fitz.Point(464, 302), color=(0.85, 0.85, 0.85), width=0.6)
    target = folder / "noisy"
    target.mkdir()
    (target / "drawing.pdf").write_bytes(document.tobytes())
    document.close()
    sheets = target / "truth" / "sheets"
    sheets.mkdir(parents=True)
    (sheets / "E1.json").write_text(
        json.dumps(
            {
                "sheet": "E1",
                "tolerance_pt": 18,
                "items": [
                    {"tag": "DUPLEX RECEPTACLE", "category": "power", "x": 200.0, "y": 300.0}
                ],
            }
        ),
        encoding="utf-8",
    )
    return target
