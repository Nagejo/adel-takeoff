"""Sheet scale is read, or the length is refused. It is not guessed."""

from decimal import Decimal
from io import BytesIO

import fitz
import openpyxl
import pytest
from fastapi.testclient import TestClient

from adel_takeoff.main import create_app
from adel_takeoff.scale import (
    ScaleInputError,
    assess_pdf,
    calibrate,
    measure_polyline,
    parse_points,
    parse_scale_text,
    polyline_length,
    scales_agree,
)
from tests.drawings import _text


def _pdf(draw) -> bytes:
    document = fitz.open()
    draw(document)
    payload = document.tobytes()
    document.close()
    return payload


def _bar(page, x, y, length, feet):
    page.draw_line(fitz.Point(x, y), fitz.Point(x + length, y), color=(0, 0, 0), width=1)
    page.draw_line(fitz.Point(x, y - 8), fitz.Point(x, y + 8), color=(0, 0, 0), width=1)
    page.draw_line(
        fitz.Point(x + length, y - 8),
        fitz.Point(x + length, y + 8),
        color=(0, 0, 0),
        width=1,
    )
    _text(page, x - 4, y + 20, "0", size=8)
    _text(page, x + length - 8, y + 20, str(feet), size=8)


def _sheet(document, name, title="PLAN"):
    page = document.new_page(width=900, height=792)
    _text(page, 72, 36, f"SHEET {name}", size=12)
    _text(page, 220, 36, f"TITLE {title}", size=12)
    return page


def test_scale_text_parses_printed_scales_and_does_not_guess():
    eighth = parse_scale_text('SCALE: 1/8" = 1\'-0"')
    assert eighth is not None
    assert eighth.label == '1/8" = 1\'-0"'
    assert eighth.feet_per_point == Decimal(8) / Decimal(72)
    assert parse_scale_text('3/16" = 1\'-0"').feet_per_point == (
        (Decimal(1) / (Decimal(3) / Decimal(16))) / Decimal(72)
    )
    assert parse_scale_text('1" = 10\'').feet_per_point == Decimal(10) / Decimal(72)
    assert parse_scale_text("1:100").label == "1:100"
    assert parse_scale_text("1:100").feet_per_point == (Decimal(100) / Decimal(12)) / Decimal(72)
    assert parse_scale_text("N.T.S.").nts
    assert parse_scale_text("NOT TO SCALE").nts
    for text in ("LP1-12", "SEE NOTE", "SHEET E1", "1/2", "GRID B-4", ""):
        assert parse_scale_text(text) is None
    eighth_feet = Decimal(8) / Decimal(72)
    assert scales_agree(eighth_feet, eighth_feet * Decimal("1.04"))
    assert not scales_agree(eighth_feet, eighth_feet * Decimal("1.06"))
    length = polyline_length(
        [(Decimal(0), Decimal(0)), (Decimal(72), Decimal(0)), (Decimal(72), Decimal(72))],
        numerator=Decimal(8),
        denominator=Decimal(72),
    )
    assert length == Decimal(16)
    with pytest.raises(ScaleInputError):
        parse_points("2 EA, 4")


def test_text_and_bar_that_agree_record_source_and_region():
    def draw(document):
        page = _sheet(document, "E1")
        _text(page, 72, 80, 'SCALE: 1/8" = 1\'-0"', size=11)
        _bar(page, 72, 160, 144, 16)

    sheet = assess_pdf(_pdf(draw))[0]
    assert sheet["sheet"] == "E1"
    assert sheet["status"] == "valid"
    region = sheet["regions"][0]
    assert region["source"] == "text and graphic bar"
    assert region["label"] == '1/8" = 1\'-0"'
    assert Decimal(region["feet_per_point"]) == Decimal(8) / Decimal(72)
    kinds = {item["kind"] for item in region["read_from"]}
    assert kinds == {"text", "graphic bar"}
    assert all(len(item["bbox"]) == 4 for item in region["read_from"])
    measured = measure_polyline(
        sheet,
        [(Decimal(100), Decimal(400)), (Decimal(172), Decimal(400))],
    )
    assert measured["ok"] is True
    assert measured["length_ft"] == Decimal(8)
    assert measured["scale_source"] == "text and graphic bar"


def test_bar_and_text_conflict_refuses_without_choosing_a_scale():
    def draw(document):
        page = _sheet(document, "E1")
        _text(page, 72, 80, 'SCALE: 1/8" = 1\'-0"', size=11)
        _bar(page, 72, 160, 144, 32)

    sheet = assess_pdf(_pdf(draw))[0]
    assert sheet["status"] == "conflict"
    assert "disagree" in sheet["reason"].lower()
    assert sheet["regions"][0]["feet_per_point"] == ""
    refused = measure_polyline(sheet, [(Decimal(0), Decimal(0)), (Decimal(72), Decimal(0))])
    assert refused["ok"] is False
    assert "disagree" in refused["reason"].lower()
    assert "length_ft" not in refused


def test_nts_and_missing_scale_refuse():
    def draw(document):
        noted = _sheet(document, "E2", "DETAIL")
        _text(noted, 72, 80, "NOT TO SCALE", size=12)
        _sheet(document, "E3", "UNTITLED PLAN")

    sheets = assess_pdf(_pdf(draw))
    assert sheets[0]["status"] == "nts"
    assert "not to scale" in sheets[0]["reason"].lower()
    assert sheets[0]["regions"][0]["feet_per_point"] == ""
    refused = measure_polyline(sheets[0], [(Decimal(0), Decimal(0)), (Decimal(72), Decimal(0))])
    assert refused["ok"] is False
    assert sheets[1]["status"] == "missing"
    assert "no scale is printed" in sheets[1]["reason"].lower()
    assert measure_polyline(sheets[1], [(Decimal(0), Decimal(0)), (Decimal(10), Decimal(0))])["ok"] is False


def test_viewports_keep_their_own_scales_and_loose_scales_stay_ambiguous():
    def draw(document):
        page = _sheet(document, "E4", "DETAILS")
        page.draw_rect(fitz.Rect(40, 80, 280, 320), color=(0, 0, 0), width=0.8)
        _text(page, 60, 110, '1/4" = 1\'-0"', size=10)
        page.draw_rect(fitz.Rect(320, 80, 620, 320), color=(0, 0, 0), width=0.8)
        _text(page, 340, 110, '1/2" = 1\'-0"', size=10)
        loose = _sheet(document, "E5", "TWO SCALES")
        _text(loose, 72, 80, '1/8" = 1\'-0"', size=11)
        _text(loose, 72, 110, '1/4" = 1\'-0"', size=11)

    detailed, loose = assess_pdf(_pdf(draw))
    assert detailed["status"] == "regional"
    valid = [region for region in detailed["regions"] if region["status"] == "valid"]
    assert {region["label"] for region in valid} == {'1/4" = 1\'-0"', '1/2" = 1\'-0"'}
    assert all(region["viewport"] for region in valid)
    quarter = measure_polyline(
        detailed,
        [(Decimal(80), Decimal(200)), (Decimal(152), Decimal(200))],
    )
    half = measure_polyline(
        detailed,
        [(Decimal(360), Decimal(200)), (Decimal(432), Decimal(200))],
    )
    outside = measure_polyline(
        detailed,
        [(Decimal(700), Decimal(500)), (Decimal(772), Decimal(500))],
    )
    assert quarter["ok"] and quarter["length_ft"] == Decimal(4)
    assert half["ok"] and half["length_ft"] == Decimal(2)
    assert outside["ok"] is False
    assert loose["status"] == "ambiguous"
    assert "not tied to separate viewports" in loose["reason"].lower()
    assert loose["regions"][0]["feet_per_point"] == ""
    assert measure_polyline(loose, [(Decimal(0), Decimal(0)), (Decimal(72), Decimal(0))])["ok"] is False


def test_manual_scale_overrides_detection_and_records_who_set_it():
    def draw(document):
        page = _sheet(document, "E6", "NO SCALE")
        _text(page, 72, 80, "NOT TO SCALE", size=12)

    feet_per_point, label = calibrate(
        Decimal(0), Decimal(0), Decimal(144), Decimal(0), Decimal(16)
    )
    assert feet_per_point == Decimal(8) / Decimal(72)
    manual = {
        0: {
            "label": '1/8" = 1\'-0"',
            "feet_per_point": format(Decimal(8) / Decimal(72), "f"),
            "set_by": "Ada Estimator",
            "set_at": "2026-10-17T12:00:00+00:00",
            "read_from": [{"kind": "picked scale", "label": label, "bbox": None}],
        }
    }
    sheet = assess_pdf(_pdf(draw), manual)[0]
    assert sheet["status"] == "manual"
    assert sheet["overridden"] is True
    assert sheet["regions"][0]["source"] == "manual"
    assert sheet["regions"][0]["set_by"] == "Ada Estimator"
    assert sheet["regions"][0]["set_at"] == "2026-10-17T12:00:00+00:00"
    assert "overrides" in sheet["reason"].lower()
    assert sheet["detected"][0]["status"] == "nts"
    measured = measure_polyline(
        sheet,
        [(Decimal(10), Decimal(10)), (Decimal(82), Decimal(10))],
    )
    assert measured["ok"] is True
    assert measured["length_ft"] == Decimal(8)
    assert measured["scale_source"] == "manual"


def _table(sheet, header_name):
    header_row = None
    for row in sheet.iter_rows():
        if header_name in [cell.value for cell in row]:
            header_row = row[0].row
            break
    assert header_row is not None
    headers = {cell.value: cell.column - 1 for cell in sheet[header_row] if cell.value}
    found = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or row[headers[header_name]] is None:
            if found:
                break
            continue
        found.append({name: row[index] for name, index in headers.items() if index < len(row)})
    return found


def test_measure_tool_exports_measured_lengths_and_refuses_nts(tmp_path):
    def draw(document):
        page = _sheet(document, "E1", "POWER")
        _text(page, 72, 80, 'SCALE: 1/8" = 1\'-0"', size=11)
        _bar(page, 72, 160, 144, 16)
        refused = _sheet(document, "E2", "DETAIL")
        _text(refused, 72, 80, "N.T.S.", size=12)

    client = TestClient(create_app(tmp_path))
    uploaded = client.post(
        "/jobs",
        files={"pdf": ("scaled.pdf", _pdf(draw), "application/pdf")},
        follow_redirects=False,
    )
    job_id = uploaded.headers["location"].rstrip("/").split("/")[-1]
    page = client.get(f"/jobs/{job_id}/scale")
    assert page.status_code == 200
    assert "1/8" in page.text
    assert "text and graphic bar" in page.text
    assert "not to scale" in page.text.lower()
    assert "data/scales/" in page.text or "scales/" in page.text

    saved = client.post(
        f"/jobs/{job_id}/measurements",
        data={"page_index": "0", "points": "100,400; 172,400", "set_by": "Ada Estimator"},
        follow_redirects=True,
    )
    assert saved.status_code == 200
    assert "method measured" in saved.text
    refused = client.post(
        f"/jobs/{job_id}/measurements",
        data={"page_index": "1", "points": "100,400; 172,400", "set_by": "Ada Estimator"},
        follow_redirects=True,
    )
    assert refused.status_code == 400
    assert "not to scale" in refused.text.lower()
    assert refused.text.count("method measured") == 1

    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    assert "Measured" in workbook.sheetnames
    assert "Routing" not in workbook.sheetnames
    lengths = _table(workbook["Measured"], "Method")
    assert len(lengths) == 1
    assert lengths[0]["Sheet"] == "E1"
    assert lengths[0]["Scale source"] == "text and graphic bar"
    assert lengths[0]["Method"] == "measured"
    assert Decimal(str(lengths[0]["Length (ft)"])) == Decimal(8)
    materials = "\n".join(
        "" if cell.value is None else str(cell.value)
        for row in workbook["Materials"].iter_rows()
        for cell in row
    )
    assert "allowance" in materials.lower()
    assert "method" not in materials.lower()
    picked = client.post(
        f"/jobs/{job_id}/scales",
        data={"page_index": "1", "mode": "picked", "picked": "1:100", "set_by": "Ada Estimator"},
        follow_redirects=True,
    )
    assert picked.status_code == 200
    assert "Ada Estimator" in picked.text
    assert "manual" in picked.text.lower()
    measured = client.post(
        f"/jobs/{job_id}/measurements",
        data={"page_index": "1", "points": "0,0; 72,0", "set_by": "Ada Estimator"},
        follow_redirects=True,
    )
    assert measured.status_code == 200
    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    lengths = _table(workbook["Measured"], "Method")
    manual = next(row for row in lengths if row["Sheet"] == "E2")
    assert manual["Scale source"] == "manual"
    assert manual["Method"] == "measured"
    assert Decimal(str(manual["Length (ft)"])).quantize(Decimal("0.0001")) == (
        Decimal(100) / Decimal(12)
    ).quantize(Decimal("0.0001"))
