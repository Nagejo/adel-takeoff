"""Assembly expansion stays on the estimator's quantities and hours."""

import json
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient

from adel_takeoff.assemblies import (
    AssemblyError,
    apply_assemblies,
    apply_overrides,
    load_starter,
    match_assembly,
    material_rows,
    starter_catalog_path,
)
from adel_takeoff.constants import LABOR_RATES, UNCONFIRMED_ASSEMBLY_WARNING
from adel_takeoff.costing import MarkupRejected, apply_markup, parse_markup
from adel_takeoff.main import create_app
from tests.drawings import build_sample_pdf
from tests.test_app import _marked, _rows


def _job():
    return {f"rate_{key}": str(amount) for key, _label, amount in LABOR_RATES}


def test_starter_values_are_examples_and_not_a_published_standard():
    text = starter_catalog_path().read_text(encoding="utf-8")
    assert "EXAMPLE" in text
    assert "NECA" not in text
    starter = load_starter()
    assert {item.id for item in starter} >= {
        "duplex",
        "gfci",
        "switch",
        "dimmer",
        "recessed-fixture",
        "junction-box",
        "data-lv",
    }
    assert all(item.example and not item.confirmed for item in starter)
    assert match_assembly("GFCI DUPLEX RECEPTACLE", starter).id == "gfci"
    assert match_assembly("DIMMER SWITCH", starter).id == "dimmer"
    assert match_assembly("DATA OUTLET", starter).id == "data-lv"
    assert match_assembly("JUNCTION BOX", starter).id == "junction-box"
    assert match_assembly("SINGLE POLE SWITCH", starter).id == "switch"
    assert match_assembly("RECESSED LIGHT", starter).id == "recessed-fixture"
    assert match_assembly("LANDSCAPE LIGHT", starter) is None


def test_expansion_uses_count_times_per_unit_and_skips_excluded():
    starter = load_starter()
    rows = [
        {"id": "1", "label": "DUPLEX RECEPTACLE", "effective": 2, "direct_cost": Decimal("0")},
        {"id": "9", "label": "DUPLEX RECEPTACLE", "effective": None, "direct_cost": None},
    ]
    apply_assemblies(rows, _job(), starter)
    assert rows[0]["assembly_name"] == "Duplex receptacle"
    assert rows[0]["assembly_version"] == "example-1"
    assert rows[0]["assembly_labor_hours"] == Decimal("2")
    assert rows[0]["direct_cost"] == Decimal("0")
    assert rows[0]["assembly_confirmed"] is False
    assert rows[1]["assembly_id"] == ""

    line = {
        "line_id": "1",
        "item_ids": ["1"],
        "label": "DUPLEX RECEPTACLE",
        "effective": 2,
        "assembly_id": "duplex",
        "assembly_name": "Duplex receptacle",
        "assembly_version": "example-1",
        "assembly_confirmed": False,
        "assembly_labor_hours": Decimal("2.00"),
        "assembly_hours_per_unit": Decimal("1.00"),
        "assembly_rate_key": "journeyman",
    }
    materials = material_rows([line], starter)
    device = next(row for row in materials if row["role"] == "device")
    raceway = next(row for row in materials if row["role"] == "raceway")
    conductor = next(row for row in materials if row["role"] == "conductor")
    assert device["quantity"] == Decimal("2")
    assert device["item_ids"] == ["1"]
    assert raceway["quantity"] == Decimal("20")
    assert raceway["allowance"] == "allowance"
    assert raceway["size"] == "1/2"
    assert conductor["quantity"] == Decimal("60")
    assert conductor["conductor_count"] == "3"
    assert "9" not in device["item_ids"]


def test_override_replaces_only_the_edited_value_and_confirmation_prices_labor():
    starter = load_starter()
    with pytest.raises(AssemblyError):
        apply_overrides(
            starter,
            {"overrides": {"duplex": {"parts": {"device": {"quantity": "2 EA"}}}}},
        )
    resolved = apply_overrides(
        starter,
        {
            "overrides": {
                "duplex": {
                    "confirmed": True,
                    "parts": {"raceway": {"quantity": "4"}},
                }
            }
        },
    )
    duplex = match_assembly("DUPLEX RECEPTACLE", resolved)
    assert duplex.confirmed is True
    assert duplex.version == "example-1+job"
    assert duplex.labor_hours == Decimal("1")
    raceway = next(part for part in duplex.parts if part.role == "raceway")
    device = next(part for part in duplex.parts if part.role == "device")
    assert raceway.quantity == Decimal("4")
    assert device.quantity == Decimal("1")
    rows = [
        {"id": "1", "label": "DUPLEX RECEPTACLE", "effective": 2, "direct_cost": Decimal("10")}
    ]
    apply_assemblies(rows, _job(), resolved)
    assert rows[0]["assembly_labor_hours"] == Decimal("2")
    assert rows[0]["direct_cost"] == Decimal("240")
    assert apply_markup(rows[0]["direct_cost"], Decimal("0.35")) == Decimal("324.00")
    assert apply_markup(rows[0]["direct_cost"], None) is None
    assert parse_markup("") is None
    assert parse_markup("5") == Decimal("5")
    with pytest.raises(MarkupRejected, match="greater than 5"):
        parse_markup("5.01")
    assert "unconfirmed assembly" in UNCONFIRMED_ASSEMBLY_WARNING.lower()
    assert "NECA" not in UNCONFIRMED_ASSEMBLY_WARNING


def _sheet_text(sheet) -> str:
    return "\n".join(
        " ".join("" if cell.value is None else str(cell.value) for cell in row)
        for row in sheet.iter_rows()
    )


def _table(sheet, header_name):
    header_row = None
    for row in sheet.iter_rows():
        values = [cell.value for cell in row]
        if header_name in values:
            header_row = row[0].row
            break
    assert header_row is not None
    headers = {cell.value: cell.column - 1 for cell in sheet[header_row] if cell.value}
    found = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or row[headers[header_name]] is None:
            if found:
                break
            continue
        found.append({name: row[index] for name, index in headers.items() if index < len(row)})
    return found


def _upload(tmp_path):
    client = TestClient(create_app(tmp_path))
    response = client.post(
        "/jobs",
        files={"pdf": ("set.pdf", build_sample_pdf(), "application/pdf")},
        follow_redirects=False,
    )
    assert response.status_code == 303
    job_id = response.headers["location"].rstrip("/").split("/")[-1]
    return client, job_id


def _counted(html, label):
    return [
        row
        for row in _rows(html)
        if row["label"] == label and row["effective"] not in ("", None)
    ]


def test_export_warns_until_confirmed_and_skips_rejected_devices(tmp_path):
    client, job_id = _upload(tmp_path)
    page = client.get(f"/jobs/{job_id}/assemblies")
    assert page.status_code == 200
    assert "unconfirmed assembly" in page.text.lower()
    assert "EXAMPLE" in page.text
    assert "assemblies/starter.json" in page.text
    assert f"assemblies/{job_id}.json" in page.text
    assert "Raceway allowance" in page.text
    review = client.get(f"/jobs/{job_id}/review").text
    assert "unconfirmed assembly" in review.lower()
    assert _marked(review, "power-device-count") == "4"
    duplexes = _counted(review, "DUPLEX RECEPTACLE")
    switches = _counted(review, "SINGLE POLE SWITCH")
    lights = _counted(review, "RECESSED LIGHT")
    assert duplexes and switches and lights

    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    assert "Materials" in workbook.sheetnames
    assert "Labor" in workbook.sheetnames
    assert "Routing" not in workbook.sheetnames
    assert "NECA" not in _sheet_text(workbook["Quantities"])
    assert "unconfirmed assembly" in _sheet_text(workbook["Quantities"]).lower()
    assert "unconfirmed assembly" in _sheet_text(workbook["Labor"]).lower()
    quantities = _table(workbook["Quantities"], "How counted")
    duplex_line = next(row for row in quantities if row["Assembly"] == "Duplex receptacle")
    assert duplex_line["Assembly version"] == "example-1"
    assert duplex_line["Assembly confirmed"] == "no"
    materials = _table(workbook["Materials"], "Source item IDs")
    device = next(
        row
        for row in materials
        if row["Assembly"] == "Duplex receptacle" and row["Part"] == "device"
    )
    assert Decimal(str(device["Quantity"])) == sum(int(row["effective"]) for row in duplexes)
    assert device["Confirmed"] == "no"
    source_ids = {item.strip() for item in str(device["Source item IDs"]).split(",")}
    assert source_ids == {row["id"] for row in duplexes}
    raceway = next(
        row
        for row in materials
        if row["Assembly"] == "Duplex receptacle" and row["Part"] == "raceway"
    )
    assert raceway["Allowance"] == "allowance"
    assert Decimal(str(raceway["Quantity"])) == Decimal(len(duplexes)) * Decimal("10")
    conductor = next(
        row
        for row in materials
        if row["Assembly"] == "Duplex receptacle" and row["Part"] == "conductor"
    )
    assert str(conductor["Conductor count"]) == "3"
    assert "Landscape" not in _sheet_text(workbook["Materials"])
    labor = _table(workbook["Labor"], "Hours per unit")
    assert any(row["Rate"] == "Journeyman" and Decimal(str(row["Hours per unit"])) == Decimal("1") for row in labor)
    rate_text = _sheet_text(workbook["Labor"])
    for amount in ("115", "130", "150", "65"):
        assert amount in rate_text
    direct = next(
        row[1].value
        for row in workbook["Quantities"].iter_rows()
        if row[0].value == "Job direct cost"
    )
    assert Decimal(str(direct)) == Decimal("0")

    rejected_id = duplexes[0]["id"]
    rejected = client.post(
        f"/jobs/{job_id}/items/{rejected_id}/review",
        data={
            "action": "reject",
            "label": "DUPLEX RECEPTACLE",
            "panel": duplexes[0]["panel"],
            "circuit": duplexes[0]["circuit"],
        },
        follow_redirects=True,
    )
    assert rejected.status_code == 200
    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    materials = _table(workbook["Materials"], "Source item IDs")
    device = next(
        row
        for row in materials
        if row["Assembly"] == "Duplex receptacle" and row["Part"] == "device"
    )
    assert Decimal(str(device["Quantity"])) == Decimal(len(duplexes) - 1)
    assert rejected_id not in str(device["Source item IDs"]).split(", ")

    starter = load_starter()
    form = {}
    for assembly in starter:
        form[f"labor_{assembly.id}"] = format(assembly.labor_hours, "f")
        form[f"rate_{assembly.id}"] = assembly.labor_rate_key
        for part in assembly.parts:
            form[f"qty_{assembly.id}_{part.role}"] = format(part.quantity, "f")
            form[f"size_{assembly.id}_{part.role}"] = part.size
            form[f"count_{assembly.id}_{part.role}"] = part.count
    form["qty_duplex_raceway"] = "4"
    for assembly_id in ("duplex", "switch", "recessed-fixture"):
        form[f"confirmed_{assembly_id}"] = "yes"
    saved = client.post(f"/jobs/{job_id}/assemblies", data=form, follow_redirects=True)
    assert saved.status_code == 200
    assert "unconfirmed assembly" not in saved.text.lower()
    stored = json.loads((tmp_path / "assemblies" / f"{job_id}.json").read_text())
    assert stored["overrides"]["duplex"]["confirmed"] is True
    assert stored["overrides"]["duplex"]["parts"]["raceway"]["quantity"] == "4"

    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    assert "unconfirmed assembly" not in _sheet_text(workbook["Quantities"]).lower()
    assert "not a published labor standard" in _sheet_text(workbook["Labor"]).lower()
    materials = _table(workbook["Materials"], "Source item IDs")
    raceway = next(
        row
        for row in materials
        if row["Assembly"] == "Duplex receptacle" and row["Part"] == "raceway"
    )
    remaining = len(duplexes) - 1
    assert Decimal(str(raceway["Quantity"])) == Decimal(remaining) * Decimal("4")
    assert raceway["Assembly version"] == "example-1+job"
    quantities = _table(workbook["Quantities"], "How counted")
    switch_line = next(row for row in quantities if row["Assembly"] == "Switch")
    assert switch_line["Assembly version"] == "example-1"
    assert switch_line["Assembly confirmed"] == "yes"
    direct = next(
        row[1].value
        for row in workbook["Quantities"].iter_rows()
        if row[0].value == "Job direct cost"
    )
    expected = (
        Decimal(remaining) * Decimal("1") * Decimal("115")
        + Decimal(len(switches)) * Decimal("0.50") * Decimal("115")
        + Decimal(len(lights)) * Decimal("1") * Decimal("115")
    )
    assert Decimal(str(direct)) == expected
    labor_total = next(
        row[1].value
        for row in workbook["Labor"].iter_rows()
        if row[0].value == "Confirmed direct labor"
    )
    assert Decimal(str(labor_total)) == expected

    refused = client.post(
        f"/jobs/{job_id}/assemblies",
        data={**form, "qty_duplex_device": "2 EA"},
        follow_redirects=True,
    )
    assert refused.status_code == 400
    assert "must be a number" in refused.text
    still = json.loads((tmp_path / "assemblies" / f"{job_id}.json").read_text())
    assert still["overrides"]["duplex"]["parts"]["raceway"]["quantity"] == "4"
