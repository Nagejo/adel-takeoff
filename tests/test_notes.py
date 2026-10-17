"""General notes, keynote links, disposition warning, and the RFI log."""

import re
import zipfile
from io import BytesIO

import fitz
import openpyxl
from fastapi.testclient import TestClient

from adel_takeoff.coordination import drawing_discrepancies, rfi_docx_bytes
from adel_takeoff.extract import PageData, Word, read_pdf
from adel_takeoff.main import create_app
from adel_takeoff.notes import attention_for, link_keynotes, parse_notes
from adel_takeoff.present import build_review, present_item
from adel_takeoff.schedule_link import attach_schedules
from tests.drawings import draw_duplex, draw_unknown

STREET_RE = re.compile(
    r"\d{1,6}\s+[A-Za-z0-9.'-]+(?:\s+[A-Za-z0-9.'-]+){0,3}\s+"
    r"(?:Street|St\.|Avenue|Ave\.|Road|Rd\.|Boulevard|Blvd\.|Drive|Dr\.|Lane|Ln\.)",
    re.IGNORECASE,
)


def _text(page, x, y, text, size=9):
    page.insert_text(fitz.Point(x, y), text, fontsize=size, fontname="helv")


def build_notes_pdf() -> bytes:
    document = fitz.open()
    legend = document.new_page(width=612, height=792)
    _text(legend, 72, 48, "SYMBOL LEGEND", size=12)
    draw_duplex(legend, 90, 100)
    _text(legend, 118, 103, "DUPLEX RECEPTACLE")

    sheet = document.new_page(width=792, height=792)
    _text(sheet, 72, 36, "SHEET E1", size=12)
    _text(sheet, 220, 36, "TITLE FIRST FLOOR POWER", size=12)
    draw_duplex(sheet, 200, 220)
    _text(sheet, 230, 223, "KN-1")
    draw_unknown(sheet, 420, 220)
    _text(sheet, 72, 470, "KEYNOTES", size=12)
    _text(sheet, 72, 492, "1. EMPTY PVC FOR TELCO AND CATV FROM THE DEMARC.")
    _text(sheet, 72, 560, "GENERAL NOTES", size=12)
    notes = (
        "1. PROVIDE EMPTY PVC FOR TELCO/CATV.",
        "2. LOW VOLTAGE PREWIRE SHALL MATCH THE FLOOR PLAN.",
        "3. ROMEX IS ALLOWED WITH OWNER APPROVAL.",
        "4. PROVIDE A 1-YEAR WARRANTY.",
        "5. FIELD CHANGES AT NO EXTRA CHARGE.",
        "6. CONNECT ITEMS SHOWN ON OTHER TRADES DRAWINGS.",
    )
    for offset, line in enumerate(notes):
        _text(sheet, 72, 584 + offset * 16, line)
    payload = document.tobytes()
    document.close()
    return payload


def build_discrepancy_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page(width=792, height=792)
    _text(page, 72, 36, "SHEET E4", size=12)
    _text(page, 220, 36, "TITLE SCHEDULES AND RISER", size=12)
    _text(page, 72, 80, "PANEL SCHEDULE LP1 42 CIRCUITS", size=12)
    for x, label in ((72, "CIRCUIT"), (180, "DESCRIPTION"), (360, "BREAKER")):
        _text(page, x, 108, label)
    _text(page, 72, 128, "1")
    _text(page, 180, 128, "LIGHTS")
    _text(page, 360, 128, "20A")
    _text(page, 72, 148, "2")
    _text(page, 180, 148, "RECEPTACLES")
    _text(page, 360, 148, "20A")
    _text(page, 72, 230, "RISER DIAGRAM", size=12)
    _text(page, 72, 252, "PANEL HP9")
    _text(page, 72, 360, 'SCALE 1/4" = 1\'-0"')
    _text(page, 480, 760, 'SCALE 1/8" = 1\'-0"')
    payload = document.tobytes()
    document.close()
    return payload


def _upload(tmp_path, payload, name="notes.pdf"):
    client = TestClient(create_app(tmp_path))
    response = client.post(
        "/jobs",
        files={"pdf": (name, payload, "application/pdf")},
        follow_redirects=False,
    )
    assert response.status_code == 303
    job_id = response.headers["location"].rstrip("/").split("/")[-1]
    return client, job_id


def test_numbered_general_notes_keep_sheet_and_location():
    notes = [note for note in parse_notes(read_pdf(build_notes_pdf())) if note["kind"] == "general"]
    assert [note["number"] for note in notes] == ["1", "2", "3", "4", "5", "6"]
    assert {note["sheet"] for note in notes} == {"E1"}
    assert all(note["x"] > 0 and note["y"] > 0 for note in notes)
    by_number = {note["number"]: note for note in notes}
    assert "TELCO/CATV" in by_number["1"]["text"]
    assert "PREWIRE" in by_number["2"]["text"]
    assert "ROMEX" in by_number["3"]["text"]
    assert "WARRANTY" in by_number["4"]["text"]
    assert "NO EXTRA CHARGE" in by_number["5"]["text"]
    assert "OTHER TRADES" in by_number["6"]["text"]
    assert by_number["3"]["disposition"] == ""
    assert "deduct alternate" in by_number["3"]["assumption"]
    assert by_number["5"]["who"] == "GC"
    assert "change-order" in by_number["5"]["why"]
    assert attention_for("hello").get("assumption") == ""


def test_ocr_words_parse_the_same_numbered_note():
    words = []
    for text, y, confidence in (
        ("GENERAL", 100, 0.4),
        ("NOTES", 100, 0.4),
        ("1.", 118, 0.95),
        ("PROVIDE", 118, 0.95),
        ("EMPTY", 118, 0.4),
        ("PVC", 118, 0.95),
    ):
        words.append(
            Word(
                x0=72 + len(words) * 40,
                y0=y,
                x1=100 + len(words) * 40,
                y1=y + 8,
                text=text,
                source="ocr",
                confidence=confidence,
            )
        )
    page = PageData(index=0, width=612, height=792, words=words, sheet="E9", text_source="ocr")
    notes = parse_notes([page])
    assert len(notes) == 1
    assert notes[0]["text"] == "PROVIDE EMPTY PVC"
    assert notes[0]["source"] == "ocr"
    assert notes[0]["ocr_check"] is True
    assert notes[0]["sheet"] == "E9"


def test_keynote_tag_attaches_to_the_nearby_item():
    pages = read_pdf(build_notes_pdf())
    notes = parse_notes(pages)
    keynotes = [note for note in notes if note["kind"] == "keynote"]
    assert len(keynotes) == 1
    assert keynotes[0]["tags"]
    row = present_item(
        {
            "id": "2",
            "label": "DUPLEX RECEPTACLE",
            "category": "power",
            "measurement": "matched",
            "measured_quantity": 1,
            "sheets": ["E1"],
            "panel": "",
            "circuit": "",
            "link_flag": "",
            "deduped": 0,
            "method": "legend_match",
            "match_score": "1",
            "review_status": "accepted",
            "exclusion": "",
            "override_quantity": None,
            "occurrences": [
                {
                    "page_index": 1,
                    "sheet": "E1",
                    "x0": 192,
                    "y0": 212,
                    "x1": 208,
                    "y1": 228,
                }
            ],
        },
        {"rate_journeyman": "115", "scope_power": "ec"},
    )
    far = present_item(
        {
            "id": "9",
            "label": "DUPLEX RECEPTACLE",
            "category": "power",
            "measurement": "matched",
            "measured_quantity": 1,
            "sheets": ["E1"],
            "panel": "",
            "circuit": "",
            "link_flag": "",
            "deduped": 0,
            "method": "legend_match",
            "match_score": "1",
            "review_status": "accepted",
            "exclusion": "",
            "override_quantity": None,
            "occurrences": [
                {
                    "page_index": 1,
                    "sheet": "E1",
                    "x0": 500,
                    "y0": 500,
                    "x1": 516,
                    "y1": 516,
                }
            ],
        },
        {"rate_journeyman": "115", "scope_power": "ec"},
    )
    link_keynotes(notes, [row, far])
    assert keynotes[0]["item_ids"] == ["2"]
    assert row["keynotes"][0]["text"].startswith("EMPTY PVC")
    assert far["keynotes"] == []


def test_discrepancy_log_covers_riser_scale_and_panel_header():
    payload = build_discrepancy_pdf()
    pages = read_pdf(payload)
    linked = attach_schedules([], pages)
    rows = drawing_discrepancies(pages, linked["discrepancies"], linked["panel_rows"], payload)
    categories = {row["category"] for row in rows}
    assert "panel header vs drawn circuits" in categories
    assert "panel schedule vs one-line/riser" in categories
    assert "title-block vs view scale" in categories
    header = next(row for row in rows if row["category"] == "panel header vs drawn circuits")
    assert header["plan"] == 2
    assert header["schedule"] == 42
    riser = [row for row in rows if row["category"] == "panel schedule vs one-line/riser"]
    assert any(row["item"] == "HP9" for row in riser)
    assert any(row["item"] == "LP1" for row in riser)
    scale = next(row for row in rows if row["category"] == "title-block vs view scale")
    assert '1/8"' in scale["detail"] and '1/4"' in scale["detail"]


def test_undisposed_notes_warn_and_a_disposition_clears_the_warning(tmp_path):
    client, job_id = _upload(tmp_path, build_notes_pdf())
    page = client.get(f"/jobs/{job_id}/notes")
    assert page.status_code == 200
    assert 'id="notes-warning"' in page.text
    assert "TELCO/CATV" in page.text
    assert "ROMEX" in page.text
    assert "KN-1" in page.text or "EMPTY PVC FOR TELCO AND CATV" in page.text
    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    assert "Notes Disposition" in workbook.sheetnames
    assert "Routing" not in workbook.sheetnames
    warning = "\n".join(
        str(cell.value)
        for row in workbook["Notes Disposition"].iter_rows()
        for cell in row
        if cell.value
    )
    assert "no disposition" in warning
    quantities = "\n".join(
        str(cell.value) for row in workbook["Quantities"].iter_rows() for cell in row if cell.value
    )
    assert "no disposition" in quantities
    for number in range(1, 7):
        saved = client.post(
            f"/jobs/{job_id}/notes/E1:general:{number}",
            data={"disposition": "included" if number != 3 else "alternate_deduct"},
            follow_redirects=False,
        )
        assert saved.status_code == 303
    ready = client.get(f"/jobs/{job_id}/notes")
    assert 'id="notes-warning"' not in ready.text
    assert 'id="notes-ready"' in ready.text
    again = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    text = "\n".join(
        str(cell.value)
        for sheet in again.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value
    )
    assert "no disposition" not in text
    assert "Alternate, deduct" in text


def test_rfi_log_excel_and_docx_have_the_columns_and_no_street(tmp_path):
    client, job_id = _upload(tmp_path, build_notes_pdf())
    marked = client.post(
        f"/jobs/{job_id}/notes/E1:general:5",
        data={"disposition": "rfi"},
        follow_redirects=False,
    )
    assert marked.status_code == 303
    added = client.post(
        f"/jobs/{job_id}/coordination",
        data={
            "sheet": "E1",
            "location": "72, 640",
            "item": "Warranty",
            "question": "Is the printed warranty one year?",
            "why": "A longer warranty changes the bid.",
            "who": "Owner",
            "assumption": "Price the printed 1-year term.",
            "status": "open",
        },
        follow_redirects=False,
    )
    assert added.status_code == 303
    page = client.get(f"/jobs/{job_id}/coordination")
    assert "Unlegended symbol" in page.text
    assert "FIELD CHANGES" in page.text
    assert "Is the printed warranty one year?" in page.text
    assert "Engineer" in page.text and "Architect" in page.text
    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    sheet = workbook["RFI log"]
    headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert headers == [
        "ID",
        "Sheet",
        "Location",
        "Item",
        "Question",
        "Why it matters",
        "Who answers",
        "Suggested pricing assumption",
        "Status",
    ]
    blob = "\n".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)
    assert "Unlegended symbol" in blob
    assert "FIELD CHANGES AT NO EXTRA CHARGE" in blob
    assert "Is the printed warranty one year?" in blob
    assert "Owner" in blob
    assert STREET_RE.search(blob) is None
    document = client.get(f"/jobs/{job_id}/rfi.docx")
    assert document.status_code == 200
    archive = zipfile.ZipFile(BytesIO(document.content))
    xml = archive.read("word/document.xml").decode()
    assert "Coordination and RFI log" in xml
    assert "Adel Electric, LLC" in xml
    assert "Why it matters" in xml
    assert "Who answers" in xml
    assert "Suggested pricing assumption" in xml
    assert "FIELD CHANGES AT NO EXTRA CHARGE" in xml
    assert "Is the printed warranty one year?" in xml
    assert STREET_RE.search(xml) is None
    assert "violation" not in xml.lower()
    assert "Home Entertainment Design" not in xml
    bare = rfi_docx_bytes("Job", [])
    bare_xml = zipfile.ZipFile(BytesIO(bare)).read("word/document.xml").decode()
    assert "Coordination and RFI log" in bare_xml
    assert STREET_RE.search(bare_xml) is None
