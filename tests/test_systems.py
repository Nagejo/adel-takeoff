"""Low voltage, fire alarm, and controls: tags, scope, rough-in, and floors."""

from io import BytesIO

import fitz
import openpyxl
from fastapi.testclient import TestClient

from adel_takeoff.assemblies import apply_assemblies, load_starter, material_rows
from adel_takeoff.classify import classify_label, posted_scope, scope_mode
from adel_takeoff.lines import quantity_lines
from adel_takeoff.main import create_app
from adel_takeoff.present import build_review, present_item
from adel_takeoff.systems import fire_alarm_by_floor, scope_split
from adel_takeoff.takeoff import run_takeoff
from tests.drawings import _text, draw_duplex, draw_switch
from tests.test_app import _save


def test_label_rules_keep_each_system_separate():
    expected = {
        "DUPLEX RECEPTACLE": ("power", "duplex"),
        "DATA OUTLET": ("low_voltage", "data"),
        "VOICE OUTLET": ("low_voltage", "voice"),
        "WAP": ("low_voltage", "wap"),
        "TV OUTLET": ("low_voltage", "tv"),
        "COAX OUTLET": ("low_voltage", "coax"),
        "SPEAKER": ("low_voltage", "speaker"),
        "CAMERA": ("low_voltage", "camera"),
        "DOOR STATION": ("low_voltage", "intercom"),
        "CARD READER": ("low_voltage", "access_control"),
        "SMOKE DETECTOR": ("fire_alarm", "smoke"),
        "HEAT DETECTOR": ("fire_alarm", "heat"),
        "CO DETECTOR": ("fire_alarm", "co"),
        "SMOKE/CO DETECTOR": ("fire_alarm", "combo"),
        "HORN/STROBE": ("fire_alarm", "horn_strobe"),
        "PULL STATION": ("fire_alarm", "pull_station"),
        "DUCT DETECTOR": ("fire_alarm", "duct_detector"),
        "MONITOR MODULE": ("fire_alarm", "module"),
        "FACP": ("fire_alarm", "facp"),
        "ANNUNCIATOR": ("fire_alarm", "annunciator"),
        "LIGHTING KEYPAD": ("controls", "keypad"),
        "CONTROL DIMMER": ("controls", "control_dimmer"),
        "SHADE CONTROL": ("controls", "shade"),
        "THERMOSTAT": ("controls", "thermostat"),
        "OCCUPANCY SENSOR": ("controls", "occupancy"),
        "VACANCY SENSOR": ("controls", "vacancy"),
        "LIGHTING CONTROL": ("lighting_control", "lighting_control"),
        "LANDSCAPE LIGHT": ("landscape_lighting", "landscape"),
        "RECESSED LIGHT": ("lighting", "lighting"),
        "DIMMER SWITCH": ("power", "power"),
    }
    for label, (category, device_type) in expected.items():
        found = classify_label(label)
        assert (found.category, found.device_type) == (category, device_type), label
    assert classify_label("SD", "fire_alarm").device_type == "smoke"
    assert classify_label("SD", "fire_alarm").category == "fire_alarm"
    assert classify_label("WIDGET", "low_voltage").category == "low_voltage"


def test_scope_defaults_and_job_overrides():
    job = {
        "include_landscape": False,
        "include_lighting_control": False,
    }
    assert scope_mode("power", job) == "ec"
    assert scope_mode("lighting", job) == "ec"
    assert scope_mode("landscape_lighting", job) == "by_others"
    assert scope_mode("lighting_control", job) == "by_others"
    assert scope_mode("low_voltage", job) == "by_others"
    assert scope_mode("fire_alarm", job) == "by_others"
    assert scope_mode("controls", job) == "by_others"
    assert scope_mode("landscape_lighting", {**job, "include_landscape": True}) == "ec"
    assert posted_scope("", checkbox_yes=True, default="by_others") == "ec"
    assert posted_scope("rough_in", checkbox_yes=True, default="by_others") == "rough_in"
    chosen = {**job, "scope_fire_alarm": "ec", "scope_low_voltage": "rough_in"}
    assert scope_mode("fire_alarm", chosen) == "ec"
    assert scope_mode("low_voltage", chosen) == "rough_in"


def _job():
    return {
        "id": "job",
        "include_landscape": False,
        "include_lighting_control": False,
        "scope_low_voltage": "rough_in",
        "rate_journeyman": "115",
        "rate_foreman": "130",
        "rate_pm": "150",
        "rate_apprentice": "65",
        "markup": None,
        "pdf_path": "",
    }


def _device(item_id, label, sheet, title):
    return {
        "id": item_id,
        "label": label,
        "category": classify_label(label).category,
        "measurement": "matched",
        "measured_quantity": 1,
        "sheets": [sheet],
        "panel": "",
        "circuit": "",
        "link_flag": "",
        "deduped": False,
        "occurrences": [
            {
                "page_index": 1,
                "sheet": sheet,
                "sheet_title": title,
                "x0": 10,
                "y0": 10,
                "x1": 20,
                "y1": 20,
                "grid": "",
            }
        ],
        "override_quantity": None,
        "material_unit_cost": None,
        "labor_hours": None,
        "labor_rate_key": "journeyman",
        "method": "legend_match",
        "match_score": 1,
        "review_status": "accepted",
        "exclusion": "",
    }


def test_rough_in_expands_box_ring_and_conduit_only():
    job = _job()
    row = present_item(_device("4", "DATA OUTLET", "E1", "FIRST FLOOR POWER"), job)
    assert row["scope"] == "rough-in only"
    assert row["effective"] == 1
    catalog = load_starter()
    apply_assemblies([row], job, catalog)
    roles = {part["role"] for part in row["assembly_parts"]}
    assert roles == {"box", "ring", "raceway"}
    assert any(part["description"] == "Sleeve, conduit, or pull string allowance" for part in row["assembly_parts"])
    materials = material_rows(quantity_lines([row]), catalog)
    assert {item["role"] for item in materials} == {"box", "ring", "raceway"}
    assert all(item["role"] != "device" for item in materials)


def test_fire_alarm_counts_are_by_device_type_and_floor():
    job = _job()
    job["scope_fire_alarm"] = "by_others"
    rows = [
        present_item(_device("1", "SMOKE DETECTOR", "E1", "FIRST FLOOR FIRE ALARM"), job),
        present_item(_device("2", "SMOKE DETECTOR", "E2", "SECOND FLOOR FIRE ALARM"), job),
        present_item(_device("3", "HORN/STROBE", "E1", "FIRST FLOOR FIRE ALARM"), job),
        present_item(_device("4", "DATA OUTLET", "E1", "FIRST FLOOR POWER"), job),
    ]
    floors = fire_alarm_by_floor(rows)
    smoke = [row for row in floors if row["device_type"] == "smoke"]
    assert {row["floor"]: row["quantity"] for row in smoke} == {
        "FIRST FLOOR": 1,
        "SECOND FLOOR": 1,
    }
    horn = next(row for row in floors if row["device_type"] == "horn_strobe")
    assert horn["floor"] == "FIRST FLOOR"
    assert horn["quantity"] == 1
    assert all(row["category"] == "fire_alarm" for row in floors)
    split = scope_split(rows)
    assert any(row["party"] == "FA vendor" and row["scope"] == "by others" for row in split)
    rough = [row for row in split if row["device_type"] == "data"]
    assert {row["scope"] for row in rough} == {"EC", "by others"}
    assert any("Rough-in only" in row["note"] for row in rough)
    assert any(row["note"] == "Device by others" for row in rough)


def _symbol(page, x, y, kind):
    if kind == "smoke":
        page.draw_circle(fitz.Point(x, y), 9, color=(0, 0, 0), width=1.2)
        page.draw_rect(fitz.Rect(x - 3, y - 3, x + 3, y + 3), color=(0, 0, 0), width=0.8)
    elif kind == "horn":
        page.draw_rect(fitz.Rect(x - 8, y - 6, x + 8, y + 6), color=(0, 0, 0), width=1)
        page.draw_line(fitz.Point(x + 8, y), fitz.Point(x + 16, y - 6), color=(0, 0, 0), width=0.8)
    elif kind == "data":
        draw_duplex(page, x, y)
    else:
        draw_switch(page, x, y)


def _systems_pdf() -> bytes:
    document = fitz.open()
    legend = document.new_page(width=612, height=792)
    _text(legend, 72, 48, "FIRE ALARM LEGEND", size=12)
    _symbol(legend, 90, 100, "smoke")
    _text(legend, 120, 104, "SMOKE DETECTOR")
    _symbol(legend, 90, 150, "horn")
    _text(legend, 120, 154, "HORN/STROBE")
    _text(legend, 72, 210, "LOW VOLTAGE LEGEND", size=12)
    _symbol(legend, 90, 260, "data")
    _text(legend, 120, 264, "DATA OUTLET")
    _text(legend, 72, 320, "CONTROLS LEGEND", size=12)
    _symbol(legend, 90, 370, "keypad")
    _text(legend, 120, 374, "LIGHTING KEYPAD")

    first = document.new_page(width=700, height=500)
    _text(first, 36, 28, "SHEET E1")
    _text(first, 160, 28, "TITLE FIRST FLOOR FIRE ALARM")
    _symbol(first, 180, 120, "smoke")
    _symbol(first, 300, 120, "horn")
    _symbol(first, 180, 220, "data")
    _symbol(first, 300, 220, "keypad")

    second = document.new_page(width=700, height=500)
    _text(second, 36, 28, "SHEET E2")
    _text(second, 160, 28, "TITLE SECOND FLOOR FIRE ALARM")
    _symbol(second, 180, 140, "smoke")

    payload = document.tobytes()
    document.close()
    return payload


def test_legend_headings_map_symbols_and_excel_splits_scope(tmp_path):
    result = run_takeoff(_systems_pdf())
    by_label = {}
    for item in result.items:
        if item.measurement != "matched":
            continue
        by_label.setdefault(item.label, []).append(item)
    assert by_label["SMOKE DETECTOR"][0].category == "fire_alarm"
    assert by_label["HORN/STROBE"][0].category == "fire_alarm"
    assert by_label["DATA OUTLET"][0].category == "low_voltage"
    assert by_label["LIGHTING KEYPAD"][0].category == "controls"
    assert len(by_label["SMOKE DETECTOR"]) == 2

    client = TestClient(create_app(tmp_path))
    uploaded = client.post(
        "/jobs",
        files={"pdf": ("systems.pdf", _systems_pdf(), "application/pdf")},
        follow_redirects=False,
    )
    job_id = uploaded.headers["location"].rstrip("/").split("/")[-1]
    page = client.get(f"/jobs/{job_id}").text
    assert "by others" in page
    assert 'id="scope_fire_alarm"' in page
    assert 'id="scope_low_voltage"' in page
    saved = _save(
        client,
        job_id,
        scope_fire_alarm="ec",
        scope_low_voltage="rough_in",
        scope_controls="by_others",
    )
    assert saved.status_code == 200
    review = client.get(f"/jobs/{job_id}/review")
    assert 'id="fire-alarm-floors"' in review.text
    assert "FIRST FLOOR" in review.text
    assert "SECOND FLOOR" in review.text
    assert "Smoke detector" in review.text
    assert "rough-in only" in review.text
    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    assert "By system" in workbook.sheetnames
    assert "Scope Split" in workbook.sheetnames
    assert "Routing" not in workbook.sheetnames
    systems = "\n".join(
        " ".join("" if cell is None else str(cell) for cell in row)
        for row in workbook["By system"].iter_rows(values_only=True)
    )
    assert "fire alarm" in systems
    assert "Smoke detector" in systems
    assert "FIRST FLOOR" in systems
    assert "SECOND FLOOR" in systems
    split = "\n".join(
        " ".join("" if cell is None else str(cell) for cell in row)
        for row in workbook["Scope Split"].iter_rows(values_only=True)
    )
    assert "FA vendor" not in split or "EC" in split
    assert "Rough-in only" in split
    assert "AV/low-voltage integrator" in split
    assert "lighting control vendor" in split
    loaded = build_review(
        {"id": job_id, "pdf_path": str(tmp_path / "pdfs" / f"{job_id}.pdf"), "include_landscape": False,
         "include_lighting_control": False, "scope_fire_alarm": "ec", "scope_low_voltage": "rough_in",
         "scope_controls": "by_others", "rate_journeyman": "115", "rate_foreman": "130",
         "rate_pm": "150", "rate_apprentice": "65", "markup": None},
        [],
        [],
        [],
        [],
    )
    assert loaded["fire_alarm_floors"] == []
