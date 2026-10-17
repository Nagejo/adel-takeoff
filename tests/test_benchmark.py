"""Accuracy benchmark: frozen synthetic truth, failing thresholds, and review export."""

import json
import shutil
import zipfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

from adel_takeoff.benchmark import (
    documents_from_rows,
    engine_items,
    include_in_truth,
    load_thresholds,
    load_truth,
    render_markdown,
    score_predictions,
    score_set,
    write_scorecard,
)
from adel_takeoff.main import create_app, ingest_pdf
from adel_takeoff.__main__ import main
from tests.drawings import build_sample_pdf

ROOT = Path(__file__).resolve().parents[1] / "benchmarks"
SCHEDULED = ROOT / "synthetic-scheduled"
TOTALS = ROOT / "synthetic-totals"


def _occurrence(sheet, x, y):
    return {
        "sheet": sheet,
        "page_index": 0,
        "x0": x - 8,
        "y0": y - 8,
        "x1": x + 8,
        "y1": y + 8,
    }


def _row(**overrides):
    row = {
        "label": "DUPLEX RECEPTACLE",
        "category": "power",
        "review_status": "accepted",
        "measurement": "matched",
        "method": "legend_match",
        "exclusion": "",
        "measured": 1,
        "panel": "",
        "circuit": "",
        "type_mark": "",
        "voltage": "",
        "mounting": "",
        "mounting_height": "",
        "occurrences": [_occurrence("E1", 150, 160)],
    }
    row.update(overrides)
    return row


def test_scheduled_set_meets_thresholds():
    report = score_set(SCHEDULED)
    thresholds = load_thresholds(SCHEDULED)
    assert report.passes(thresholds)
    assert report.precision == 1
    assert report.recall == 1
    assert report.max_abs_count_error == 0
    assert report.schedule_recall == 1
    assert report.missed == []
    assert report.extra == []
    assert report.schedule_mismatches == []
    assert report.matched == 12
    duplex = next(row for row in report.totals if row["tag"] == "DUPLEX RECEPTACLE")
    assert duplex["expected"] == 3
    assert duplex["engine"] == 3
    assert duplex["count_error"] == 0


def test_totals_only_set_meets_thresholds():
    report = score_set(TOTALS)
    assert report.location_scored is False
    assert report.passes(load_thresholds(TOTALS))
    assert report.precision == 1
    assert report.recall == 1
    assert report.max_abs_count_error == 0
    assert report.extra == []
    assert report.missed == []
    assert report.by_sheet == []
    assert "tag totals" in render_markdown(report)


def test_moved_truth_point_fails_recall():
    predicted, unmatched = engine_items((SCHEDULED / "drawing.pdf").read_bytes())
    truth = deepcopy(load_truth(SCHEDULED))
    truth["sheets"][0]["items"][0]["y"] = float(truth["sheets"][0]["items"][0]["y"]) + 200
    report = score_predictions(predicted, unmatched, truth, set_name="moved")
    assert report.recall < 1
    assert report.missed
    assert report.passes(load_thresholds(SCHEDULED)) is False


def test_wrong_total_fails_count_error():
    predicted, unmatched = engine_items((SCHEDULED / "drawing.pdf").read_bytes())
    truth = deepcopy(load_truth(SCHEDULED))
    truth["totals"][0]["quantity"] = int(truth["totals"][0]["quantity"]) + 5
    report = score_predictions(predicted, unmatched, truth, set_name="wrong-total")
    assert report.max_abs_count_error >= 5
    assert report.passes({"max_abs_count_error": 0, "precision": 1, "recall": 1}) is False


def test_scorer_is_deterministic():
    first = render_markdown(score_set(SCHEDULED))
    second = render_markdown(score_set(SCHEDULED))
    assert first == second


def test_external_folder_scores_without_living_in_the_repo(tmp_path):
    external = tmp_path / "client-drawings" / "riverside"
    shutil.copytree(SCHEDULED, external)
    assert ROOT not in external.parents
    report = score_set(external)
    assert report.passes(load_thresholds(external))
    markdown_path, excel_path = write_scorecard(report, external / "scorecard")
    assert markdown_path.is_file()
    workbook = openpyxl.load_workbook(excel_path)
    assert workbook.sheetnames == [
        "Summary",
        "By tag",
        "By sheet",
        "By category",
        "Totals",
        "Missed",
        "Extra",
        "Schedule",
    ]
    assert "Routing" not in workbook.sheetnames


def test_review_decisions_export_to_truth():
    rows = [
        _row(),
        _row(
            label="Unmatched mark",
            category="unmatched",
            measurement="unmatched",
            method="unmatched",
            measured=None,
            occurrences=[_occurrence("E1", 150, 519)],
        ),
        _row(
            label="SINGLE POLE SWITCH",
            review_status="rejected",
            occurrences=[_occurrence("E1", 150, 430)],
        ),
        _row(
            label="RECESSED LIGHT",
            category="lighting",
            review_status="corrected",
            measurement="unmatched",
            method="unmatched",
            occurrences=[_occurrence("E1", 150, 340)],
        ),
        _row(
            label="DATA OUTLET",
            category="low_voltage",
            measurement="unmatched",
            method="manual_add",
            occurrences=[_occurrence("E2", 80, 80)],
        ),
        _row(
            label="LANDSCAPE LIGHT",
            category="landscape_lighting",
            occurrences=[_occurrence("E1", 320, 160)],
        ),
        _row(
            label="HORN STROBE",
            category="fire_alarm",
            exclusion="by_others",
            occurrences=[_occurrence("E1", 40, 40)],
        ),
        _row(label="", review_status="pending"),
    ]
    assert include_in_truth(rows[0])
    assert include_in_truth(rows[1]) is False
    assert include_in_truth(rows[2]) is False
    assert include_in_truth(rows[3])
    assert include_in_truth(rows[4])
    assert include_in_truth(rows[5])
    assert include_in_truth(rows[6]) is False
    documents = documents_from_rows(rows)
    payload = json.dumps(documents)
    assert "DUPLEX RECEPTACLE" in payload
    assert "RECESSED LIGHT" in payload
    assert "DATA OUTLET" in payload
    assert "LANDSCAPE LIGHT" in payload
    assert "Unmatched" not in payload
    assert "SINGLE POLE SWITCH" not in payload
    assert "HORN STROBE" not in payload
    duplex = json.loads(
        json.dumps(documents["truth/totals/duplex-receptacle.json"])
    )
    assert duplex["quantity"] == 1
    assert "truth/sheets/e1.json" in documents
    assert "truth/sheets/e2.json" in documents


def test_truth_download_drops_rejected_and_unmatched(tmp_path):
    client = TestClient(create_app(tmp_path))
    uploaded = client.post(
        "/jobs",
        files={"pdf": ("set.pdf", build_sample_pdf(), "application/pdf")},
        follow_redirects=False,
    )
    job_id = uploaded.headers["location"].rstrip("/").split("/")[-1]
    page = client.get(f"/jobs/{job_id}/review")
    assert page.status_code == 200
    assert 'id="benchmark-truth"' in page.text
    assert 'id="benchmark-note"' in page.text
    downloaded = client.get(f"/jobs/{job_id}/truth.zip")
    assert downloaded.status_code == 200
    archive = zipfile.ZipFile(BytesIO(downloaded.content))
    assert "drawing.pdf" in archive.namelist()
    body = "\n".join(
        archive.read(name).decode() for name in archive.namelist() if name.endswith(".json")
    )
    assert "DUPLEX RECEPTACLE" in body
    assert "Unmatched" not in body
    import re

    match = re.search(
        r'id="item-(?P<id>[^"]+)" class="[^"]*" data-label="SINGLE POLE SWITCH"',
        page.text,
    )
    assert match
    rejected = client.post(
        f"/jobs/{job_id}/items/{match.group('id')}/review",
        data={"action": "reject", "label": "SINGLE POLE SWITCH", "panel": "", "circuit": ""},
        follow_redirects=False,
    )
    assert rejected.status_code == 303
    again = zipfile.ZipFile(BytesIO(client.get(f"/jobs/{job_id}/truth.zip").content))
    body = "\n".join(
        again.read(name).decode() for name in again.namelist() if name.endswith(".json")
    )
    assert "SINGLE POLE SWITCH" not in body
    assert "DUPLEX RECEPTACLE" in body


def test_cli_writes_scorecard_and_exits_when_thresholds_fail(tmp_path):
    out = tmp_path / "card"
    main(["benchmark", str(TOTALS), "--out", str(out)])
    assert (out / "scorecard.md").read_text(encoding="utf-8").startswith("# Benchmark")
    assert (out / "scorecard.xlsx").is_file()

    bad = tmp_path / "bad-set"
    shutil.copytree(TOTALS, bad)
    path = bad / "truth" / "totals" / "duplex-receptacle.json"
    record = json.loads(path.read_text(encoding="utf-8"))
    record["quantity"] = 99
    path.write_text(json.dumps(record), encoding="utf-8")
    try:
        main(["benchmark", str(bad), "--out", str(bad / "scorecard")])
    except SystemExit as exc:
        assert exc.code == 1
    else:
        raise AssertionError("a wrong total should fail the benchmark")
    assert (bad / "scorecard" / "scorecard.md").is_file()


def test_truth_command_writes_a_set_and_keeps_an_existing_pdf(tmp_path):
    job_id = ingest_pdf(tmp_path, "set.pdf", build_sample_pdf())
    out = tmp_path / "reviewed-set"
    main(["--data-dir", str(tmp_path), "truth", job_id, "--out", str(out)])
    assert (out / "drawing.pdf").is_file()
    assert (out / "truth" / "totals" / "duplex-receptacle.json").is_file()
    marker = b"%PDF-existing"
    (out / "drawing.pdf").write_bytes(marker)
    main(["--data-dir", str(tmp_path), "truth", job_id, "--out", str(out)])
    assert (out / "drawing.pdf").read_bytes() == marker
