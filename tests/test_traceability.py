"""Stable mark ids, sheet references, confidence bands, and review decisions."""

import json
import subprocess
import sys
from io import BytesIO
from pathlib import Path

import pytest

import fitz
import openpyxl
from fastapi.testclient import TestClient

from adel_takeoff.confidence import ConfidenceThresholds, confidence_band, combine_confidence
from adel_takeoff.constants import (
    CONFIDENCE_CHECK_MIN,
    CONFIDENCE_HIGH_MIN,
    EXACT_MATCH_SCORE,
)
from adel_takeoff.main import create_app
from adel_takeoff.takeoff import run_takeoff
from tests.drawings import build_sample_pdf
from tests.test_app import _marked, _rows, build_markup_text


ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def uploaded(tmp_path):
    client = TestClient(create_app(tmp_path))
    response = client.post(
        "/jobs",
        files={"pdf": ("set.pdf", build_sample_pdf(), "application/pdf")},
        follow_redirects=False,
    )
    assert response.status_code == 303
    job_id = response.headers["location"].rstrip("/").split("/")[-1]
    return client, job_id, tmp_path


def _header_map(sheet):
    for row in sheet.iter_rows():
        values = [cell.value for cell in row]
        if values and "How counted" in values:
            return {value: index for index, value in enumerate(values) if value}
    raise AssertionError("header row not found")


def _data_rows(sheet, headers):
    found = []
    header_seen = False
    for row in sheet.iter_rows(values_only=True):
        if not header_seen:
            if row and "How counted" in row:
                header_seen = True
            continue
        if not row or row[headers["Item"]] is None:
            if found:
                break
            continue
        found.append({name: row[index] for name, index in headers.items() if index < len(row)})
    return found


def test_confidence_bands_follow_config_thresholds():
    custom = ConfidenceThresholds(high_min=0.97, check_min=0.42)
    assert confidence_band(0.97, custom) == "high"
    assert confidence_band(0.96, custom) == "check"
    assert confidence_band(0.42, custom) == "check"
    assert confidence_band(0.41, custom) == "low"
    assert confidence_band(None, custom) == ""

    assert confidence_band(EXACT_MATCH_SCORE) == "high"
    assert confidence_band(CONFIDENCE_HIGH_MIN) == "high"
    assert confidence_band(CONFIDENCE_CHECK_MIN) == "check"
    between = (CONFIDENCE_HIGH_MIN + CONFIDENCE_CHECK_MIN) / 2
    assert CONFIDENCE_CHECK_MIN <= between < CONFIDENCE_HIGH_MIN
    assert confidence_band(between) == "check"
    assert confidence_band(CONFIDENCE_CHECK_MIN - 0.05) == "low"
    assert combine_confidence(["high", "check"]) == "check"
    assert combine_confidence(["high", ""]) == ""
    assert combine_confidence(["high", "high"]) == "high"


def test_item_ids_are_stable_and_scores_are_not_invented():
    first = run_takeoff(build_sample_pdf())
    second = run_takeoff(build_sample_pdf())
    assert [item.item_id for item in first.items] == [item.item_id for item in second.items]
    assert len({item.item_id for item in first.items}) == len(first.items)
    matched = [item for item in first.items if item.method == "legend_match"]
    unmatched = [item for item in first.items if item.method == "unmatched"]
    assert matched
    assert unmatched
    assert all(item.match_score == EXACT_MATCH_SCORE for item in matched)
    assert all(item.match_score is None for item in unmatched)
    assert all(item.item_id.isdigit() for item in first.items)


def test_marks_sheet_refs_and_confidence_export(uploaded):
    client, job_id, _tmp = uploaded
    review = client.get(f"/jobs/{job_id}/review")
    rows = _rows(review.text)
    duplex = next(row for row in rows if row["panel"] == "LP1" and row["circuit"] == "12")
    assert "grid B-4" in review.text
    assert "FIRST FLOOR POWER" in review.text
    assert "legend match" in review.text
    assert _marked(review.text, "power-device-count") == "4"

    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    assert "Excluded" in workbook.sheetnames
    headers = _header_map(workbook["Quantities"])
    for name in (
        "Line",
        "Item IDs",
        "Sheet number",
        "Sheet title",
        "Location",
        "How counted",
        "Confidence",
        "Detector score",
    ):
        assert name in headers
    quantities = _data_rows(workbook["Quantities"], headers)
    tagged = next(
        row for row in quantities if row["Panel"] == "LP1" and str(row["Circuit"]) == "12"
    )
    assert tagged["Item IDs"] == duplex["id"]
    assert tagged["Line"] == duplex["id"]
    assert tagged["Sheet number"] == "E1"
    assert tagged["Sheet title"] == "FIRST FLOOR POWER"
    assert "grid B-4" in tagged["Location"]
    assert "legend match" in tagged["How counted"]
    assert f"item {duplex['id']}" in tagged["How counted"]
    assert tagged["How counted"].rstrip().endswith("= 1")
    assert tagged["Confidence"] == "high"
    assert tagged["Detector score"] == 1

    overlap = next(row for row in quantities if "overlap" in str(row["How counted"]))
    assert "E1" in overlap["Sheet number"] and "E2" in overlap["Sheet number"]
    assert "FIRST FLOOR POWER" in overlap["Sheet title"]
    assert "FIRST FLOOR LIGHTING" in overlap["Sheet title"]
    assert overlap["Effective quantity"] == 1
    assert overlap["Confidence"] == "high"

    excluded_headers = _header_map(workbook["Excluded"])
    excluded = _data_rows(workbook["Excluded"], excluded_headers)
    reasons = {row["Reason"] for row in excluded}
    assert "by others" in reasons
    assert any("not counted" in str(row["Reason"]) for row in excluded)
    assert all(row["Confidence"] in (None, "", "high") for row in excluded)

    markup = build_markup_text(client.get(f"/jobs/{job_id}/export.pdf").content)
    assert f"{duplex['id']} measured" in markup
    assert "counted once" in markup
    assert "uncounted" in markup
    # The bare item id is also drawn inside the numbered mark.
    assert duplex["id"] in markup


def test_review_reject_and_accept_persist(uploaded):
    client, job_id, tmp_path = uploaded
    html = client.get(f"/jobs/{job_id}/review").text
    target = next(row for row in _rows(html) if row["panel"] == "LP1" and row["circuit"] == "12")
    rejected = client.post(
        f"/jobs/{job_id}/items/{target['id']}/review",
        data={"action": "reject", "label": "DUPLEX RECEPTACLE", "panel": "LP1", "circuit": "12"},
        follow_redirects=True,
    )
    assert rejected.status_code == 200
    assert _marked(rejected.text, "power-device-count") == "3"
    kept = next(row for row in _rows(rejected.text) if row["id"] == target["id"])
    assert kept["scope"] == "rejected"
    assert kept["effective"] == ""
    assert "rejected" in rejected.text

    document = json.loads((tmp_path / "reviews" / f"{job_id}.json").read_text())
    assert document["decisions"][target["id"]]["review_status"] == "rejected"

    again = TestClient(create_app(tmp_path))
    refreshed = again.get(f"/jobs/{job_id}/review")
    assert _marked(refreshed.text, "power-device-count") == "3"
    workbook = openpyxl.load_workbook(BytesIO(again.get(f"/jobs/{job_id}/export.xlsx").content))
    excluded = _data_rows(workbook["Excluded"], _header_map(workbook["Excluded"]))
    match = next(row for row in excluded if row["Item ID"] == target["id"])
    assert match["Reason"] == "rejected"
    assert "not in the counted total" in match["How counted"]
    quantities = _data_rows(workbook["Quantities"], _header_map(workbook["Quantities"]))
    assert all(row["Item IDs"] != target["id"] for row in quantities)
    markup = build_markup_text(again.get(f"/jobs/{job_id}/export.pdf").content)
    assert f"{target['id']} rejected" in markup

    accepted = again.post(
        f"/jobs/{job_id}/items/{target['id']}/review",
        data={"action": "accept", "label": "", "panel": "", "circuit": ""},
        follow_redirects=True,
    )
    assert _marked(accepted.text, "power-device-count") == "4"


def test_exclude_add_and_correct_flow_into_totals(uploaded):
    client, job_id, tmp_path = uploaded
    html = client.get(f"/jobs/{job_id}/review").text
    rows = _rows(html)
    switch = next(row for row in rows if row["label"] == "SINGLE POLE SWITCH")
    blank = next(row for row in rows if row["label"] == "DUPLEX RECEPTACLE" and row["panel"] == "")
    unmatched = next(row for row in rows if row["category"] == "unmatched")

    nic = client.post(
        f"/jobs/{job_id}/items/{switch['id']}/review",
        data={"action": "nic", "label": switch["label"], "panel": "", "circuit": ""},
        follow_redirects=True,
    )
    assert _marked(nic.text, "power-device-count") == "3"
    remain = client.post(
        f"/jobs/{job_id}/items/{blank['id']}/review",
        data={"action": "existing_to_remain", "label": blank["label"], "panel": "", "circuit": ""},
        follow_redirects=True,
    )
    assert _marked(remain.text, "power-device-count") == "2"
    by_others = client.post(
        f"/jobs/{job_id}/items/{unmatched['id']}/review",
        data={"action": "by_others", "label": "", "panel": "", "circuit": ""},
        follow_redirects=False,
    )
    assert by_others.status_code == 303

    added = client.post(
        f"/jobs/{job_id}/items",
        data={
            "page_index": "1",
            "x": "420",
            "y": "180",
            "label": "DUPLEX RECEPTACLE",
            "panel": "LP1",
            "circuit": "9",
        },
        follow_redirects=True,
    )
    assert added.status_code == 200
    assert _marked(added.text, "power-device-count") == "3"
    added_row = next(row for row in _rows(added.text) if row["panel"] == "LP1" and row["circuit"] == "9")
    assert "manual add" in added_row["flag"] or "manual add" in added.text
    assert added_row["measured"] == "1"

    missing = client.post(
        f"/jobs/{job_id}/items",
        data={"page_index": "1", "x": "10", "y": "10", "label": "   ", "panel": "", "circuit": ""},
        follow_redirects=False,
    )
    assert missing.status_code == 400

    outside = client.post(
        f"/jobs/{job_id}/items",
        data={
            "page_index": "1",
            "x": "5000",
            "y": "10",
            "label": "DUPLEX RECEPTACLE",
            "panel": "",
            "circuit": "",
        },
        follow_redirects=False,
    )
    assert outside.status_code == 400
    assert _marked(client.get(f"/jobs/{job_id}/review").text, "power-device-count") == "3"

    corrected = client.post(
        f"/jobs/{job_id}/items/{added_row['id']}/review",
        data={"action": "correct", "label": "RECESSED LIGHT", "panel": "LP2", "circuit": "8"},
        follow_redirects=True,
    )
    assert corrected.status_code == 200
    moved = next(row for row in _rows(corrected.text) if row["id"] == added_row["id"])
    assert moved["label"] == "RECESSED LIGHT"
    assert moved["category"] == "lighting"
    assert moved["panel"] == "LP2"
    assert moved["circuit"] == "8"
    assert _marked(corrected.text, "power-device-count") == "2"
    assert _marked(corrected.text, "lighting-fixture-count") == "3"

    empty_type = client.post(
        f"/jobs/{job_id}/items/{switch['id']}/review",
        data={"action": "correct", "label": "  ", "panel": "", "circuit": ""},
        follow_redirects=False,
    )
    assert empty_type.status_code == 400

    document = json.loads((tmp_path / "reviews" / f"{job_id}.json").read_text())
    assert document["decisions"][switch["id"]]["exclusion"] == "nic"
    assert document["decisions"][blank["id"]]["exclusion"] == "existing_to_remain"
    assert document["decisions"][unmatched["id"]]["exclusion"] == "by_others"
    assert document["added"][0]["id"] == added_row["id"]
    assert document["added"][0]["method"] == "manual_add"
    assert document["decisions"][added_row["id"]]["corrected_label"] == "RECESSED LIGHT"

    exported = TestClient(create_app(tmp_path)).get(f"/jobs/{job_id}/export.xlsx")
    workbook = openpyxl.load_workbook(BytesIO(exported.content))
    text = "\n".join(
        " ".join("" if value is None else str(value) for value in row)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows(values_only=True)
    )
    assert "Power devices: 2" in text
    assert "not in contract" in text
    assert "existing to remain" in text
    assert "by others" in text
    assert "corrected type" in text or "legend match" in text
    manual = _data_rows(workbook["Quantities"], _header_map(workbook["Quantities"]))
    light = next(row for row in manual if row["Item"] == "RECESSED LIGHT" and str(row["Circuit"]) == "8")
    assert light["Confidence"] in (None, "")
    assert light["Detector score"] in (None, "")
    assert "corrected type" in light["How counted"]
    assert light["Sheet number"] == "E1"
    assert light["Effective quantity"] == 1

    markup = build_markup_text(
        TestClient(create_app(tmp_path)).get(f"/jobs/{job_id}/export.pdf").content
    )
    assert f"{switch['id']} NIC" in markup
    assert f"{blank['id']} existing to remain" in markup
    assert f"{added_row['id']} type corrected" in markup


def test_correct_unmatched_mark_counts_one_without_a_score(uploaded):
    client, job_id, _tmp = uploaded
    html = client.get(f"/jobs/{job_id}/review").text
    unmatched = next(row for row in _rows(html) if row["category"] == "unmatched")
    corrected = client.post(
        f"/jobs/{job_id}/items/{unmatched['id']}/review",
        data={"action": "correct", "label": "DUPLEX RECEPTACLE", "panel": "LP1", "circuit": "1"},
        follow_redirects=True,
    )
    row = next(item for item in _rows(corrected.text) if item["id"] == unmatched["id"])
    assert row["label"] == "DUPLEX RECEPTACLE"
    assert row["category"] == "power"
    assert row["measured"] == "1"
    assert row["effective"] == "1"
    assert _marked(corrected.text, "power-device-count") == "5"
    assert "corrected type" in corrected.text
    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    quantities = _data_rows(workbook["Quantities"], _header_map(workbook["Quantities"]))
    line = next(item for item in quantities if item["Item IDs"] == unmatched["id"])
    assert line["Confidence"] in (None, "")
    assert line["Detector score"] in (None, "")
    assert "corrected type" in line["How counted"]
    assert line["How counted"].rstrip().endswith("= 1") or "total 1" in line["How counted"]


def test_cli_review_round_trip(tmp_path):
    pdf_path = tmp_path / "set.pdf"
    pdf_path.write_bytes(build_sample_pdf())
    data_dir = tmp_path / "data"
    job_id = _run(
        ["--data-dir", str(data_dir), "takeoff", str(pdf_path)]
    ).strip()
    assert (data_dir / "reviews" / f"{job_id}.json").is_file()
    _run(
        ["--data-dir", str(data_dir), "review", job_id, "--item", "1", "--action", "reject"]
    )
    out = tmp_path / "out"
    exported = _run(
        ["--data-dir", str(data_dir), "export", job_id, "--out", str(out)]
    )
    assert f"{job_id}-quantities.xlsx" in exported
    assert f"{job_id}-markup.pdf" in exported
    workbook = openpyxl.load_workbook(out / f"{job_id}-quantities.xlsx")
    excluded = _data_rows(workbook["Excluded"], _header_map(workbook["Excluded"]))
    assert any(row["Item ID"] == "1" and row["Reason"] == "rejected" for row in excluded)
    markup = build_markup_text((out / f"{job_id}-markup.pdf").read_bytes())
    assert "1 rejected" in markup
    document = fitz.open(out / f"{job_id}-markup.pdf")
    try:
        assert any(page.get_drawings() for page in document)
    finally:
        document.close()


def _run(args: list[str]) -> str:
    completed = subprocess.run(
        [sys.executable, "-m", "adel_takeoff", *args],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    return completed.stdout
