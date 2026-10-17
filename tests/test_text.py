"""Scanned sheets are read with OCR. Search covers the whole set."""

from io import BytesIO

import fitz
import openpyxl
from fastapi.testclient import TestClient

from adel_takeoff.extract import (
    Cluster,
    PageData,
    Word,
    legend_entries,
    pdf_text_words,
    read_pdf,
    usable_text_layer,
)
from adel_takeoff.main import create_app
from adel_takeoff.scale import regions_from_words
from adel_takeoff.schedules import parse_schedules
from adel_takeoff.text_index import build_text_index, render_hit, search_index
from tests.drawings import build_sample_pdf, build_scheduled_pdf


def _rasterize(payload: bytes, dpi: int = 200) -> bytes:
    source = fitz.open(stream=payload, filetype="pdf")
    output = fitz.open()
    matrix = fitz.Matrix(dpi / 72, dpi / 72)
    try:
        for page in source:
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            raster = output.new_page(width=page.rect.width, height=page.rect.height)
            raster.insert_image(raster.rect, pixmap=pixmap)
        return output.tobytes()
    finally:
        source.close()
        output.close()


def _scale_words(confidence: float) -> list[Word]:
    tokens = (
        (72, "SHEET"),
        (130, "E1"),
        (220, "TITLE"),
        (280, "FIRST"),
        (340, "FLOOR"),
        (410, "POWER"),
    )
    words = [
        Word(x, 28, x + 40, 40, text, source="pdf")
        for x, text in tokens
    ]
    scale = (
        (72, "SCALE:"),
        (140, '1/8"'),
        (180, "="),
        (200, "1'-0\""),
    )
    for x, text in scale:
        words.append(
            Word(x, 70, x + 36, 84, text, source="ocr", confidence=confidence)
        )
    return words


def test_vector_pages_keep_the_pdf_text_layer():
    pages = read_pdf(build_sample_pdf())
    assert pages
    assert all(page.text_source == "pdf" for page in pages)
    assert all(word.source == "pdf" and word.confidence is None for page in pages for word in page.words)
    assert usable_text_layer(pages[0].words)


def test_low_confidence_ocr_does_not_set_a_scale_or_schedule_value():
    regions = regions_from_words(_scale_words(0.4))
    assert regions[0]["status"] == "check"
    assert regions[0]["source"] == "ocr"
    assert regions[0]["feet_per_point"] == ""
    assert regions[0]["confidence"] == "check"
    assert "check" in regions[0]["reason"].lower()

    accepted = regions_from_words(_scale_words(0.95))
    assert accepted[0]["status"] == "valid"
    assert accepted[0]["source"] == "ocr"
    assert accepted[0]["label"] == '1/8" = 1\'-0"'

    low_page = _schedule_page(0.3)
    low = parse_schedules([low_page])
    assert low.fixture_rows == []
    assert low.ocr_checks
    assert low.ocr_checks[0]["kind"] == "schedule"
    assert "check" in low.ocr_checks[0]["note"].lower()

    high = parse_schedules([_schedule_page(0.96)])
    assert high.fixture_rows[0].mark == "A"
    assert high.fixture_rows[0].description == "DOWNLIGHT"
    assert high.fixture_rows[0].source == "ocr"
    assert high.ocr_checks == []


def test_low_confidence_ocr_does_not_name_a_legend_symbol():
    page = _legend_page(0.2)
    assert legend_entries(page) == []
    assert page.ocr_checks
    assert page.ocr_checks[0]["kind"] == "legend"
    assert "DUPLEX" in page.ocr_checks[0]["text"]

    named = _legend_page(0.97)
    entries = legend_entries(named)
    assert entries[0].label == "DUPLEX RECEPTACLE"
    assert entries[0].source == "ocr"
    assert named.ocr_checks == []


def test_rasterized_synthetic_set_is_indexed_and_searchable():
    vector = build_scheduled_pdf()
    raster = _rasterize(vector)
    document = fitz.open(stream=raster, filetype="pdf")
    try:
        assert pdf_text_words(document[0]) == []
        assert not usable_text_layer(pdf_text_words(document[1]))
    finally:
        document.close()

    index = build_text_index(raster)
    assert index["pages"]
    assert all(page["source"] == "ocr" for page in index["pages"])
    assert all(word["source"] == "ocr" and word["confidence"] is not None for page in index["pages"] for word in page["words"])
    e1 = next(page for page in index["pages"] if page["sheet"] == "E1")
    assert "FLOOR" in e1["sheet_title"]
    legend = index["pages"][0]
    legend_text = " ".join(word["text"] for word in legend["words"])
    assert "LEGEND" in legend_text
    assert "DUPLEX" in legend_text

    pages = read_pdf(raster)
    parsed = parse_schedules(pages)
    marks = {row.mark: row for row in parsed.fixture_rows}
    assert marks["A"].source == "ocr"
    assert "DOWNLIGHT" in marks["A"].description
    assert marks["A"].count == 3

    hits = search_index(index, "LP1")
    assert hits
    assert all(hit["source"] == "ocr" for hit in hits)
    assert any(hit["sheet"] in {"E1", "E3"} for hit in hits)
    assert all(hit["x1"] > hit["x0"] and hit["y1"] > hit["y0"] for hit in hits)
    phrase = search_index(index, "FIRST FLOOR")
    assert any("FIRST" in hit["text"] and "FLOOR" in hit["text"] for hit in phrase)

    png = render_hit(
        raster,
        hits[0]["page_index"],
        (hits[0]["x0"], hits[0]["y0"], hits[0]["x1"], hits[0]["y1"]),
    )
    assert png.startswith(b"\x89PNG")


def test_search_page_marks_the_hit_and_excel_lists_it(tmp_path):
    raster = _rasterize(build_scheduled_pdf())
    client = TestClient(create_app(tmp_path))
    uploaded = client.post(
        "/jobs",
        files={"pdf": ("scanned.pdf", raster, "application/pdf")},
        follow_redirects=False,
    )
    job_id = uploaded.headers["location"].rstrip("/").split("/")[-1]
    page = client.get(f"/jobs/{job_id}/search")
    assert page.status_code == 200
    assert "source ocr" in page.text
    assert ">ocr<" in page.text
    found = client.get(f"/jobs/{job_id}/search", params={"q": "LP1"})
    assert 'id="text-hits"' in found.text
    assert "Show on sheet" in found.text
    spot = client.get(f"/jobs/{job_id}/search", params={"q": "LP1", "hit": 0})
    assert 'id="text-spot"' in spot.text
    image = client.get(spot.text.split('src="')[1].split('"')[0].replace("&amp;", "&"))
    assert image.status_code == 200
    assert image.content.startswith(b"\x89PNG")
    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    assert "Text search" in workbook.sheetnames
    assert "OCR checks" in workbook.sheetnames
    assert "Routing" not in workbook.sheetnames
    text = "\n".join(
        " ".join("" if cell is None else str(cell) for cell in row)
        for row in workbook["Text search"].iter_rows(values_only=True)
    )
    assert "LP1" in text
    assert "ocr" in text


def _schedule_page(confidence: float) -> PageData:
    rows = [
        (40, [("72", "LIGHTING"), ("160", "FIXTURE"), ("250", "SCHEDULE")]),
        (70, [("72", "TYPE"), ("200", "DESCRIPTION"), ("400", "LAMP"), ("520", "VOLTAGE")]),
        (92, [("72", "A"), ("200", "DOWNLIGHT"), ("400", "LED"), ("520", "120")]),
    ]
    words = []
    for y, tokens in rows:
        for x_text, text in tokens:
            x = float(x_text)
            words.append(
                Word(
                    x,
                    y,
                    x + 8 * len(text),
                    y + 10,
                    text,
                    source="ocr",
                    confidence=confidence,
                )
            )
    return PageData(
        index=0,
        width=800,
        height=400,
        words=words,
        sheet="E3",
        sheet_title="FIXTURE AND PANEL SCHEDULES",
        text_source="ocr",
    )


def _legend_page(confidence: float) -> PageData:
    words = [
        Word(70, 36, 140, 48, "SYMBOL", source="ocr", confidence=confidence),
        Word(150, 36, 220, 48, "LEGEND", source="ocr", confidence=confidence),
        Word(120, 96, 180, 108, "DUPLEX", source="ocr", confidence=confidence),
        Word(186, 96, 280, 108, "RECEPTACLE", source="ocr", confidence=confidence),
    ]
    cluster = Cluster(
        page_index=0,
        primitives=[],
        bbox=(80, 90, 100, 110),
        fingerprint=(("duplex",),),
        center=(90, 100),
    )
    return PageData(
        index=0,
        width=612,
        height=400,
        words=words,
        clusters=[cluster],
        sheet="E0",
        text_source="ocr",
    )
