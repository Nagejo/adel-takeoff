"""Upload, review, corrections, and exports against a synthetic drawing set."""

import re
from io import BytesIO

import fitz
import openpyxl
import pytest
from fastapi.testclient import TestClient

from adel_takeoff.constants import DISCLAIMER, JOURNEYMAN_WARNING
from adel_takeoff.main import create_app
from tests.drawings import build_sample_pdf

STREET_RE = re.compile(
    r"\d{1,6}\s+[A-Za-z0-9.'-]+(?:\s+[A-Za-z0-9.'-]+){0,3}\s+"
    r"(?:Street|St\.|Avenue|Ave\.|Road|Rd\.|Boulevard|Blvd\.|Drive|Dr\.|Lane|Ln\.)",
    re.IGNORECASE,
)
ROW_RE = re.compile(
    r'<tr id="item-(?P<id>[^"]+)" class="(?P<class>[^"]*)" '
    r'data-label="(?P<label>[^"]*)" data-category="(?P<category>[^"]*)">'
    r'(?P<body>.*?)</tr>',
    re.DOTALL,
)


def _cell(body: str, name: str) -> str:
    match = re.search(rf'class="{name}">(.*?)</td>', body, re.DOTALL)
    assert match, name
    text = re.sub(r"<[^>]+>", " ", match.group(1))
    return " ".join(text.split())


def _rows(html: str) -> list[dict]:
    found = []
    for match in ROW_RE.finditer(html):
        body = match.group("body")
        found.append(
            {
                "id": match.group("id"),
                "class": match.group("class"),
                "label": match.group("label"),
                "category": match.group("category"),
                "measured": _cell(body, "measured"),
                "sheets": _cell(body, "sheets"),
                "panel": _cell(body, "panel"),
                "circuit": _cell(body, "circuit"),
                "flag": _cell(body, "link-flag"),
                "scope": _cell(body, "scope"),
                "effective": _cell(body, "effective"),
            }
        )
    return found


def _marked(html: str, element_id: str) -> str:
    match = re.search(rf'id="{element_id}"[^>]*>([^<]*)', html)
    assert match, element_id
    return match.group(1).strip()


@pytest.fixture
def uploaded(tmp_path):
    client = TestClient(create_app(tmp_path))
    response = client.post(
        "/jobs",
        files={"pdf": ("set.pdf", build_sample_pdf(), "application/pdf")},
        follow_redirects=False,
    )
    assert response.status_code == 303
    job_id = response.headers["location"].rstrip("/").split("/")[-1]
    return client, job_id, tmp_path


def _save(client, job_id, **extra):
    form = {
        "permit_issue_date": "",
        "markup": "",
        "rate_journeyman": "115",
        "rate_foreman": "130",
        "rate_pm": "150",
        "rate_apprentice": "65",
    }
    form.update(extra)
    return client.post(f"/jobs/{job_id}", data=form, follow_redirects=True)


def test_legend_match_overlap_links_and_unmatched(uploaded):
    client, job_id, _tmp = uploaded
    review = client.get(f"/jobs/{job_id}/review")
    assert review.status_code == 200
    html = review.text
    rows = _rows(html)
    duplexes = [row for row in rows if row["label"] == "DUPLEX RECEPTACLE"]
    lights = [row for row in rows if row["label"] == "RECESSED LIGHT"]
    switches = [row for row in rows if row["label"] == "SINGLE POLE SWITCH"]
    unmatched = [row for row in rows if row["category"] == "unmatched"]
    panels = [row for row in rows if row["label"] == "PANELBOARD"]
    assert len(duplexes) == 3
    assert len(lights) == 2
    assert len(switches) == 1
    assert len(unmatched) == 1
    assert unmatched[0]["measured"] == ""
    assert "not counted" in unmatched[0]["flag"]
    assert _marked(html, "power-device-count") == "4"
    assert _marked(html, "lighting-fixture-count") == "2"
    deduped = [row for row in duplexes if "deduped" in row["class"]]
    assert len(deduped) == 1
    assert "E1" in deduped[0]["sheets"] and "E2" in deduped[0]["sheets"]
    assert "counted once" in deduped[0]["sheets"]
    assert deduped[0]["measured"] == "1"
    tagged = [row for row in duplexes if row["panel"] == "LP1" and row["circuit"] == "12"]
    assert len(tagged) == 1
    blank = [row for row in duplexes if row["panel"] == "" and row["circuit"] == ""]
    assert len(blank) == 1
    assert "not readable" in blank[0]["flag"]
    recessed = [row for row in lights if row["panel"] == "LP2" and row["circuit"] == "4"]
    assert len(recessed) == 1
    other_light = [row for row in lights if row["panel"] == "HP1" and row["circuit"] == "3"]
    assert len(other_light) == 1
    assert switches[0]["panel"] == "" and "not readable" in switches[0]["flag"]
    assert len(panels) == 1
    assert panels[0]["scope"] == "not counted"
    assert "unmatched" in html


def test_landscape_and_lighting_control_default_to_by_others(uploaded):
    client, job_id, _tmp = uploaded
    html = client.get(f"/jobs/{job_id}/review").text
    rows = _rows(html)
    landscape = [row for row in rows if row["category"] == "landscape_lighting"]
    control = [row for row in rows if row["category"] == "lighting_control"]
    assert len(landscape) == 1 and landscape[0]["scope"] == "by others"
    assert landscape[0]["measured"] == "1"
    assert len(control) == 1 and control[0]["scope"] == "by others"
    assert _marked(html, "lighting-fixture-count") == "2"
    assert _marked(html, "lighting-control-count") == "0"
    job = client.get(f"/jobs/{job_id}").text
    assert "by others" in job
    saved = _save(
        client,
        job_id,
        include_landscape="yes",
        include_lighting_control="yes",
    )
    assert saved.status_code == 200
    html = client.get(f"/jobs/{job_id}/review").text
    assert _marked(html, "lighting-fixture-count") == "3"
    assert _marked(html, "lighting-control-count") == "1"


def test_correction_and_allowance_persist_and_export(uploaded):
    client, job_id, tmp_path = uploaded
    html = client.get(f"/jobs/{job_id}/review").text
    target = next(row for row in _rows(html) if row["panel"] == "LP1" and row["circuit"] == "12")
    corrected = client.post(
        f"/jobs/{job_id}/items/{target['id']}/correction",
        data={
            "override_quantity": "2",
            "material_unit_cost": "10",
            "labor_hours": "2",
            "labor_rate_key": "journeyman",
        },
        follow_redirects=True,
    )
    assert corrected.status_code == 200
    assert _marked(corrected.text, "power-device-count") == "5"
    assert _marked(corrected.text, "direct-cost") == "250"
    assert _marked(corrected.text, "marked-up-total") == "not applied"
    allowance = client.post(
        f"/jobs/{job_id}/allowances",
        data={
            "description": "Spare duplex",
            "quantity": "4",
            "reason": "spare devices requested by the owner",
            "category": "allowance",
        },
        follow_redirects=True,
    )
    assert "spare devices requested by the owner" in allowance.text
    assert _marked(allowance.text, "power-device-count") == "5"

    again = TestClient(create_app(tmp_path))
    refreshed = again.get(f"/jobs/{job_id}/review")
    assert refreshed.status_code == 200
    kept = next(row for row in _rows(refreshed.text) if row["id"] == target["id"])
    assert kept["measured"] == "1"
    assert kept["effective"] == "2"
    assert "spare devices requested by the owner" in refreshed.text

    workbook = openpyxl.load_workbook(BytesIO(again.get(f"/jobs/{job_id}/export.xlsx").content))
    quantities = _sheet_text(workbook["Quantities"])
    assert "Spare duplex" in quantities
    assert "spare devices requested by the owner" in quantities
    assert "Power devices: 5" in quantities
    markup = build_markup_text(again.get(f"/jobs/{job_id}/export.pdf").content)
    assert "corrected count 2; measured 1" in markup
    assert "counted once" in markup
    assert "uncounted" in markup


def test_markup_and_labor_rates(uploaded):
    client, job_id, _tmp = uploaded
    page = client.get(f"/jobs/{job_id}").text
    assert 'id="markup"' in page
    assert 'value=""' in page or "blank means do not apply" in page
    assert "No markup is applied until you enter one." in page
    assert JOURNEYMAN_WARNING in page
    for label in ("Journeyman", "Foreman", "PM or master", "Apprentice"):
        assert label in page
    assert 'id="rate_journeyman"' in page
    assert ">115<" in page or 'value="115"' in page
    assert 'value="130"' in page
    assert 'value="150"' in page
    assert 'value="65"' in page

    rejected = client.post(
        f"/jobs/{job_id}",
        data={
            "permit_issue_date": "",
            "markup": "6",
            "rate_journeyman": "115",
            "rate_foreman": "130",
            "rate_pm": "150",
            "rate_apprentice": "65",
        },
    )
    assert rejected.status_code == 400
    assert "Markup greater than 5 is rejected." in rejected.text
    still_blank = client.get(f"/jobs/{job_id}").text
    assert 'id="markup" name="markup" type="text" inputmode="decimal" value=""' in still_blank

    accepted = _save(client, job_id, markup="0.35")
    assert "Markup greater than 5" not in accepted.text
    assert 'value="0.35"' in accepted.text
    review = client.get(f"/jobs/{job_id}/review").text
    assert "Markup: 0.35" in review

    over = client.post(
        f"/jobs/{job_id}",
        data={
            "permit_issue_date": "",
            "markup": "5.01",
            "rate_journeyman": "115",
            "rate_foreman": "130",
            "rate_pm": "150",
            "rate_apprentice": "65",
        },
    )
    assert over.status_code == 400
    kept = client.get(f"/jobs/{job_id}").text
    assert 'value="0.35"' in kept

    cleared = _save(client, job_id, markup="")
    assert 'value=""' in cleared.text or "Markup: not applied" in client.get(f"/jobs/{job_id}/review").text
    assert _marked(client.get(f"/jobs/{job_id}/review").text, "marked-up-total") == "not applied"

    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    text = _sheet_text(workbook["Quantities"])
    assert JOURNEYMAN_WARNING in text
    assert "Journeyman" in text and "Foreman" in text
    assert "PM or master" in text and "Apprentice" in text
    assert "Markup: not applied" in text

    _save(client, job_id, rate_journeyman="120")
    warned = client.get(f"/jobs/{job_id}").text
    assert JOURNEYMAN_WARNING not in warned
    exported = _sheet_text(
        openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))["Quantities"]
    )
    assert JOURNEYMAN_WARNING not in exported
    assert "120" in exported


def test_code_edition_rfi_and_claims(uploaded):
    client, job_id, _tmp = uploaded
    _save(client, job_id, permit_issue_date="2026-12-31")
    job = client.get(f"/jobs/{job_id}").text
    assert "FBC 8th Edition" in job
    assert "NEC 2020" in job
    review = client.get(f"/jobs/{job_id}/review").text
    assert "Drawing text says NEC 2023" in review
    assert 'data-disposition="rfi"' in review
    assert 'data-disposition="violation"' not in review
    assert "725.136" in review and "800.133" in review
    assert 'data-section="unknown"' in review
    assert "NEC 300.11" in review
    assert "violation" not in review.lower()

    _save(client, job_id, permit_issue_date="2027-01-01")
    later = client.get(f"/jobs/{job_id}").text
    assert "FBC 9th Edition" in later
    assert "NEC 2023" in later
    assert "permit-issue date" in later
    still = client.get(f"/jobs/{job_id}/review").text
    assert "Drawing text says NEC 2023" in still
    assert "It is not recorded as the applicable code edition." in still


def test_exports_identity_and_files(uploaded):
    client, job_id, _tmp = uploaded
    excel = client.get(f"/jobs/{job_id}/export.xlsx")
    pdf = client.get(f"/jobs/{job_id}/export.pdf")
    assert excel.status_code == 200
    assert pdf.status_code == 200
    assert excel.content[:2] == b"PK"
    assert pdf.content[:4] == b"%PDF"
    workbook = openpyxl.load_workbook(BytesIO(excel.content))
    assert "Routing" not in workbook.sheetnames
    excel_text = "\n".join(_sheet_text(sheet) for sheet in workbook.worksheets)
    pdf_text = build_markup_text(pdf.content)
    for blob in (excel_text, pdf_text):
        assert "Adel Electric, LLC" in blob
        assert DISCLAIMER in blob
        assert "Home Entertainment Design" not in blob
        assert "insurance" not in blob.lower()
        assert STREET_RE.search(blob) is None
    assert "counted once" in pdf_text
    assert "measured" in pdf_text


def _sheet_text(sheet) -> str:
    lines = []
    for row in sheet.iter_rows(values_only=True):
        lines.append(" ".join("" if value is None else str(value) for value in row))
    return "\n".join(lines)


def build_markup_text(payload: bytes) -> str:
    document = fitz.open(stream=payload, filetype="pdf")
    try:
        return "\n".join(page.get_text("text") for page in document)
    finally:
        document.close()
