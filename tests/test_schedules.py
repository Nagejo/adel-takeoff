"""Fixture and panel schedules stay tied to the cells the drawing printed."""

from io import BytesIO

import fitz
import openpyxl
import pytest
from fastapi.testclient import TestClient

from adel_takeoff.confidence import apply_discrepancy_cap
from adel_takeoff.extract import PageData, Word, read_pdf
from adel_takeoff.main import create_app
from adel_takeoff.schedule_link import attach_schedules, floor_from_title, resolve_height
from adel_takeoff.schedules import normalize_circuit, parse_schedules, whole_number
from tests.drawings import build_sample_pdf, build_scheduled_pdf
from tests.test_app import _marked, _rows
from tests.test_traceability import _data_rows, _header_map


def test_whole_number_count_does_not_parse_a_unit():
    assert whole_number("3") == 3
    assert whole_number("0") == 0
    assert whole_number("2 EA") is None
    assert whole_number("") is None
    assert normalize_circuit("012") == "12"
    assert normalize_circuit("12") == "12"


def test_discrepancy_cap_does_not_invent_or_raise_a_flag():
    assert apply_discrepancy_cap("high") == "check"
    assert apply_discrepancy_cap("check") == "check"
    assert apply_discrepancy_cap("low") == "low"
    assert apply_discrepancy_cap("") == ""
    assert floor_from_title("FIRST FLOOR POWER") == "FIRST FLOOR"
    assert floor_from_title("LEVEL 2") == "LEVEL 2"
    assert floor_from_title("FIXTURE AND PANEL SCHEDULES") == ""
    assert resolve_height("9'-0\"", ["9'-0\""]) == ("9'-0\"", False)
    assert resolve_height("9'-0\"", ["7'-0\""]) == ('schedule: 9\'-0"; tag: 7\'-0"', True)
    assert resolve_height("", ["6'-0\""]) == ("6'-0\"", False)
    assert resolve_height("", []) == ("", False)


def _place(words: list[Word], x: float, y: float, text: str) -> None:
    cursor = x
    for token in text.split():
        width = max(8.0, len(token) * 4.5)
        words.append(Word(cursor, y, cursor + width, y + 8, token))
        cursor += width + 3


def _page(lines: list[tuple[float, list[tuple[float, str]]]], sheet: str, index: int, title: str = "") -> PageData:
    words: list[Word] = []
    for y, cells in lines:
        for x, text in cells:
            _place(words, x, y, text)
    return PageData(
        index=index,
        width=1100,
        height=800,
        words=words,
        sheet=sheet,
        sheet_title=title,
        text="",
    )


def _device(item_id: str, page: PageData, x: float, y: float, panel: str = "", circuit: str = "", confidence: str = "high") -> dict:
    return {
        "id": item_id,
        "label": "RECESSED LIGHT",
        "category": "lighting",
        "panel": panel,
        "circuit": circuit,
        "confidence": confidence,
        "sheet_numbers": [page.sheet],
        "occurrences": [
            {
                "page_index": page.index,
                "sheet": page.sheet,
                "sheet_title": page.sheet_title,
                "x0": x - 8,
                "y0": y - 8,
                "x1": x + 8,
                "y1": y + 8,
            }
        ],
    }


def _synthetic_pages() -> list[PageData]:
    schedule = _page(
        [
            (40, [(72, "LIGHTING FIXTURE SCHEDULE")]),
            (
                60,
                [
                    (72, "TYPE"),
                    (160, "DESCRIPTION"),
                    (420, "MANUFACTURER"),
                    (560, "LAMP"),
                    (640, "VOLTAGE"),
                    (740, "MOUNTING"),
                    (880, "HEIGHT"),
                    (1000, "COUNT"),
                ],
            ),
            (
                80,
                [
                    (72, "A"),
                    (160, "ONE"),
                    (420, "ACME"),
                    (560, "LED"),
                    (640, "120"),
                    (740, "WALL"),
                    (880, "9'-0\""),
                    (1000, "2 EA"),
                ],
            ),
            (
                100,
                [
                    (72, "A"),
                    (160, "TWO"),
                    (420, "ACME"),
                    (560, "LED"),
                    (640, "277"),
                    (740, "PENDANT"),
                    (880, "8'-0\""),
                    (1000, "4"),
                ],
            ),
            (
                120,
                [
                    (72, "B"),
                    (160, "SCONCE"),
                    (420, "ACME"),
                    (560, "LED"),
                    (640, "120"),
                    (740, "WALL"),
                    (880, "6'-0\""),
                    (1000, "1"),
                ],
            ),
            (
                140,
                [
                    (72, "D"),
                    (160, "EMPTY"),
                    (420, "ACME"),
                    (560, "LED"),
                    (640, "120"),
                    (740, "PENDANT"),
                    (880, "10'-0\""),
                    (1000, "0"),
                ],
            ),
            (300, [(72, "PANEL SCHEDULE")]),
            (320, [(72, "CIRCUIT"), (180, "DESCRIPTION"), (460, "BREAKER")]),
            (340, [(72, "1"), (180, "SHOULD NOT READ"), (460, "20A")]),
            (420, [(72, "PANEL SCHEDULE LP1")]),
            (448, [(72, "CIRCUIT"), (180, "DESCRIPTION"), (460, "BREAKER"), (580, "POLES"), (680, "LOAD")]),
            (468, [(72, "012"), (180, "DUPLEX RECEPTACLE"), (460, "20A"), (580, "1"), (680, "180")]),
        ],
        sheet="E9",
        index=0,
    )
    plan = _page(
        [
            (80, [(130, "TYPE A TYPE B")]),
            (160, [(130, "B")]),
            (240, [(130, "Q")]),
            (320, [(130, "GRID B-4")]),
            (400, [(130, "LP1-12")]),
        ],
        sheet="E1",
        index=1,
        title="FIRST FLOOR POWER",
    )
    return [schedule, plan]


def test_duplicate_and_non_integer_rows_are_not_chosen():
    pages = _synthetic_pages()
    parsed = parse_schedules(pages)
    assert parsed.fixture_table_count == 1
    assert parsed.panel_table_count == 1
    assert "A" in parsed.duplicate_marks
    assert "A" not in parsed.fixture_by_mark
    assert parsed.fixture_by_mark["B"].description == "SCONCE"
    assert parsed.fixture_by_mark["B"].count == 1
    duplicate = [row for row in parsed.fixture_rows if row.mark == "A"]
    assert [row.count_text for row in duplicate] == ["2 EA", "4"]
    assert duplicate[0].count is None
    assert duplicate[1].count == 4
    assert parsed.fixture_by_mark["D"].count == 0
    assert all(row.description != "SHOULD NOT READ" for row in parsed.panel_rows)
    assert parsed.panel_by_key[("LP1", "12")].description == "DUPLEX RECEPTACLE"

    plan = pages[1]
    rows = [
        _device("1", plan, 100, 80),
        _device("2", plan, 100, 160),
        _device("3", plan, 100, 240),
        _device("4", plan, 100, 320),
        _device("5", plan, 100, 400, panel="LP1", circuit="12"),
        _device("6", plan, 100, 400, panel="LP8", circuit="1"),
    ]
    linked = attach_schedules(rows, pages)
    by_id = {row["id"]: row for row in rows}
    assert by_id["1"]["type_mark"] == ""
    assert by_id["1"]["type_description"] == ""
    assert by_id["1"]["schedule_type_conflict"] is True
    assert by_id["2"]["type_mark"] == "B"
    assert by_id["2"]["type_description"] == "SCONCE"
    assert by_id["2"]["voltage"] == "120"
    assert by_id["2"]["mounting"] == "WALL"
    assert by_id["3"]["type_mark"] == ""
    assert by_id["4"]["type_mark"] == ""
    assert by_id["5"]["circuit_description"] == "DUPLEX RECEPTACLE"
    assert by_id["5"]["circuit"] == "12"
    assert by_id["6"]["circuit_description"] == ""

    kinds = {item["kind"] for item in linked["discrepancies"]}
    assert "type tag conflicts; left blank" in kinds
    assert "duplicate schedule rows; neither used" in kinds
    assert "panel schedule not found" in kinds
    assert "schedule row with no plan marks" not in {
        item["kind"]
        for item in linked["discrepancies"]
        if "type D" in item["detail"] or "type B" in item["detail"]
    }
    duplicate_note = next(
        item for item in linked["discrepancies"] if item["kind"] == "duplicate schedule rows; neither used"
    )
    assert duplicate_note["plan"] == 0
    assert "2 EA" in str(duplicate_note["schedule"])
    assert "4" in str(duplicate_note["schedule"])
    assert "2 EA" in duplicate_note["detail"]
    assert "4" in duplicate_note["detail"]
    panel_note = next(item for item in linked["discrepancies"] if item["kind"] == "panel schedule not found")
    assert panel_note["item_ids"] == ["6"]
    assert all(item["kind"] != "circuit on plan missing from panel schedule" for item in linked["discrepancies"])


def test_no_panel_schedule_does_not_flag_every_circuit():
    page = _page([(80, [(130, "LP1-12")])], sheet="E1", index=0, title="FIRST FLOOR POWER")
    rows = [_device("1", page, 100, 80, panel="LP1", circuit="12")]
    linked = attach_schedules(rows, [page])
    assert linked["discrepancies"] == []
    assert linked["panel_table_count"] == 0
    assert rows[0]["circuit_description"] == ""


def _upload(tmp_path, payload: bytes):
    client = TestClient(create_app(tmp_path))
    response = client.post(
        "/jobs",
        files={"pdf": ("set.pdf", payload, "application/pdf")},
        follow_redirects=False,
    )
    assert response.status_code == 303
    job_id = response.headers["location"].rstrip("/").split("/")[-1]
    return client, job_id


def _kind_rows(sheet):
    headers = {cell.value: cell.column - 1 for cell in sheet[2] if cell.value}
    found = []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        if not row or row[headers["Kind"]] is None:
            if found:
                break
            continue
        found.append({name: row[index] for name, index in headers.items() if index < len(row)})
    return found


def _type_rows(sheet):
    headers = {cell.value: cell.column - 1 for cell in sheet[3] if cell.value}
    found = []
    for row in sheet.iter_rows(min_row=4, values_only=True):
        if not row or row[headers["Type"]] is None:
            if found:
                break
            continue
        found.append({name: row[index] for name, index in headers.items() if index < len(row)})
    return found


def test_sample_without_schedules_stays_unflagged(tmp_path):
    client, job_id = _upload(tmp_path, build_sample_pdf())
    review = client.get(f"/jobs/{job_id}/review").text
    assert "No fixture schedule detected." in review
    assert "No panel schedule detected." in review
    assert 'data-kind=' not in review
    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    assert _kind_rows(workbook["Discrepancies"]) == []
    headers = _header_map(workbook["Quantities"])
    quantities = _data_rows(workbook["Quantities"], headers)
    tagged = next(row for row in quantities if row["Panel"] == "LP1" and str(row["Circuit"]) == "12")
    assert tagged["Confidence"] == "high"
    assert tagged["Type description"] is None


def test_scheduled_drawing_links_rows_and_keeps_both_counts(tmp_path):
    client, job_id = _upload(tmp_path, build_scheduled_pdf())
    review = client.get(f"/jobs/{job_id}/review")
    assert review.status_code == 200
    html = review.text
    assert 'id="fixture-schedule"' in html
    assert 'id="panel-schedule"' in html
    assert "6 INCH DOWNLIGHT" in html
    assert "WALL SCONCE" in html
    assert "PENDANT" in html
    assert "DUPLEX RECEPTACLE" in html
    assert "SPARE" in html
    assert _marked(html, "lighting-fixture-count") == "4"
    assert _marked(html, "power-device-count") == "4"

    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    assert "By type" in workbook.sheetnames
    assert "Discrepancies" in workbook.sheetnames
    assert "Routing" not in workbook.sheetnames
    headers = _header_map(workbook["Quantities"])
    for name in (
        "Type",
        "Type description",
        "Mounting",
        "Mounting height",
        "Voltage",
        "Circuit description",
        "Breaker",
        "Poles",
        "Load",
    ):
        assert name in headers
    quantities = _data_rows(workbook["Quantities"], headers)

    def line(**wanted):
        for row in quantities:
            if all(row.get(key) == value for key, value in wanted.items()):
                return row
        raise AssertionError(wanted)

    type_b = line(**{"Type": "B"})
    assert type_b["Item"] == "RECESSED LIGHT"
    assert type_b["Type description"] == "WALL SCONCE"
    assert type_b["Mounting"] == "WALL"
    assert type_b["Voltage"] == "120"
    assert "6'-0" in str(type_b["Mounting height"])
    assert "schedule:" not in str(type_b["Mounting height"])
    assert type_b["Effective quantity"] == 1
    assert type_b["Confidence"] == "high"

    type_z = line(**{"Type": "Z"})
    assert type_z["Type description"] is None
    assert type_z["Voltage"] is None
    assert type_z["Effective quantity"] == 1
    assert type_z["Confidence"] == "check"

    type_a_lp2 = line(**{"Type": "A", "Panel": "LP2"})
    assert type_a_lp2["Type description"] == "6 INCH DOWNLIGHT"
    assert type_a_lp2["Mounting"] == "RECESSED"
    assert type_a_lp2["Voltage"] == "120"
    assert "9'-0" in str(type_a_lp2["Mounting height"])
    assert "schedule:" not in str(type_a_lp2["Mounting height"])
    assert type_a_lp2["Circuit description"] == "RECESSED FIXTURE"
    assert type_a_lp2["Breaker"] == "20A"
    assert str(type_a_lp2["Poles"]) == "1"
    assert str(type_a_lp2["Load"]) == "100"
    assert type_a_lp2["Effective quantity"] == 1
    assert type_a_lp2["Confidence"] == "check"

    type_a_hp1 = line(**{"Type": "A", "Panel": "HP1"})
    assert "schedule:" in str(type_a_hp1["Mounting height"])
    assert "tag:" in str(type_a_hp1["Mounting height"])
    assert "9'-0" in str(type_a_hp1["Mounting height"])
    assert "7'-0" in str(type_a_hp1["Mounting height"])
    assert type_a_hp1["Circuit description"] == "RECESSED FIXTURE"
    assert type_a_hp1["Effective quantity"] == 1

    duplex = line(**{"Panel": "LP1", "Circuit": "12"})
    assert duplex["Circuit description"] == "DUPLEX RECEPTACLE"
    assert duplex["Breaker"] == "20A"
    assert str(duplex["Poles"]) == "1"
    assert str(duplex["Load"]) == "180"
    assert duplex["Confidence"] == "high"
    assert duplex["Effective quantity"] == 1

    missing = line(**{"Panel": "LP1", "Circuit": "7"})
    assert missing["Circuit description"] is None
    assert missing["Confidence"] == "check"
    assert missing["Effective quantity"] == 1

    orphan = line(**{"Panel": "LP9", "Circuit": "1"})
    assert orphan["Confidence"] == "check"
    assert orphan["Circuit description"] is None

    notes = _kind_rows(workbook["Discrepancies"])
    by_kind = {}
    for note in notes:
        by_kind.setdefault(note["Kind"], []).append(note)

    assert "type tag with no schedule row" in by_kind
    missing_type = by_kind["type tag with no schedule row"][0]
    assert missing_type["Plan"] == 1
    assert missing_type["Schedule"] is None
    assert "Z" in missing_type["Detail"]

    zero = next(note for note in by_kind["schedule row with no plan marks"] if "type C" in note["Detail"])
    assert zero["Plan"] == 0
    assert zero["Schedule"] == 1
    mismatch_c = next(
        note
        for note in by_kind["plan count differs from schedule count"]
        if note["Detail"].startswith("type C ")
    )
    assert mismatch_c["Plan"] == 0
    assert mismatch_c["Schedule"] == 1

    mismatch_a = next(
        note
        for note in by_kind["plan count differs from schedule count"]
        if note["Detail"].startswith("type A ")
    )
    assert mismatch_a["Plan"] == 2
    assert mismatch_a["Schedule"] == 3
    assert "2" in mismatch_a["Detail"] and "3" in mismatch_a["Detail"]
    assert not any("type A" in note["Detail"] for note in by_kind["schedule row with no plan marks"])
    assert not any("type B" in note["Detail"] for note in by_kind.get("schedule row with no plan marks", []))
    assert not any("SPARE" in note["Detail"] or "circuit 15" in note["Detail"] for note in notes)

    height = by_kind["mounting height differs"][0]
    assert "7'-0" in str(height["Plan"])
    assert "9'-0" in str(height["Schedule"])

    missing_circuit = by_kind["circuit on plan missing from panel schedule"][0]
    assert "LP1" in missing_circuit["Detail"] and "7" in missing_circuit["Detail"]
    assert missing_circuit["Plan"] == 1
    assert missing_circuit["Schedule"] is None

    missing_panel = by_kind["panel schedule not found"][0]
    assert "LP9" in missing_panel["Detail"]

    rollup = _type_rows(workbook["By type"])
    assert workbook["By type"]["A1"].value.startswith("Each row is one fixture type")
    type_a_sheets = [row for row in rollup if row["Type"] == "A"]
    assert {row["Sheet number"] for row in type_a_sheets} == {"E1", "E2"}
    assert {row["Floor"] for row in type_a_sheets} == {"FIRST FLOOR"}
    assert {row["Plan count on sheet"] for row in type_a_sheets} == {1}
    assert {row["Plan count all sheets"] for row in type_a_sheets} == {2}
    assert {row["Schedule count"] for row in type_a_sheets} == {3}
    assert all(row["Type description"] == "6 INCH DOWNLIGHT" for row in type_a_sheets)
    assert all(row["Type"] != "C" for row in rollup)
    type_b_sheet = next(row for row in rollup if row["Type"] == "B")
    assert type_b_sheet["Plan count on sheet"] == 1
    assert type_b_sheet["Schedule count"] == 1
    assert type_b_sheet["Floor"] == "FIRST FLOOR"
    type_z_sheet = next(row for row in rollup if row["Type"] == "Z")
    assert type_z_sheet["Type description"] is None
    assert type_z_sheet["Schedule count"] is None

    lighting_effective = sum(
        row["Effective quantity"] for row in quantities if row["Category"] == "lighting"
    )
    assert lighting_effective == 4

    rejected = client.post(
        f"/jobs/{job_id}/items/{type_b['Item IDs']}/review",
        data={"action": "reject", "label": "RECESSED LIGHT", "panel": "", "circuit": ""},
        follow_redirects=True,
    )
    assert _marked(rejected.text, "lighting-fixture-count") == "3"
    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    rollup = _type_rows(workbook["By type"])
    type_b_sheet = next(row for row in rollup if row["Type"] == "B")
    assert type_b_sheet["Plan count on sheet"] == 1
    assert type_b_sheet["Plan count all sheets"] == 1
    notes = _kind_rows(workbook["Discrepancies"])
    assert not any(note["Kind"] == "schedule row with no plan marks" and "type B" in note["Detail"] for note in notes)

    markup = fitz.open(stream=client.get(f"/jobs/{job_id}/export.pdf").content, filetype="pdf")
    try:
        schedule_text = markup[3].get_text("text")
    finally:
        markup.close()
    assert "fixture schedule" in schedule_text
    assert "panel schedule" in schedule_text
