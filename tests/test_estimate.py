"""Estimator quantities beside the engine count. Synthetic files only."""

from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from adel_takeoff.estimate import (
    FUZZY_AUTO_THRESHOLD,
    EstimateError,
    apply_mapping,
    compare_estimate,
    parse_upload,
    set_pair,
)
from adel_takeoff.main import create_app
from tests.drawings import build_sample_pdf


def _engine(line_id, label, qty, **extra):
    return {
        "line_id": line_id,
        "label": label,
        "effective": qty,
        "sheet_number": extra.get("sheet_number", "E1"),
        "category": extra.get("category", "power"),
        "counts_in_base": extra.get("counts_in_base", True),
        "type_mark": extra.get("type_mark", ""),
        "type_description": extra.get("type_description", ""),
    }


def _store(rows, *, kind="quantities", pairs=None):
    stored = []
    if kind == "proposal":
        headers = ["line", "description", "qty", "unit price"]
        mapping = {"tag": 0, "description": 1, "quantity": 2, "unit_price": 3}
    else:
        headers = ["tag", "description", "quantity", "unit", "system", "sheet"]
        mapping = {"tag": 0, "description": 1, "quantity": 2, "unit": 3, "system": 4, "sheet": 5}
    for index, row in enumerate(rows, start=2):
        if kind == "proposal":
            cells = [
                row.get("tag", ""),
                row.get("description", ""),
                str(row.get("quantity", "")),
                str(row.get("unit_price", "")),
            ]
        else:
            cells = [
                row.get("tag", ""),
                row.get("description", ""),
                str(row.get("quantity", "")),
                row.get("unit", "EA"),
                row.get("system", "power"),
                row.get("sheet", ""),
            ]
        stored.append({"n": index, "cells": cells})
    return {
        "kind": kind,
        "source_name": "estimate.csv",
        "headers": headers,
        "rows": stored,
        "mapping": mapping,
        "pairs": pairs or {},
    }


def _by_tag(result, tag):
    return next(row for row in result["rows"] if row["tag"] == tag)


def test_quantities_wait_for_a_column_map_and_reject_a_bad_number():
    payload = (
        "Item / tag,Description,Quantity,Unit,System,Sheet\n"
        "DUPLEX,Duplex receptacle,3,EA,power,E1\n"
    ).encode()
    store = parse_upload("quantities", "takeoff.csv", payload)
    assert store["mapping"] is None
    assert store["headers"][0] == "Item / tag"
    compared = compare_estimate(store, [_engine("a", "DUPLEX", 3)])
    assert compared["mapped"] is False
    assert compared["rows"] == []
    mapped = apply_mapping(
        store,
        {"tag": "0", "description": "1", "quantity": "2", "unit": "3", "system": "4", "sheet": "5"},
    )
    assert mapped["mapping"]["sheet"] == 5
    assert mapped["pairs"] == {}
    with pytest.raises(EstimateError, match="Each column can be used once"):
        apply_mapping(store, {"tag": "0", "description": "0", "quantity": "2", "unit": "3", "system": "4"})
    bad = (
        "Item,Description,Quantity,Unit,System\n"
        "\n"
        "A,Widget,nope,EA,power\n"
    ).encode()
    waiting = parse_upload("quantities", "bad.csv", bad)
    with pytest.raises(EstimateError, match="Row 3 quantity is not a number"):
        apply_mapping(
            waiting,
            {"tag": "0", "description": "1", "quantity": "2", "unit": "3", "system": "4"},
        )


def test_excel_quantities_use_the_same_mapping_step():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Mark", "Desc", "Count", "UOM", "Sys", "Dwg"])
    sheet.append(["SWITCH", "A switch", 2, "EA", "power", "E2"])
    buffer = BytesIO()
    workbook.save(buffer)
    store = parse_upload("quantities", "takeoff.xlsx", buffer.getvalue())
    assert store["mapping"] is None
    mapped = apply_mapping(
        store,
        {"tag": "0", "description": "1", "quantity": "2", "unit": "3", "system": "4", "sheet": "5"},
    )
    result = compare_estimate(mapped, [_engine("s", "SWITCH", 2, sheet_number="E2")])
    row = _by_tag(result, "SWITCH")
    assert row["method"] == "tag"
    assert row["status"] == "match"
    assert row["review"] == "firm"
    assert row["sheets"] == "E2"
    assert row["estimator_sheet"] == "E2"


def test_tag_description_and_type_mark_pair_in_that_order():
    engines = [
        _engine("d1", "DUPLEX RECEPTACLE", 4, sheet_number="E1"),
        _engine("r1", "RECESSED LIGHT", 2, type_mark="L1", sheet_number="E3"),
    ]
    result = compare_estimate(
        _store(
            [
                {
                    "tag": "DUPLEX RECEPTACLE",
                    "description": "not the same words",
                    "quantity": "2",
                    "sheet": "E9",
                },
                {"tag": "L1", "description": "downlight", "quantity": "2"},
                {"tag": "A1", "description": "Recessed light", "quantity": "1"},
            ]
        ),
        engines,
    )
    duplex = _by_tag(result, "DUPLEX RECEPTACLE")
    assert duplex["method"] == "tag"
    assert duplex["review"] == "firm"
    assert duplex["status"] == "engine higher"
    assert duplex["delta"] == "2"
    assert duplex["estimator_qty"] == "2"
    assert duplex["engine_qty"] == "4"
    assert duplex["sheets"] == "E1"
    marked = _by_tag(result, "L1")
    assert marked["method"] == "tag"
    assert marked["engine_line_id"] == "r1"
    assert marked["status"] == "match"
    described = next(row for row in result["rows"] if row["tag"] == "A1")
    assert described["method"] == "only"
    assert described["status"] == "only in estimator"
    assert described["review"] == "firm"


def test_description_matches_when_the_tag_does_not():
    result = compare_estimate(
        _store(
            [
                {
                    "tag": "A1",
                    "description": "Duplex receptacle",
                    "quantity": "3",
                }
            ]
        ),
        [_engine("d1", "DUPLEX RECEPTACLE", 3, type_description="Duplex receptacle")],
    )
    row = _by_tag(result, "A1")
    assert row["method"] == "description"
    assert row["review"] == "firm"
    assert row["status"] == "match"
    assert row["engine_line_id"] == "d1"


def test_fuzzy_below_the_threshold_is_not_auto_paired():
    words = "alpha bravo charlie delta echo foxtrot golf hotel".split()
    result = compare_estimate(
        _store(
            [
                {
                    "tag": "",
                    "description": " ".join(words),
                    "quantity": "2",
                    "unit_price": None,
                }
            ]
        ),
        [_engine("long", " ".join(words + ["india", "juliet"]), 1, sheet_number="E5")],
    )
    row = result["rows"][0]
    assert FUZZY_AUTO_THRESHOLD == Decimal("0.90")
    assert row["method"] == "suggestion"
    assert row["review"] == "needs-review"
    assert row["score"] == "0.80"
    assert row["engine_line_id"] == "long"
    assert row["engine_qty"] == "1"
    assert row["estimator_qty"] == "2"
    assert row["status"] == "estimator higher"
    assert row["sheets"] == "E5"
    assert [item["status"] for item in result["rows"]].count("only in engine") == 0


def test_fuzzy_at_the_threshold_is_auto_paired_when_it_is_unique():
    words = "alpha bravo charlie delta echo foxtrot golf hotel india".split()
    result = compare_estimate(
        _store([{"tag": "", "description": " ".join(words), "quantity": "1"}]),
        [_engine("long", " ".join(words + ["juliet"]), 1, sheet_number="E6")],
    )
    row = result["rows"][0]
    assert row["method"] == "fuzzy"
    assert row["review"] == "firm"
    assert row["score"] == "0.90"
    assert row["status"] == "match"
    assert row["sheets"] == "E6"


def test_a_fuzzy_tie_is_not_auto_paired():
    words = "alpha bravo charlie delta echo foxtrot golf hotel india".split()
    result = compare_estimate(
        _store([{"tag": "", "description": " ".join(words), "quantity": "1"}]),
        [
            _engine("b", " ".join(words + ["juliet"]), 1),
            _engine("a", " ".join(words + ["kilo"]), 1),
        ],
    )
    row = next(item for item in result["rows"] if item["estimator_id"] == "e1")
    assert row["method"] == "ambiguous"
    assert row["review"] == "needs-review"
    assert row["engine_line_id"] == ""
    assert row["status"] == "only in estimator"
    assert {item["line_id"] for item in row["candidates"]} == {"a", "b"}
    engine_only = [item for item in result["rows"] if item["status"] == "only in engine"]
    assert [item["engine_line_id"] for item in engine_only] == ["a", "b"]
    assert {item["review"] for item in engine_only} == {"firm"}


def test_an_ambiguous_tag_is_not_auto_paired():
    result = compare_estimate(
        _store([{"tag": "DUPLEX RECEPTACLE", "description": "Duplex", "quantity": "3"}]),
        [
            _engine("b", "DUPLEX RECEPTACLE", 1, sheet_number="E2"),
            _engine("a", "DUPLEX RECEPTACLE", 1, sheet_number="E1"),
        ],
    )
    row = next(item for item in result["rows"] if item["estimator_id"] == "e1")
    assert row["method"] == "ambiguous"
    assert row["review"] == "needs-review"
    assert row["status"] == "only in estimator"
    assert row["engine_qty"] == ""
    assert [item["line_id"] for item in row["candidates"]] == ["a", "b"]
    store = _store([{"tag": "DUPLEX RECEPTACLE", "description": "Duplex", "quantity": "3"}])
    paired_store = set_pair(store, "e1", "a", {"a", "b"})
    paired = next(
        item
        for item in compare_estimate(paired_store, [
            _engine("b", "DUPLEX RECEPTACLE", 1, sheet_number="E2"),
            _engine("a", "DUPLEX RECEPTACLE", 1, sheet_number="E1"),
        ])["rows"]
        if item["estimator_id"] == "e1"
    )
    assert paired["method"] == "manual"
    assert paired["review"] == "firm"
    assert paired["engine_line_id"] == "a"
    assert paired["sheets"] == "E1"
    assert paired["status"] == "estimator higher"


def test_manual_pair_is_firm_and_survives_a_remap():
    engines = [
        _engine("keep", "PANELBOARD", 1, sheet_number="E1"),
        _engine("other", "SWITCH", 1),
    ]
    store = _store(
        [{"tag": "CUSTOM", "description": "no shared tokens zz", "quantity": "4"}],
        pairs={"e1": "keep"},
    )
    result = compare_estimate(store, engines)
    row = _by_tag(result, "CUSTOM")
    assert row["method"] == "manual"
    assert row["review"] == "firm"
    assert row["status"] == "estimator higher"
    assert row["delta"] == "-3"
    assert row["sheets"] == "E1"
    remapped = apply_mapping(
        store,
        {"tag": "0", "description": "1", "quantity": "2", "unit": "3", "system": "4", "sheet": "5"},
    )
    assert remapped["pairs"] == {"e1": "keep"}
    again = compare_estimate(remapped, list(reversed(engines)))
    assert _by_tag(again, "CUSTOM")["method"] == "manual"
    assert compare_estimate(remapped, engines)["rows"] == again["rows"]


def test_proposal_price_delta_splits_firm_and_needs_review():
    payload = (
        "Line,Description,Qty,Unit Price\n"
        'PANELBOARD,Panel,"2","$100.00"\n'
        "alpha,alpha bravo charlie delta echo foxtrot golf hotel,2,4.00\n"
    ).encode()
    store = parse_upload("proposal", "proposal.csv", payload)
    assert store["mapping"]["unit_price"] == 3
    words = "alpha bravo charlie delta echo foxtrot golf hotel".split()
    result = compare_estimate(
        store,
        [
            _engine("p", "PANELBOARD", 3, sheet_number="E1"),
            _engine("long", " ".join(words + ["india", "juliet"]), 1, sheet_number="E8"),
        ],
    )
    panel = _by_tag(result, "PANELBOARD")
    assert panel["method"] == "tag"
    assert panel["review"] == "firm"
    assert panel["status"] == "engine higher"
    assert panel["delta"] == "1"
    assert panel["unit"] == "EA"
    assert panel["unit_price"] == "100.00"
    assert panel["price_delta"] == "100.00"
    assert panel["sheets"] == "E1"
    suggested = next(row for row in result["rows"] if row["method"] == "suggestion")
    assert suggested["review"] == "needs-review"
    assert suggested["score"] == "0.80"
    assert suggested["price_delta"] == "-4.00"
    assert result["firm_price_delta"] == "100.00"
    assert result["review_price_delta"] == "-4.00"
    assert result["threshold"] == "0.90"


def test_proposal_headers_that_are_not_unique_wait_for_a_map():
    payload = b"Line,Item,Description,Qty,Unit Price\nA,B,C,1,2\n"
    store = parse_upload("proposal", "proposal.csv", payload)
    assert store["mapping"] is None
    with pytest.raises(EstimateError, match="proposal import is a CSV"):
        parse_upload("proposal", "proposal.xlsx", payload)


def test_a_held_zone_line_is_not_an_engine_line():
    result = compare_estimate(
        _store([{"tag": "LANDSCAPE LIGHT", "description": "Landscape", "quantity": "1"}]),
        [_engine("z", "LANDSCAPE LIGHT", 1, counts_in_base=False)],
    )
    row = _by_tag(result, "LANDSCAPE LIGHT")
    assert row["status"] == "only in estimator"
    assert row["review"] == "firm"
    assert row["engine_line_id"] == ""


def test_comparison_order_is_deterministic():
    store = _store(
        [
            {"tag": "MATCH-B", "description": "b", "quantity": "1"},
            {"tag": "MATCH-A", "description": "a", "quantity": "1"},
            {"tag": "NOPE", "description": "qq", "quantity": "1"},
        ]
    )
    engines = [
        _engine("z-only", "ZULU", 2),
        _engine("a-only", "ALPHA", 1, sheet_number="E4"),
        _engine("mb", "MATCH-B", 1),
        _engine("ma", "MATCH-A", 1),
    ]
    first = compare_estimate(store, engines)
    second = compare_estimate(store, list(reversed(engines)))
    assert [row["tag"] for row in first["rows"]] == [
        "MATCH-B",
        "MATCH-A",
        "NOPE",
        "ALPHA",
        "ZULU",
    ]
    assert first["rows"] == second["rows"]
    assert [row["status"] for row in first["rows"][:2]] == ["match", "match"]
    assert first["rows"][2]["status"] == "only in estimator"
    assert first["rows"][3]["status"] == "only in engine"
    assert first["rows"][3]["sheets"] == "E4"


def test_sample_job_import_leaves_the_engine_count_and_writes_excel(tmp_path):
    client = TestClient(create_app(tmp_path))
    uploaded = client.post(
        "/jobs",
        files={"pdf": ("set.pdf", build_sample_pdf(), "application/pdf")},
        follow_redirects=False,
    )
    job_id = uploaded.headers["location"].rstrip("/").split("/")[-1]
    before = client.get(f"/jobs/{job_id}/review").text
    assert 'id="power-device-count">4<' in before
    empty = client.get(f"/jobs/{job_id}/estimate").text
    assert 'id="estimate-empty"' in empty
    assert "0.90" in empty
    quantities = (
        "Item / tag,Description,Quantity,Unit,System,Sheet\n"
        "SINGLE POLE SWITCH,Switch,2,EA,power,E1\n"
        "DUPLEX RECEPTACLE,Duplex,3,EA,power,E1\n"
        "NOBODY,Unlisted feeder,4,EA,power,E9\n"
    ).encode()
    imported = client.post(
        f"/jobs/{job_id}/estimate",
        data={"kind": "quantities"},
        files={"file": ("takeoff.csv", quantities, "text/csv")},
        follow_redirects=False,
    )
    assert imported.status_code == 303
    mapping = client.get(f"/jobs/{job_id}/estimate").text
    assert 'id="mapping-form"' in mapping
    assert 'id="estimate-comparison"' not in mapping
    assert "SINGLE POLE SWITCH" in mapping
    saved = client.post(
        f"/jobs/{job_id}/estimate/map",
        data={"tag": "0", "description": "1", "quantity": "2", "unit": "3", "system": "4", "sheet": "5"},
        follow_redirects=False,
    )
    assert saved.status_code == 303
    page = client.get(f"/jobs/{job_id}/estimate").text
    assert 'data-status="estimator higher"' in page
    assert 'data-status="only in estimator"' in page
    assert 'data-review="needs-review"' in page
    assert 'data-method="tag"' in page
    assert 'data-method="ambiguous"' in page
    assert 'id="firm-price-delta">0.00<' in page
    assert 'id="review-price-delta">0.00<' in page
    assert 'id="power-device-count">4<' in client.get(f"/jobs/{job_id}/review").text

    proposal = (
        "Line,Description,Qty,Unit Price\n"
        "SINGLE POLE SWITCH,Switch,2,100.00\n"
        "DUPLEX RECEPTACLE,Duplex,3,15.00\n"
        "NOBODY,Unlisted feeder,4,10.00\n"
    ).encode()
    priced = client.post(
        f"/jobs/{job_id}/estimate",
        data={"kind": "proposal"},
        files={"file": ("proposal.csv", proposal, "text/csv")},
        follow_redirects=False,
    )
    assert priced.status_code == 303
    priced_page = client.get(f"/jobs/{job_id}/estimate").text
    assert 'id="mapping-form"' not in priced_page
    assert 'id="firm-price-delta">-140.00<' in priced_page
    assert 'id="review-price-delta">-45.00<' in priced_page
    assert "100.00" in priced_page
    select = priced_page.split('aria-label="Pair DUPLEX RECEPTACLE"', 1)[1].split("</select>", 1)[0]
    line_id = next(value for value in _values(select) if value)
    paired = client.post(
        f"/jobs/{job_id}/estimate/pair",
        data={"estimator_id": "e2", "engine_line_id": line_id},
        follow_redirects=False,
    )
    assert paired.status_code == 303
    confirmed = client.get(f"/jobs/{job_id}/estimate").text
    assert 'data-method="manual"' in confirmed
    assert "Clear pair" in confirmed
    assert 'id="power-device-count">4<' in client.get(f"/jobs/{job_id}/review").text

    workbook = openpyxl.load_workbook(BytesIO(client.get(f"/jobs/{job_id}/export.xlsx").content))
    assert "Estimator vs Engine" in workbook.sheetnames
    assert "Routing" not in workbook.sheetnames
    assert workbook.properties.title == "Quantity takeoff"
    quantities_sheet = workbook["Quantities"]
    quantity_text = " ".join(
        str(cell.value) for row in quantities_sheet.iter_rows() for cell in row if cell.value
    )
    assert "How counted" in quantity_text
    sheet = workbook["Estimator vs Engine"]
    assert sheet.cell(3, 1).value == "Estimator tag"
    assert sheet.cell(3, 9).value == "Status"
    assert sheet.cell(3, 10).value == "Firm or needs-review"
    assert sheet.cell(3, 11).value == "Engine sheet refs"
    assert "0.90" in sheet["A1"].value
    totals = next(row for row in sheet.iter_rows(values_only=True) if row[0] == "Totals")
    assert totals[1] == "Firm price delta"
    assert totals[3] == "Needs-review price delta"
    assert Decimal(str(totals[2])) != Decimal("0") or Decimal(str(totals[4])) != Decimal("0")
    statuses = {
        row[8]
        for row in sheet.iter_rows(min_row=4, values_only=True)
        if row[0] not in {None, "Totals", "No estimator import"}
    }
    assert "only in engine" in statuses
    assert "match" in statuses or "engine higher" in statuses or "estimator higher" in statuses


def _values(html: str) -> list[str]:
    import re

    return re.findall(r'value="([^"]*)"', html)
