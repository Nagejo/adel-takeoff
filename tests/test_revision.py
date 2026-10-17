"""Sheet matching, count deltas, review carry-forward, and alignment refusal."""

from io import BytesIO

import re

import fitz
import openpyxl
from fastapi.testclient import TestClient

from adel_takeoff.main import create_app
from adel_takeoff.revision import (
    align_pages,
    build_overlay,
    diff_counts,
    diff_schedules,
    match_sheets,
)
from adel_takeoff.review_store import load_document, review_path
from tests.drawings import _text, draw_duplex, draw_switch


def _sheet(page_index, number, title=""):
    return {"page_index": page_index, "number": number, "title": title}


def test_sheets_match_on_number_then_title_and_report_the_rest():
    old = [
        _sheet(0, "E1", "POWER PLAN"),
        _sheet(1, "E2", "LIGHTING PLAN"),
        _sheet(2, "E3", "PANEL SCHEDULE"),
    ]
    new = [
        _sheet(0, "E1", "POWER PLAN"),
        _sheet(1, "E2A", "LIGHTING PLAN"),
        _sheet(2, "E4", "DETAILS"),
    ]
    pairs = match_sheets(old, new)
    by_status = {pair["status"]: pair for pair in pairs}
    assert by_status["matched"]["old_number"] == "E1"
    assert by_status["matched"]["new_number"] == "E1"
    assert by_status["renumbered"]["old_number"] == "E2"
    assert by_status["renumbered"]["new_number"] == "E2A"
    assert by_status["renumbered"]["new_title"] == "LIGHTING PLAN"
    assert by_status["added"]["new_number"] == "E4"
    assert by_status["removed"]["old_number"] == "E3"


def test_duplicate_sheet_numbers_are_not_guessed():
    old = [
        _sheet(0, "E1", "POWER PLAN"),
        _sheet(1, "E1", "POWER PLAN"),
    ]
    new = [_sheet(0, "E1", "POWER PLAN")]
    pairs = match_sheets(old, new)
    assert [pair["status"] for pair in pairs].count("matched") == 0
    assert [pair["status"] for pair in pairs].count("renumbered") == 0
    assert [pair["status"] for pair in pairs].count("removed") == 2
    assert [pair["status"] for pair in pairs].count("added") == 1


def _item(item_id, label, page, sheet, x, y, panel="", circuit="", grid="A1"):
    return {
        "id": item_id,
        "label": label,
        "panel": panel,
        "circuit": circuit,
        "measurement": "matched",
        "measured_quantity": 1,
        "occurrences": [
            {
                "page_index": page,
                "sheet": sheet,
                "sheet_title": "POWER PLAN",
                "x0": x - 8,
                "y0": y - 8,
                "x1": x + 8,
                "y1": y + 8,
                "grid": grid,
            }
        ],
    }


def _matched_sheet():
    return {
        "status": "matched",
        "old_page_index": 0,
        "new_page_index": 0,
        "old_number": "E1",
        "new_number": "E1",
        "old_title": "POWER PLAN",
        "new_title": "POWER PLAN",
        "alignment": {
            "ok": True,
            "reason": "",
            "confidence": 0.99,
            "dx": 0,
            "dy": 0,
            "scale_x": 1,
            "scale_y": 1,
        },
    }


def _decision(status="accepted"):
    return {
        "corrected_circuit": None,
        "corrected_label": None,
        "corrected_panel": None,
        "exclusion": "",
        "review_status": status,
    }


def test_count_deltas_and_carry_forward_only_unchanged_items():
    old = [
        _item("1", "DUPLEX RECEPTACLE", 0, "E1", 100, 100, "LP1", "1", "A1"),
        _item("2", "SINGLE POLE SWITCH", 0, "E1", 160, 100, "LP1", "2", "A2"),
        _item("3", "RECESSED LIGHT", 0, "E1", 220, 180, "LP1", "3", "B1"),
    ]
    new = [
        _item("8", "DUPLEX RECEPTACLE", 0, "E1", 100, 102, "LP1", "1", "A1"),
        _item("9", "SINGLE POLE SWITCH", 0, "E1", 160, 112, "LP1", "9", "A2"),
        _item("4", "DUPLEX RECEPTACLE", 0, "E1", 300, 200, "LP1", "4", "C3"),
    ]
    decisions = {"1": _decision("accepted"), "2": _decision("rejected"), "3": _decision("accepted")}
    result = diff_counts(old, new, [_matched_sheet()], decisions)
    by_new = {link["new_id"]: link for link in result["links"]}
    assert by_new["8"]["old_id"] == "1"
    assert by_new["8"]["unchanged"] is True
    assert by_new["9"]["old_id"] == "2"
    assert by_new["9"]["unchanged"] is False
    assert result["id_map"]["8"] == "1"
    assert result["id_map"]["9"] == "2"
    assert result["id_map"]["4"] not in {"1", "2", "3"}
    assert result["kept_decisions"]["1"]["review_status"] == "accepted"
    assert "2" not in result["kept_decisions"]
    assert "3" not in result["kept_decisions"]
    assert "1" not in result["revision_checks"]
    assert result["id_map"]["9"] in result["revision_checks"]
    assert result["id_map"]["4"] in result["revision_checks"]
    rows = [row for row in result["changes"] if row["source"] == "takeoff"]
    recessed = next(row for row in rows if row["tag"] == "RECESSED LIGHT")
    assert recessed["old_qty"] == 1
    assert recessed["new_qty"] == 0
    assert recessed["delta"] == -1
    assert recessed["sheet"] == "E1"
    added = next(row for row in rows if row["area"] == "C3")
    assert added["old_qty"] == 0
    assert added["new_qty"] == 1
    assert added["delta"] == 1
    moved = next(row for row in rows if row["tag"] == "SINGLE POLE SWITCH")
    assert moved["old_qty"] == 1
    assert moved["new_qty"] == 1
    assert moved["delta"] == 0
    assert "panel/circuit" in moved["note"]


def test_refused_alignment_does_not_emit_a_count_diff():
    old = [_item("1", "DUPLEX RECEPTACLE", 0, "E1", 100, 100)]
    new = [_item("7", "DUPLEX RECEPTACLE", 0, "E1", 280, 200)]
    sheet = _matched_sheet()
    sheet["alignment"] = {
        "ok": False,
        "reason": "Alignment refused: agreement 0.10 is below 0.72. No overlay was drawn.",
        "confidence": 0.10,
        "dx": 0,
        "dy": 0,
        "scale_x": 1,
        "scale_y": 1,
    }
    result = diff_counts(old, new, [sheet], {"1": _decision()})
    assert result["links"] == []
    assert result["kept_decisions"] == {}
    takeoff = [row for row in result["changes"] if row["source"] == "takeoff"]
    assert takeoff == []
    refused = next(row for row in result["changes"] if row["source"] == "alignment")
    assert "not compared" in refused["note"]
    assert "7" in result["revision_checks"] or result["id_map"]["7"] in result["revision_checks"]


def _page(width, height, draw) -> bytes:
    document = fitz.open()
    page = document.new_page(width=width, height=height)
    draw(page)
    payload = document.tobytes()
    document.close()
    return payload


def _plan(page, ox=0, oy=0, number="E1", title="POWER PLAN", extra=False):
    _text(page, 24 + ox, 28 + oy, f"SHEET {number}")
    _text(page, 24 + ox, 46 + oy, f"TITLE {title}")
    page.draw_rect(
        fitz.Rect(80 + ox, 90 + oy, 170 + ox, 150 + oy),
        color=(0, 0, 0),
        fill=(0, 0, 0),
        width=1,
    )
    page.draw_circle(fitz.Point(230 + ox, 170 + oy), 16, color=(0, 0, 0), fill=(0, 0, 0), width=1)
    if extra:
        page.draw_rect(
            fitz.Rect(250 + ox, 70 + oy, 310 + ox, 110 + oy),
            color=(0, 0, 0),
            fill=(0, 0, 0),
            width=1,
        )


def _one(draw, width=360, height=260) -> bytes:
    return _page(width, height, draw)


def test_shifted_sheet_aligns_and_a_far_shift_or_blank_is_refused():
    old = _one(lambda page: _plan(page))
    shifted = _one(lambda page: _plan(page, ox=12, oy=8))
    aligned = align_pages(old, shifted, match_sheets(
        [{"page_index": 0, "number": "E1", "title": "POWER PLAN"}],
        [{"page_index": 0, "number": "E1", "title": "POWER PLAN"}],
    ))
    alignment = aligned[0]["alignment"]
    assert alignment["ok"] is True
    assert abs(alignment["dx"] - 12) <= 2
    assert abs(alignment["dy"] - 8) <= 2
    assert alignment["confidence"] >= 0.72

    far = _one(lambda page: _plan(page, ox=200, oy=0))
    refused = align_pages(old, far, match_sheets(
        [{"page_index": 0, "number": "E1", "title": "POWER PLAN"}],
        [{"page_index": 0, "number": "E1", "title": "POWER PLAN"}],
    ))[0]["alignment"]
    assert refused["ok"] is False
    assert "below" in refused["reason"]
    assert "No overlay was drawn" in refused["reason"]

    blank = _one(lambda page: None)
    no_ink = align_pages(old, blank, match_sheets(
        [{"page_index": 0, "number": "E1", "title": "POWER PLAN"}],
        [{"page_index": 0, "number": "E1", "title": "POWER PLAN"}],
    ))[0]["alignment"]
    assert no_ink["ok"] is False
    assert "no ink" in no_ink["reason"]

    tall = _page(360, 420, lambda page: _plan(page))
    proportions = align_pages(old, tall, match_sheets(
        [{"page_index": 0, "number": "E1", "title": "POWER PLAN"}],
        [{"page_index": 0, "number": "E1", "title": "POWER PLAN"}],
    ))[0]["alignment"]
    assert proportions["ok"] is False
    assert "proportions" in proportions["reason"]

    overlay = build_overlay(old, blank, align_pages(old, blank, match_sheets(
        [{"page_index": 0, "number": "E1", "title": "POWER PLAN"}],
        [{"page_index": 0, "number": "E1", "title": "POWER PLAN"}],
    )))
    document = fitz.open(stream=overlay, filetype="pdf")
    try:
        assert document.page_count == 1
        text = document[0].get_text()
        assert "No overlay was drawn" in text
        assert document[0].get_drawings() == []
    finally:
        document.close()


def test_overlay_colors_additions_and_deletions():
    old = _one(lambda page: _plan(page))
    new = _one(lambda page: _plan(page, extra=True))
    # Remove the circle from the new sheet by drawing the plan without it, then the extra.
    document = fitz.open()
    page = document.new_page(width=360, height=260)
    _text(page, 24, 28, "SHEET E1")
    _text(page, 24, 46, "TITLE POWER PLAN")
    page.draw_rect(fitz.Rect(80, 90, 170, 150), color=(0, 0, 0), fill=(0, 0, 0), width=1)
    page.draw_rect(fitz.Rect(250, 70, 310, 110), color=(0, 0, 0), fill=(0, 0, 0), width=1)
    new = document.tobytes()
    document.close()
    sheets = align_pages(old, new, match_sheets(
        [{"page_index": 0, "number": "E1", "title": "POWER PLAN"}],
        [{"page_index": 0, "number": "E1", "title": "POWER PLAN"}],
    ))
    assert sheets[0]["alignment"]["ok"] is True
    overlay = build_overlay(old, new, sheets)
    rendered = fitz.open(stream=overlay, filetype="pdf")
    try:
        pixmap = rendered[0].get_pixmap(colorspace=fitz.csRGB, alpha=False)
        samples = pixmap.samples
        blue = 0
        red = 0
        for index in range(0, len(samples), 3):
            red_channel, green_channel, blue_channel = samples[index], samples[index + 1], samples[index + 2]
            if blue_channel > 140 and blue_channel > red_channel + 40:
                blue += 1
            if red_channel > 140 and red_channel > blue_channel + 40:
                red += 1
        assert blue > 30
        assert red > 30
        assert rendered[0].get_drawings()
    finally:
        rendered.close()


def _schedule(count: str, description: str, breaker: str) -> bytes:
    document = fitz.open()
    page = document.new_page(width=720, height=400)
    _text(page, 40, 36, "SHEET E3")
    _text(page, 40, 54, "TITLE FIXTURE AND PANEL SCHEDULES")
    _text(page, 40, 90, "LIGHTING FIXTURE SCHEDULE")
    _text(page, 40, 112, "TYPE")
    _text(page, 120, 112, "DESCRIPTION")
    _text(page, 280, 112, "QTY")
    _text(page, 40, 132, "A")
    _text(page, 120, 132, description)
    _text(page, 280, 132, count)
    _text(page, 40, 180, "PANEL SCHEDULE LP1")
    _text(page, 40, 202, "CIRCUIT")
    _text(page, 140, 202, "DESCRIPTION")
    _text(page, 300, 202, "BREAKER")
    _text(page, 40, 222, "1")
    _text(page, 140, 222, "RECEPTACLES")
    _text(page, 300, 222, breaker)
    payload = document.tobytes()
    document.close()
    return payload


def test_schedule_rows_report_quantity_deltas():
    old = _schedule("4", "DOWNLIGHT", "20A")
    new = _schedule("6", "DOWNLIGHT", "30A")
    rows = diff_schedules(old, new)
    fixture = next(row for row in rows if row["source"] == "fixture schedule")
    panel = next(row for row in rows if row["source"] == "panel schedule")
    assert fixture["tag"] == "A"
    assert fixture["old_qty"] == 4
    assert fixture["new_qty"] == 6
    assert fixture["delta"] == 2
    assert panel["old_qty"] is None
    assert "breaker" in panel["note"]
    assert panel["tag"] == "LP1 / 1"


def _legend_set(move_switch=False, add_duplex=False) -> bytes:
    document = fitz.open()
    legend = document.new_page(width=500, height=360)
    _text(legend, 36, 28, "SHEET E0")
    _text(legend, 36, 46, "TITLE SYMBOL LEGEND")
    _text(legend, 36, 72, "SYMBOL LEGEND")
    draw_duplex(legend, 70, 110)
    _text(legend, 100, 114, "DUPLEX RECEPTACLE")
    draw_switch(legend, 70, 160)
    _text(legend, 100, 164, "SINGLE POLE SWITCH")
    plan = document.new_page(width=500, height=360)
    _text(plan, 36, 28, "SHEET E1")
    _text(plan, 36, 46, "TITLE POWER PLAN")
    draw_duplex(plan, 180, 160)
    _text(plan, 196, 150, "LP1-1")
    switch_y = 172 if move_switch else 160
    draw_switch(plan, 280, switch_y)
    _text(plan, 296, 150, "LP1-2")
    if add_duplex:
        draw_duplex(plan, 180, 260)
        _text(plan, 196, 250, "LP1-4")
    payload = document.tobytes()
    document.close()
    return payload


def test_revision_page_keeps_an_unchanged_decision_and_writes_excel(tmp_path):
    client = TestClient(create_app(tmp_path))
    uploaded = client.post(
        "/jobs",
        files={"pdf": ("previous.pdf", _legend_set(), "application/pdf")},
        follow_redirects=False,
    )
    assert uploaded.status_code == 303
    job_id = uploaded.headers["location"].rstrip("/").split("/")[-1]
    review = client.get(f"/jobs/{job_id}/review")
    assert "Revision" in review.text
    duplex_id = re.search(
        r'id="item-([^"]+)"[^>]*data-label="DUPLEX RECEPTACLE"',
        review.text,
    ).group(1)
    switch_id = re.search(
        r'id="item-([^"]+)"[^>]*data-label="SINGLE POLE SWITCH"',
        review.text,
    ).group(1)
    accepted = client.post(
        f"/jobs/{job_id}/items/{duplex_id}/review",
        data={"action": "accept"},
        follow_redirects=False,
    )
    assert accepted.status_code == 303
    rejected = client.post(
        f"/jobs/{job_id}/items/{switch_id}/review",
        data={"action": "reject"},
        follow_redirects=False,
    )
    assert rejected.status_code == 303
    compared = client.post(
        f"/jobs/{job_id}/revision",
        files={"pdf": ("revised.pdf", _legend_set(move_switch=True, add_duplex=True), "application/pdf")},
        follow_redirects=False,
    )
    assert compared.status_code == 303
    page = client.get(f"/jobs/{job_id}/revision")
    assert page.status_code == 200
    assert "Aligned" in page.text
    assert "Alignment refused" not in page.text
    document = load_document(review_path(tmp_path, job_id))
    assert document["decisions"][duplex_id]["review_status"] == "accepted"
    assert switch_id not in document["decisions"]
    assert switch_id in document["revision_checks"]
    assert duplex_id not in document["revision_checks"]
    workbook_response = client.get(f"/jobs/{job_id}/export.xlsx")
    workbook = openpyxl.load_workbook(BytesIO(workbook_response.content))
    assert "Revision Changes" in workbook.sheetnames
    assert "Routing" not in workbook.sheetnames
    sheet = workbook["Revision Changes"]
    text = "\n".join(
        " ".join("" if cell.value is None else str(cell.value) for cell in row)
        for row in sheet.iter_rows(max_row=20)
    )
    assert "Old qty" in text
    assert "Delta" in text
    assert "takeoff" in text
    overlay = client.get(f"/jobs/{job_id}/revision/overlay.pdf")
    assert overlay.status_code == 200
    assert overlay.content.startswith(b"%PDF")
